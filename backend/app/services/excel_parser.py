"""Excel/CSV 解析：表头探测、样例提取、整表读取。"""
from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path

import openpyxl

MAX_ROWS = 100_000
SAMPLE_SIZE = 20


class ExcelParseError(Exception):
    pass


def cell_str(v) -> str:
    """单元格值转字符串（JSON 安全）。"""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def _read_csv(path: Path) -> dict[str, list[list]]:
    for enc in ("utf-8-sig", "gbk"):
        try:
            with open(path, newline="", encoding=enc) as f:
                rows = [row for row in csv.reader(f)]
            return {path.stem: rows}
        except UnicodeDecodeError:
            continue
    raise ExcelParseError("无法识别 CSV 文件编码（仅支持 UTF-8 / GBK）")


def _xlsx_sheet_rows(ws) -> list[list]:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    # 合并单元格：用左上角值填充整个区域，保证表头探测和数据对齐
    for mr in ws.merged_cells.ranges:
        tl = ws.cell(mr.min_row, mr.min_col).value
        for r in range(mr.min_row - 1, min(mr.max_row, len(rows))):
            for c in range(mr.min_col - 1, mr.max_col):
                if c < len(rows[r]) and rows[r][c] is None:
                    rows[r][c] = tl
    return rows


def _read_xlsx(path: Path) -> dict[str, list[list]]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise ExcelParseError(f"Excel 文件解析失败：{e}")
    return {name: _xlsx_sheet_rows(wb[name]) for name in wb.sheetnames}


def read_all_sheets(path: Path) -> dict[str, list[list]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv(path)
    if suffix == ".xls":
        raise ExcelParseError("暂不支持旧版 .xls 格式，请用 Excel/WPS 另存为 .xlsx 后重新上传")
    return _read_xlsx(path)


def read_sheet(path: Path, sheet_name: str) -> list[list]:
    sheets = read_all_sheets(path)
    if sheet_name not in sheets:
        raise ExcelParseError(f"工作表不存在：{sheet_name}")
    return trim_rows(sheets[sheet_name])


def trim_rows(rows: list[list]) -> list[list]:
    """裁掉全空的尾部行和尾部空列，补齐短行。"""
    if not rows:
        return rows
    width = 0
    for row in rows:
        for i in range(len(row) - 1, -1, -1):
            if row[i] not in (None, ""):
                width = max(width, i + 1)
                break
    out = []
    for row in rows:
        row = (row + [None] * width)[:width]
        out.append(row)
    while out and all(v in (None, "") for v in out[-1]):
        out.pop()
    if len(out) > MAX_ROWS:
        raise ExcelParseError(f"行数超过上限 {MAX_ROWS}，请拆分文件后导入")
    return out


def detect_header_row(rows: list[list], max_scan: int = 10) -> int:
    """在前 max_scan 行里找最像表头的一行（返回 0-based 下标）：
    字符串多、取值互不重复、非空多的行更像表头。"""
    best, best_score = 0, -1
    for i in range(min(len(rows), max_scan)):
        vals = [v for v in rows[i] if v not in (None, "")]
        if not vals:
            continue
        strings = sum(1 for v in vals if isinstance(v, str))
        distinct = len({str(v) for v in vals})
        score = strings * 2 + distinct + len(vals)
        if score > best_score:
            best, best_score = i, score
    return best


def build_headers(row: list) -> list[str]:
    """表头清洗：空表头补"列N"，重名加后缀。"""
    headers, seen = [], {}
    for i, v in enumerate(row):
        name = str(v).strip() if v not in (None, "") else f"列{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        headers.append(name)
    return headers


def parse_file(path: Path) -> dict:
    """解析整个文件，返回每个 sheet 的表头探测结果和样例。"""
    sheets = []
    for name, raw in read_all_sheets(path).items():
        rows = trim_rows(raw)
        if not rows:
            sheets.append({"name": name, "header_row": 1, "headers": [], "sample": [], "total_rows": 0})
            continue
        h_idx = detect_header_row(rows)
        headers = build_headers(rows[h_idx])
        data_rows = rows[h_idx + 1:]
        sample = [[cell_str(v) for v in r] for r in data_rows[:SAMPLE_SIZE]]
        sheets.append({
            "name": name,
            "header_row": h_idx + 1,   # 返回给前端的是 1-based 行号
            "headers": headers,
            "sample": sample,
            "total_rows": len(data_rows),
        })
    return {"sheets": sheets}
