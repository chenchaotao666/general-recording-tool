"""Luckysheet 编辑器 JSON → xlsx（openpyxl）。

浏览器内 Excel 模板编辑器的保存链路：前端 luckysheet.getAllSheets() 的结果
在这里还原成带样式的 xlsx，供 print_excel.fill_workbook 做占位符填充。

覆盖的样式：字体（名/号/粗/斜/下划线/删除线/颜色）、填充色、对齐/换行、
合并单元格、行高列宽、边框（13 种线型近似映射）。
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# luckysheet 字体下标 → 字体名（与 luckysheet 默认字体列表一致）
_FONT_LIST = [
    "Times New Roman", "Arial", "Tahoma", "Verdana", "Microsoft YaHei",
    "SimSun", "SimHei", "KaiTi", "FangSong", "NSimSun",
    "STXinwei", "STXihei", "STXingkai",
]

# luckysheet 边框线型 1~13 → openpyxl 线型
_BORDER_STYLES = {
    1: "thin", 2: "hair", 3: "dotted", 4: "dashed", 5: "dashDot",
    6: "dashDotDot", 7: "double", 8: "medium", 9: "mediumDashed",
    10: "mediumDashDot", 11: "slantDashDot", 12: "mediumDashDotDot", 13: "thick",
}

_HT = {0: "center", 1: "left", 2: "right"}     # 水平对齐
_VT = {0: "center", 1: "top", 2: "bottom"}     # 垂直对齐
_TR = {1: 45, 2: 135, 3: 255, 4: 90, 5: 180}   # 旋转（3=竖排文字）

# 空边（style=None）：序列化时该边不输出边框，用于「清除边框」的显式哨兵
_EMPTY_SIDE = Side()


def _argb(color) -> str | None:
    if not isinstance(color, str) or not color.startswith("#"):
        return None
    h = color.lstrip("#")
    if len(h) == 3:
        h = "".join(c * 2 for c in h)
    if len(h) not in (6, 8):
        return None
    return ("FF" + h if len(h) == 6 else h).upper()


def _font_name(ff):
    if ff is None:
        return None
    try:
        return _FONT_LIST[int(ff)]
    except (TypeError, ValueError, IndexError):
        return str(ff)


def _side(spec) -> Side | None:
    """borderInfo cell 值里的单边 {style, color} → openpyxl Side。
    style=0（无边框）→ 空 Side（显式清除）；注意数字 0 是 falsy，不能用 `or 1` 兜底。"""
    if not isinstance(spec, dict):
        return None
    raw = spec.get("style")
    try:
        s = 1 if raw in (None, "") else int(raw)
    except (TypeError, ValueError):
        s = 1
    if s == 0:
        return _EMPTY_SIDE
    return Side(style=_BORDER_STYLES.get(s, "thin"), color=_argb(spec.get("color")) or "FF000000")


def _expand_range(rng: dict) -> tuple[range, range]:
    rows, cols = rng.get("row") or [0, 0], rng.get("column") or [0, 0]
    return range(rows[0], rows[-1] + 1), range(cols[0], cols[-1] + 1)


def _apply_borders(ws, border_info: list) -> None:
    """把 luckysheet 的 borderInfo 展开为逐单元格边框（模板体量小，直接展开）。

    后写的覆盖先写的，与 luckysheet「后设置优先」一致。
    """
    def set_border(r, c, **sides):
        cur = ws.cell(row=r + 1, column=c + 1).border
        # None = 该边未提及（保留现状）；_EMPTY_SIDE = 显式清除（style=0/border-none）
        merged = {k: v for k, v in {
            "left": cur.left, "right": cur.right, "top": cur.top, "bottom": cur.bottom,
            **sides}.items() if v is not None}
        ws.cell(row=r + 1, column=c + 1).border = Border(**merged)

    for info in border_info or []:
        if not isinstance(info, dict):
            continue
        if info.get("rangeType") == "cell":
            v = info.get("value") or {}
            sides = {k: _side(v.get(k)) for k in ("l", "r", "t", "b")}
            # LuckyExcel 导出用 col_index（部分版本写 column_index）
            set_border(v.get("row_index", 0), v.get("col_index", v.get("column_index", 0)),
                       left=sides["l"], right=sides["r"], top=sides["t"], bottom=sides["b"])
            continue
        if info.get("rangeType") != "range":
            continue
        btype, style, color = info.get("borderType"), info.get("style"), info.get("color")
        side = _side({"style": style, "color": color}) or Side(style="thin", color="FF000000")
        for rng in info.get("range") or []:
            rows, cols = _expand_range(rng)
            for r in rows:
                for c in cols:
                    kw = {}
                    if btype == "border-none":   # 「无边框」：显式清除四边
                        kw = {"left": _EMPTY_SIDE, "right": _EMPTY_SIDE,
                              "top": _EMPTY_SIDE, "bottom": _EMPTY_SIDE}
                    elif btype in ("border-all", "border-outside"):
                        if r == rows.start: kw["top"] = side
                        if r == rows.stop - 1: kw["bottom"] = side
                        if c == cols.start: kw["left"] = side
                        if c == cols.stop - 1: kw["right"] = side
                    if btype != "border-none" and btype in ("border-all", "border-inside"):
                        if r > rows.start: kw["top"] = side
                        if r < rows.stop - 1: kw["bottom"] = side
                        if c > cols.start: kw["left"] = side
                        if c < cols.stop - 1: kw["right"] = side
                    if btype == "border-top" and r == rows.start: kw["top"] = side
                    if btype == "border-bottom" and r == rows.stop - 1: kw["bottom"] = side
                    if btype == "border-left" and c == cols.start: kw["left"] = side
                    if btype == "border-right" and c == cols.stop - 1: kw["right"] = side
                    if btype == "border-horizontal" and r < rows.stop - 1: kw["bottom"] = side
                    if btype == "border-vertical" and c < cols.stop - 1: kw["right"] = side
                    if kw:
                        set_border(r, c, **kw)


def _cell_value(cell: dict):
    v = cell.get("v")
    if v is None or v == "":
        m = cell.get("m")
        if m not in (None, ""):
            return m
        # 编辑器里直接输入的多行/富文本：cell.ct = {t:'inlineStr', s:[{v:...}, ...]}，
        # 正文在分段 runs 里（v/m 为空），不读会整格丢失
        ct = cell.get("ct") or {}
        if ct.get("t") == "inlineStr":
            parts = [(r.get("v") if isinstance(r, dict) else r) or "" for r in ct.get("s") or []]
            text = "".join(str(p) for p in parts)
            if text:
                return text
        return None
    ct = cell.get("ct") or {}
    if ct.get("t") == "n" or isinstance(v, (int, float)):
        try:
            return float(v)
        except (TypeError, ValueError):
            pass
    return str(v)


def _write_images(ws, sheet: dict) -> None:
    """luckysheet 工具栏插入的图片（sheet.images）→ xlsx 浮动图片。
    模型：{imgId: {src: dataURL, default: {left, top, width, height}, ...}}（位置为表内绝对像素）。
    锚点必须写 twoCellAnchor：luckyexcel 重新导入编辑器时只解析这种锚点，
    absoluteAnchor 的图会整个丢掉（保存后重开看不到图片，但预览/打印正常）。"""
    images = sheet.get("images")
    if not images:
        return
    import base64
    import io as _io
    import urllib.request

    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
    from openpyxl.utils.units import pixels_to_EMU

    config = sheet.get("config") or {}
    col_len = config.get("columnlen") or {}
    row_len = config.get("rowlen") or {}
    def_col = float(sheet.get("defaultColWidth") or 73)    # luckysheet 默认列宽 73px
    def_row = float(sheet.get("defaultRowHeight") or 19)   # luckysheet 默认行高 19px

    def _size(lens, idx, default):
        try:
            return float(lens[str(idx)])
        except (KeyError, TypeError, ValueError):
            return default

    def _locate(px, lens, default):
        """表内绝对像素 → (单元格索引, 格内偏移像素)，坐标口径与 luckysheet 画布一致。"""
        idx = 0
        px = max(0.0, float(px))
        while idx < 10000:
            s = _size(lens, idx, default)
            if px < s:
                break
            px -= s
            idx += 1
        return idx, int(round(px))

    for im in images.values():
        try:
            src = (im or {}).get("src") or ""
            if src.startswith("data:"):
                data = base64.b64decode(src.split(",", 1)[1])
            elif src.startswith(("http://", "https://")):
                data = urllib.request.urlopen(src, timeout=10).read()
            else:
                continue
            dft = im.get("default") or {}
            w = float(dft.get("width") or im.get("originWidth") or 100)
            h = float(dft.get("height") or im.get("originHeight") or 100)
            left = float(dft.get("left") or 0)
            top = float(dft.get("top") or 0)
            fc, off_x = _locate(left, col_len, def_col)
            fr, off_y = _locate(top, row_len, def_row)
            tc, off_x2 = _locate(left + w, col_len, def_col)
            tr, off_y2 = _locate(top + h, row_len, def_row)
            xi = XlImage(_io.BytesIO(data))
            xi.anchor = TwoCellAnchor(
                editAs="twoCell",
                _from=AnchorMarker(col=fc, colOff=pixels_to_EMU(off_x),
                                   row=fr, rowOff=pixels_to_EMU(off_y)),
                to=AnchorMarker(col=tc, colOff=pixels_to_EMU(off_x2),
                                row=tr, rowOff=pixels_to_EMU(off_y2)),
            )
            ws.add_image(xi)
        except Exception:
            continue   # 单张图片失败不影响整体转换


def _write_sheet(ws, sheet: dict) -> None:
    ws.title = (sheet.get("name") or "Sheet")[:31]
    config = sheet.get("config") or {}

    # data 缺失时用 celldata 稀疏还原
    data = sheet.get("data")
    if not data and sheet.get("celldata"):
        n_r, n_c = int(sheet.get("row") or 0), int(sheet.get("column") or 0)
        data = [[None] * n_c for _ in range(n_r)]
        for cd in sheet["celldata"]:
            if cd.get("r", 0) < n_r and cd.get("c", 0) < n_c:
                data[cd["r"]][cd["c"]] = cd.get("v")
    data = data or []

    for r, row in enumerate(data):
        if not row:
            continue
        for c, cell in enumerate(row):
            if not isinstance(cell, dict):
                continue
            xc = ws.cell(row=r + 1, column=c + 1)
            val = _cell_value(cell)
            if val is not None:
                xc.value = val
            fc, bg = _argb(cell.get("fc")), _argb(cell.get("bg"))
            xc.font = Font(
                name=_font_name(cell.get("ff")),
                size=cell.get("fs") if isinstance(cell.get("fs"), (int, float)) else None,
                bold=bool(cell.get("bl")), italic=bool(cell.get("it")),
                underline="single" if cell.get("un") else None,
                strike=bool(cell.get("cl")), color=fc,
            )
            if bg:
                xc.fill = PatternFill("solid", fgColor=bg)
            xc.alignment = Alignment(
                horizontal=_HT.get(cell.get("ht")), vertical=_VT.get(cell.get("vt")),
                wrap_text=cell.get("tb") == 2, text_rotation=_TR.get(cell.get("tr"), 0),
            )

    for key, m in (config.get("merge") or {}).items():
        if not isinstance(m, dict):
            continue
        r, c, rs, cs = (int(m.get(k) or 0) for k in ("r", "c", "rs", "cs"))
        if rs and cs and (rs > 1 or cs > 1):
            ws.merge_cells(start_row=r + 1, start_column=c + 1,
                           end_row=r + rs, end_column=c + cs)

    for idx, px in (config.get("columnlen") or {}).items():
        try:
            width = max(1.0, (float(px) - 5) / 7)   # px ≈ 字符宽×7+5
        except (TypeError, ValueError):
            continue
        ws.column_dimensions[get_column_letter(int(idx) + 1)].width = width
    for idx, px in (config.get("rowlen") or {}).items():
        try:
            ws.row_dimensions[int(idx) + 1].height = float(px) * 0.75   # px → pt
        except (TypeError, ValueError):
            continue

    _apply_borders(ws, config.get("borderInfo"))
    _write_images(ws, sheet)


def extract_images(ws) -> list:
    """xlsx 工作表浮动图片 → 编辑器注入模型（luckysheet 像素坐标）。
    换算口径与 _write_images / _write_sheet 完全互逆：显式列宽(字符)→px=width*7+5、
    显式行高(pt)→px=pt/0.75，未配置的用编辑器默认值 73/19，保证保存→重开位置不偏。"""
    import base64

    from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, OneCellAnchor, TwoCellAnchor
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import coordinate_to_tuple

    EMU = 9525

    def col_px(idx):
        dim = ws.column_dimensions.get(get_column_letter(idx + 1))
        return float(dim.width) * 7 + 5 if dim is not None and dim.width else 73.0

    def row_px(idx):
        dim = ws.row_dimensions.get(idx + 1)
        return float(dim.height) / 0.75 if dim is not None and dim.height else 19.0

    def col_pos(idx):
        return sum(col_px(i) for i in range(max(0, idx)))

    def row_pos(idx):
        return sum(row_px(i) for i in range(max(0, idx)))

    out = []
    for img in getattr(ws, "_images", []):
        try:
            a = img.anchor
            if isinstance(a, TwoCellAnchor):
                left = col_pos(a._from.col) + a._from.colOff / EMU
                top = row_pos(a._from.row) + a._from.rowOff / EMU
                width = col_pos(a.to.col) + a.to.colOff / EMU - left
                height = row_pos(a.to.row) + a.to.rowOff / EMU - top
            elif isinstance(a, OneCellAnchor):
                left = col_pos(a._from.col) + a._from.colOff / EMU
                top = row_pos(a._from.row) + a._from.rowOff / EMU
                width, height = a.ext.cx / EMU, a.ext.cy / EMU
            elif isinstance(a, AbsoluteAnchor):
                left, top = a.pos.x / EMU, a.pos.y / EMU
                width, height = a.ext.cx / EMU, a.ext.cy / EMU
            elif isinstance(a, str):   # "A1" 单元格形式
                r, c = coordinate_to_tuple(a)
                left, top = col_pos(c - 1), row_pos(r - 1)
                width, height = float(img.width or 100), float(img.height or 100)
            else:
                continue
            data = img._data()
            mime = "image/png"
            if data[:3] == b"\xff\xd8\xff":
                mime = "image/jpeg"
            elif data[:6] in (b"GIF87a", b"GIF89a"):
                mime = "image/gif"
            out.append({
                "src": f"data:{mime};base64," + base64.b64encode(data).decode(),
                "default": {"left": round(left, 1), "top": round(top, 1),
                            "width": round(max(1.0, width), 1), "height": round(max(1.0, height), 1)},
            })
        except Exception:
            continue   # 单张图片失败不影响其他图片
    return out


def sheets_to_xlsx(sheets: list) -> bytes:
    if not sheets:
        raise ValueError("编辑器内容为空")
    wb = Workbook()
    wb.remove(wb.active)
    for i, sheet in enumerate(sheets):
        _write_sheet(wb.create_sheet(f"Sheet{i + 1}"), sheet or {})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
