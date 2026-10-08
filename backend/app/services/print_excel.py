"""Excel 打印模板：起始模板生成 + 占位符填充。

占位符语法（写在单元格文本里）：
  {字段名}              主表字段值（单元格内容恰好是单个占位符且值为数值时写为数字，便于 Excel 再计算）
  {_id}                 记录 ID
  {_today}              打印当天（YYYY-MM-DD）
  {_total_cn}           明细金额的人民币大写（金额列求和；无金额列时 数量×单价 兜底）
  {子表.列名}           明细循环：含此类占位符的行按明细行数向下复制展开（样式/合并单元格随行复制）
  {子表._index}         明细行序号（1 起）
  {子表._count}         明细行数
  {子表.列名#sum}       明细某列合计（写在非循环行，如合计行）
  {_row.字段名}         整表明细模式：模板含此类占位符时，全部筛选记录作为同一张单据的
                        明细行渲染（主表占位符取第一条记录）；{_row._index}/{_row.列名#sum} 同上
"""
import io
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from ..config import settings
from ..models import MetaField, MetaTable
from .print_presets import _AMOUNT_HINTS, _PRICE_HINTS, _match, _pick_date_field, _pick_no_field

TOKEN_RE = re.compile(r"\{([^{}]+)\}")
MAX_FILL_RECORDS = 200   # 单次填充记录数上限（批量打印与列表口径一致）


def tpl_dir() -> Path:
    d = settings.upload_dir / "print_templates"
    d.mkdir(parents=True, exist_ok=True)
    return d


def tpl_path(tpl_id: int) -> Path:
    return tpl_dir() / f"{tpl_id}.xlsx"


# ---------- 值格式化（与前端 printDoc.formatCell 同规则） ----------

def _fmt(field: MetaField | dict | None, value) -> str:
    if value is None:
        return ""
    t = (field or {}).get("data_type") if isinstance(field, dict) else getattr(field, "data_type", None)
    if t == "bool":
        return "是" if value else "否"
    if t == "date":
        return str(value)[:10]
    if t == "datetime":
        return str(value).replace("T", " ")[:19]
    if t in ("int", "decimal") and value != "":
        try:
            return str(round(float(value), 4))
        except (TypeError, ValueError):
            pass
    if isinstance(value, (list, dict)):
        return ""
    return str(value)


# ---------- 人民币大写 ----------

_CN_DIGIT = "零壹贰叁肆伍陆柒捌玖"


def to_chinese_upper(n: float) -> str:
    if n is None or n < 0:
        return ""
    total = round(n * 100)   # 总分，避免浮点误差
    integer, cents = divmod(total, 100)
    jiao, fen = divmod(cents, 10)
    s = ""
    if integer:
        groups = []
        while integer:
            groups.append(integer % 10000)
            integer //= 10000
        big_units = ["", "万", "亿", "万亿"]
        parts = []
        for i in range(len(groups) - 1, -1, -1):
            g = groups[i]
            if g == 0:
                parts.append("")
                continue
            units = ["", "拾", "佰", "仟"]
            ds = [g // 1000 % 10, g // 100 % 10, g // 10 % 10, g % 10]
            gs = ""
            zero = False
            for k in range(4):
                if ds[k] == 0:
                    if gs:
                        zero = True
                    continue
                if zero:
                    gs += "零"
                    zero = False
                gs += _CN_DIGIT[ds[k]] + units[3 - k]
            parts.append(gs + big_units[i])
        for i, p in enumerate(parts):
            if not p:
                continue
            if s and groups[len(parts) - 1 - i] < 1000:
                s += "零"
            s += p
        s += "元"
    if not jiao and not fen:
        return (s or "零元") + "整"
    if s and not jiao and fen:
        s += "零"
    return s + (_CN_DIGIT[jiao] + "角" if jiao else "") + (_CN_DIGIT[fen] + "分" if fen else "")


# ---------- 占位符解析 ----------

def _tokens(cell_value) -> list[str]:
    return TOKEN_RE.findall(cell_value) if isinstance(cell_value, str) else []


def _split_token(token: str):
    """'子表.列#sum' → ('子表', '列', 'sum')；'字段' → (None, '字段', None)；'_id' → (None, '_id', None)"""
    field, _, agg = token.partition("#")
    obj, dot, col = field.partition(".")
    if dot:
        return obj, col, agg or None
    return None, field, agg or None


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _typed(field, value):
    """占位符取值：数值字段保留数字类型（整单元格占位符写入 Excel 才是可计算的数字；
    内联到文本时 str(int/float) 同样干净），其余走 _fmt 文本。"""
    t = (field.get("data_type") if isinstance(field, dict) else getattr(field, "data_type", None)) if field is not None else None
    if t in ("int", "decimal") and value is not None and value != "":
        try:
            f = round(float(value), 4)
            return int(f) if float(f).is_integer() else f
        except (TypeError, ValueError):
            pass
    return _fmt(field, value)


def _resolve(token: str, rec: dict, fields_by_name: dict, subform_defs: dict, detail_ctx=None):
    """解析单个占位符。detail_ctx=(子表名, 行号) 时解析循环行内的 子表.列 引用。
    返回 (found: bool, value)。"""
    obj, col, agg = _split_token(token)
    if obj is None:
        if col == "_id":
            return True, rec.get("id")
        if col == "_today":
            return True, date.today().isoformat()
        if col == "_total_cn":
            return True, to_chinese_upper(_total_amount(rec, subform_defs))
        f = fields_by_name.get(col)
        if f is not None:
            return True, _typed(f, rec.get(col))
        return False, None
    # 子表引用
    if obj not in subform_defs:
        return False, None
    cols = subform_defs[obj]["columns"]
    rows = rec.get(obj) if isinstance(rec.get(obj), list) else []
    if col == "_count":
        return True, len(rows)
    if col == "_index":
        return True, (detail_ctx[1] + 1) if detail_ctx else ""
    if agg == "sum":
        return True, round(sum(_num(r.get(col)) for r in rows), 4)
    if detail_ctx and detail_ctx[0] == obj and detail_ctx[1] < len(rows):
        return True, _typed(cols.get(col), rows[detail_ctx[1]].get(col))
    return False, None


def _total_amount(rec: dict, subform_defs: dict) -> float:
    """大写金额口径：优先「金额/amount」列求和；否则 数量列×单价列 逐行求积。"""
    for sub, d in subform_defs.items():
        rows = rec.get(sub) if isinstance(rec.get(sub), list) else []
        cols = d["columns"]
        amount = next((n for n, c in cols.items() if _match(_AMOUNT_HINTS, n, c.get("label", ""))), None)
        if amount:
            return round(sum(_num(r.get(amount)) for r in rows), 2)
        price = next((n for n, c in cols.items() if _match(_PRICE_HINTS, n, c.get("label", ""))), None)
        qty = next((n for n, c in cols.items()
                    if c.get("data_type") in ("int", "decimal") and n != price and not _match(_PRICE_HINTS, n, c.get("label", ""))), None)
        if price and qty:
            return round(sum(_num(r.get(qty)) * _num(r.get(price)) for r in rows), 2)
    return 0


def _fill_cell_text(text: str, rec, fields_by_name, subform_defs, detail_ctx=None):
    """替换单元格文本中的全部占位符。内容恰好是单个占位符时尽量保留值类型（数字写数字）。"""
    tokens = _tokens(text)
    if not tokens:
        return text
    whole = tokens[0] if len(tokens) == 1 and text.strip() == "{" + tokens[0] + "}" else None
    out = text
    for token in tokens:
        found, value = _resolve(token, rec, fields_by_name, subform_defs, detail_ctx)
        if not found:
            continue   # 未识别占位符原样保留（用户可见，便于排查）
        if whole:
            return value
        out = out.replace("{" + token + "}", "" if value is None else str(value))
    return out


# ---------- 明细循环行展开 ----------

def _loop_row_of(ws, subform_names: set[str]) -> dict[int, str]:
    """找出含 子表.列 循环占位符的行 → {行号: 子表名}（#sum/_count 不算循环标记）。"""
    out = {}
    for row in ws.iter_rows():
        for cell in row:
            for token in _tokens(cell.value):
                obj, col, agg = _split_token(token)
                if obj in subform_names and col not in ("_count",) and agg != "sum" and col != "_index":
                    out[cell.row] = obj
    return out


def _expand_loop_row(ws, row_idx: int, count: int) -> None:
    """把模板循环行复制成 count 行：先手动平移下方合并区域，再插行补样式，最后复制单行合并。"""
    if count <= 1:
        return
    shift = count - 1
    single_row_merges = []
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == mr.max_row == row_idx:
            single_row_merges.append(mr)
        elif mr.min_row > row_idx:
            mr.shift(0, shift)
        elif mr.min_row <= row_idx < mr.max_row:
            mr.max_row += shift
    ws.insert_rows(row_idx + 1, shift)
    height = ws.row_dimensions[row_idx].height
    for i in range(1, shift + 1):
        target = row_idx + i
        for cell in ws[row_idx]:
            nc = ws.cell(row=target, column=cell.column, value=cell.value)
            if cell.has_style:
                nc._style = copy(cell._style)
        if height:
            ws.row_dimensions[target].height = height
        for mr in single_row_merges:
            ws.merge_cells(start_row=target, end_row=target,
                           start_column=mr.min_col, end_column=mr.max_col)


# ---------- 主流程 ----------

def _fill_cells(ws, rec, regions, fields_by_name, subform_defs) -> None:
    """逐单元格填充：循环区（regions=[(起点, 终点, 子表名)]）内的带 detail_ctx，其余全局替换。"""
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str) or "{" not in cell.value:
                continue
            detail_ctx = None
            for lr, end, name in regions:
                if lr <= cell.row < end:
                    detail_ctx = (name, cell.row - lr)
                    break
            cell.value = _fill_cell_text(cell.value, rec, fields_by_name, subform_defs, detail_ctx)


def _has_flat_tokens(ws) -> bool:
    """模板含 {_row.字段} 占位符 → 整表明细模式（全部筛选记录渲染为同一张单据的明细行）。"""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "{_row." in cell.value:
                return True
    return False


def _fill_flat_sheet(ws, fields, fields_by_name, records) -> tuple | None:
    """整表明细模式：{_row.列} 循环区按记录数展开；主表占位符取第一条记录；#sum/_total_cn 按全记录计。
    返回循环区 (起点, 终点)（行号 1 起，终点不含），供分页切片用。"""
    cols = {f.field_name: {"field_name": f.field_name, "label": f.label, "data_type": f.data_type}
            for f in fields if f.data_type not in ("subform", "image")}
    subform_defs = {"_row": {"label": "明细", "columns": cols}}
    rec = dict(records[0]) if records else {}
    rec["_row"] = records   # 借用子表解析机制：_row 即整条记录列表
    loops = _loop_row_of(ws, {"_row"})
    n = max(len(records), 1)
    for row_idx in sorted(loops, reverse=True):
        _expand_loop_row(ws, row_idx, n)
    regions = sorted((lr, lr + n, "_row") for lr in loops)
    _fill_cells(ws, rec, regions, fields_by_name, subform_defs)
    no_field = _pick_no_field(fields)
    base = str(rec.get(no_field) or rec.get("id") or "") if no_field else str(rec.get("id") or "")
    ws.title = re.sub(r"[\\/*?\[\]:]", "", base)[:28] or "打印"
    if regions:
        return (min(r[0] for r in regions), max(r[1] for r in regions))
    return None


def xlsx_media_bytes(data: bytes) -> set:
    """xlsx 字节流里 xl/media/* 的内容集合。图片去重用它判断，
    不碰 openpyxl 图片对象的流（从 xlsx 加载的图片流读一次就关，再读/再 save 会崩）。"""
    import zipfile

    out = set()
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for n in z.namelist():
            if n.startswith("xl/media/"):
                out.add(z.read(n))
    return out


def _copy_images(src_ws, dst_ws, skip: set | None = None) -> None:
    """copy_worksheet 不复制图片（logo 等），手动带过去。
    同一源图可能被复制到多个工作表，图片数据按源对象缓存（openpyxl 图片流只能读一次；
    用 WeakKeyDictionary 避免 id() 复用后命中脏缓存）。
    skip: 目标里已有图片的内容集合（编辑器 JSON 已带回的），命中则跳过，避免一份图叠出两张。"""
    if not hasattr(_copy_images, "_cache"):
        import weakref

        _copy_images._cache = weakref.WeakKeyDictionary()
    cache = _copy_images._cache
    for img in getattr(src_ws, "_images", []):
        try:
            from openpyxl.drawing.image import Image as XlImage

            data = cache.get(img)
            if data is None:
                data = cache[img] = img._data()
            if skip and data in skip:
                continue   # 编辑器已带回同一张图
            new = XlImage(io.BytesIO(data))
            new.width, new.height = img.width, img.height
            new.anchor = copy(img.anchor)
            dst_ws.add_image(new)
        except Exception:
            continue   # 单张图片失败不影响整体填充


def fill_workbook_meta(tpl_file: Path, mt: MetaTable, fields: list[MetaField], records: list[dict]):
    """fill_workbook 的元数据版：返回 (Workbook, [每表明细区 (start, end) 或 None])。

    明细区 = 循环占位符展开后的行范围（1 起，end 不含），供 fill-view 按纸张高度切片。
    """
    fields_by_name = {f.field_name: f for f in fields}
    subform_defs = {
        f.field_name: {"label": f.label,
                       "columns": {c.get("field_name"): c for c in (f.options or {}).get("columns") or []}}
        for f in fields if f.data_type == "subform"
    }
    wb = load_workbook(tpl_file)
    tmpl = wb.active
    if _has_flat_tokens(tmpl):
        region = _fill_flat_sheet(tmpl, fields, fields_by_name, records[:MAX_FILL_RECORDS])
        return wb, [region]
    metas = []
    for idx, rec in enumerate(records[:MAX_FILL_RECORDS]):
        ws = wb.copy_worksheet(tmpl)
        _copy_images(tmpl, ws)
        # 1) 明细循环行展开（自下而上，行号才不受插行影响）
        loops = _loop_row_of(ws, set(subform_defs))
        for row_idx in sorted(loops, reverse=True):
            sub = loops[row_idx]
            n = len(rec.get(sub)) if isinstance(rec.get(sub), list) else 0
            _expand_loop_row(ws, row_idx, max(n, 1))
        # 展开后的循环区：[起点, 起点+行数)，按起点升序定位每个单元格所属区域
        regions = sorted(
            (lr, lr + max(len(rec.get(name)) if isinstance(rec.get(name), list) else 0, 1), name)
            for lr, name in loops.items()
        )
        # 2) 逐单元格填充：循环区内的带 detail_ctx，其余全局替换
        _fill_cells(ws, rec, regions, fields_by_name, subform_defs)
        metas.append((min(r[0] for r in regions), max(r[1] for r in regions)) if regions else None)
        # 工作表命名：单号字段值 → id；去非法字符，限 31 字符，重名加序号
        no_field = _pick_no_field(fields)
        base = str(rec.get(no_field) or rec.get("id") or idx + 1) if no_field else str(rec.get("id") or idx + 1)
        base = re.sub(r"[\\/*?\[\]:]", "", base)[:28] or f"记录{idx + 1}"
        name, n = base, 2
        while name in wb.sheetnames:
            name = f"{base}-{n}"
            n += 1
        ws.title = name
    wb.remove(tmpl)
    return wb, metas


def fill_workbook(tpl_file: Path, mt: MetaTable, fields: list[MetaField], records: list[dict]) -> bytes:
    """模板 xlsx + 记录列表 → 填充后的 xlsx（每条记录一个工作表，按 单号/id 命名）。"""
    wb, _ = fill_workbook_meta(tpl_file, mt, fields, records)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------- 起始模板 ----------

def starter_workbook(mt: MetaTable, fields: list[MetaField]) -> bytes:
    """按表结构生成可编辑的起始 xlsx 模板（占位符已布好，用户在 Excel/WPS 里继续排版）。"""
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    sub = next((f for f in fields if f.data_type == "subform"), None)
    sub_cols = (sub.options or {}).get("columns") or [] if sub else []
    # 无子表（平表）：取 varchar/int/decimal 字段当整表明细列，{_row.字段} 循环。
    # 按单据明细常用角色挑选（每角色限量），避免同类字段（如多个"数量"）挤掉金额等关键列
    _FLAT_COL_RULES = [
        ("名称", ("名称", "name"), 2),
        ("规格", ("规格", "spec", "型号"), 1),
        ("单位", ("单位", "unit"), 1),
        ("数量", ("数量", "qty"), 1),
        ("单价", ("单价", "price"), 1),
        ("金额", ("金额", "amount", "合计", "小计"), 1),
        ("备注", ("备注", "remark"), 1),
    ]

    def _flat_pick() -> list[dict]:
        usable = [f for f in fields if f.data_type in ("varchar", "int", "decimal")]
        picked, used = [], set()
        for _, keywords, cap in _FLAT_COL_RULES:
            for f in usable:
                if len(picked) >= 7:
                    break
                if f.field_name in used:
                    continue
                blob = f"{f.field_name} {f.label}".lower()
                if any(k in blob for k in keywords) and \
                        sum(1 for p in picked if p[0] is keywords) < cap:
                    picked.append((keywords, f))
                    used.add(f.field_name)
        for f in usable:   # 名额未满则用其余字段补齐
            if len(picked) >= 7:
                break
            if f.field_name not in used:
                picked.append((None, f))
                used.add(f.field_name)
        return [{"field_name": f.field_name, "label": f.label, "data_type": f.data_type}
                for _, f in picked]

    flat_cols = [] if sub else _flat_pick()
    detail_cols = sub_cols or flat_cols
    detail_ref = sub.field_name if sub else "_row"
    no_field = _pick_no_field(fields)
    date_field = _pick_date_field(fields)
    info = [f for f in fields if f.data_type in ("varchar", "text", "serial")][:4]
    title = mt.label if "单" in (mt.label or "") else ("送货单" if sub else "单据")

    ncols = max(1 + len(detail_cols), 4)   # 序号 + 明细列；至少 4 列撑版式
    last = get_column_letter(ncols)

    wb = Workbook()
    ws = wb.active
    ws.title = "打印模板"
    r = 1
    ws.merge_cells(f"A{r}:{last}{r}")
    c = ws.cell(r, 1, title)
    c.font = Font(size=16, bold=True)
    c.alignment = center
    ws.row_dimensions[r].height = 30
    r += 1
    ws.cell(r, 1, f"NO.：{{{no_field or '_id'}}}")
    ws.cell(r, ncols, f"日期：{{{date_field or '_today'}}}").alignment = Alignment(horizontal="right")
    r += 1
    for f in info:
        ws.cell(r, 1, f"{f.label}：{{{f.field_name}}}")
        r += 1
    r += 1   # 空一行
    if detail_cols:
        head_row = r
        ws.cell(r, 1, "序号")
        for i, col in enumerate(detail_cols):
            ws.cell(r, 2 + i, col.get("label") or col.get("field_name"))
        r += 1
        loop_row = r
        ws.cell(r, 1, f"{{{detail_ref}._index}}")
        for i, col in enumerate(detail_cols):
            ws.cell(r, 2 + i, f"{{{detail_ref}.{col.get('field_name')}}}")
        r += 1
        sum_cols = [c for c in detail_cols if c.get("data_type") in ("int", "decimal")]
        if sum_cols:
            ws.cell(r, 1, "合计")
            for i, col in enumerate(detail_cols):
                if col in sum_cols:
                    ws.cell(r, 2 + i, f"{{{detail_ref}.{col.get('field_name')}#sum}}")
            r += 1
        ws.cell(r, 1, f"合计金额（大写）：{{_total_cn}}")
        for rr in range(head_row, r):
            for cc in range(1, ncols + 1):
                cell = ws.cell(rr, cc)
                cell.border = border
                if rr == head_row:
                    cell.font = Font(bold=True)
                    cell.alignment = center
                elif cc > 1:
                    cell.alignment = center
        r += 1
    ws.cell(r, 1, "备注：")
    r += 2
    ws.cell(r, 1, "收货单位及经手人（签章）：")
    ws.cell(r, max(ncols - 2, 2), "送货单位及经手人（签章）：")
    for i in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.column_dimensions["B"].width = 24
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
