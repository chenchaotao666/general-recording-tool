"""填充结果 → HTML 打印页：每条记录一个 .page（A4），浏览器直接 window.print()。

在 fill_workbook 产出的 xlsx 上按单元格还原版式：列宽行高、合并单元格、
字体/字号/粗斜/下划线/删除线/颜色、填充色、对齐/换行、边框。
"""
import html as _html
import io
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..models import MetaField, MetaTable
from .print_excel import fill_workbook

_BORDER_CSS = {
    "hair": "1px solid", "thin": "1px solid", "medium": "2px solid", "thick": "3px solid",
    "double": "3px double", "dashed": "1px dashed", "dotted": "1px dotted",
    "dashDot": "1px dashed", "dashDotDot": "1px dotted",
    "mediumDashed": "2px dashed", "mediumDashDot": "2px dashed", "mediumDashDotDot": "2px dotted",
    "slantDashDot": "2px dashed",
}
_VALIGN = {"center": "middle", "top": "top", "bottom": "bottom", "justify": "middle", "distributed": "middle"}


def _rgb(color) -> str | None:
    """openpyxl Color → #rrggbb（仅 rgb 型；主题色/索引色放弃，用默认黑）。"""
    if color is None or getattr(color, "type", None) != "rgb" or not isinstance(color.rgb, str):
        return None
    h = color.rgb.lstrip("#")
    if len(h) == 8:
        h = h[2:]   # ARGB → RGB
    return f"#{h}" if len(h) == 6 else None


def _font_css(cell) -> str:
    f = cell.font
    css = []
    if f.name:
        css.append(f"font-family:'{f.name}'")
    if f.size:
        css.append(f"font-size:{f.size}pt")
    if f.bold:
        css.append("font-weight:bold")
    if f.italic:
        css.append("font-style:italic")
    deco = []
    if f.underline:
        deco.append("underline")
    if f.strike:
        deco.append("line-through")
    if deco:
        css.append(f"text-decoration:{' '.join(deco)}")
    color = _rgb(f.color)
    if color:
        css.append(f"color:{color}")
    return ";".join(css)


def _cell_css(cell) -> str:
    css = []
    if cell.fill and cell.fill.patternType == "solid":
        bg = _rgb(cell.fill.fgColor)
        if bg and bg.lower() != "#ffffff":
            css.append(f"background:{bg}")
    al = cell.alignment
    if al.horizontal in ("center", "right", "left", "justify"):
        css.append(f"text-align:{al.horizontal}")
    if al.vertical in _VALIGN:
        css.append(f"vertical-align:{_VALIGN[al.vertical]}")
    css.append(f"white-space:{'pre-wrap' if al.wrap_text else 'nowrap'}")
    font = _font_css(cell)
    if font:
        css.append(font)
    b = cell.border
    for side_name, css_prop in (("left", "border-left"), ("right", "border-right"),
                                ("top", "border-top"), ("bottom", "border-bottom")):
        side = getattr(b, side_name, None)
        if side and side.style:
            line = _BORDER_CSS.get(side.style, "1px solid")
            color = _rgb(side.color) or "#000000"
            css.append(f"{css_prop}:{line} {color}")
    return ";".join(css)


def _vt_stack(text: str) -> str:
    """255 竖排文字 → 单列逐字堆叠 HTML。源文本换行 → 列内额外的空行。
    不用 CSS writing-mode：vertical-rl 会把换行变成「向左的新列」，
    与 WPS/luckysheet 编辑器里 255 竖排的显示（同列留空）不一致，且窄单元格
    overflow:hidden 下多行内容会被裁得只剩中间一行。"""
    lines = ["<br>".join(_html.escape(ch) for ch in ln) for ln in text.split("\n")]
    return "<br>".join(lines)


def _images_html(ws, row_of) -> str:
    """工作表图片（logo 等）→ 绝对定位的 <img>（base64 内联）。
    row_of: 行号(1 起) → 该行的渲染 y 偏移(px)，补空行/切片时由调用方传入。"""
    import base64

    from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, OneCellAnchor, TwoCellAnchor

    EMU = 9525   # EMU per px (96dpi)

    def col_pos(idx: int) -> int:
        """列 idx(0 起) 左边缘的 x(px)；未配置列按 Excel 默认 64px。"""
        x = 0
        for c in range(1, idx + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            x += round(dim.width * 7 + 5) if dim and dim.width else 64
        return x

    out = []
    for img in getattr(ws, "_images", []):
        try:
            a = img.anchor
            if isinstance(a, AbsoluteAnchor):
                left, top = round(a.pos.x / EMU), round(a.pos.y / EMU)
                w, h = round(a.ext.cx / EMU), round(a.ext.cy / EMU)
            elif isinstance(a, (TwoCellAnchor, OneCellAnchor)):
                fr = a._from
                left = col_pos(fr.col) + round((fr.colOff or 0) / EMU)
                top = round(row_of(fr.row) + (fr.rowOff or 0) / EMU)
                if isinstance(a, TwoCellAnchor):
                    # 尺寸 = from→to 跨度（编辑器里的显示尺寸；这种锚点没有 ext，拿原图
                    # 像素尺寸当显示尺寸是错的——预览里图片大小不对的根因）
                    w = col_pos(a.to.col) + round((a.to.colOff or 0) / EMU) - left
                    h = round(row_of(a.to.row) + (a.to.rowOff or 0) / EMU) - top
                else:
                    w, h = round(a.ext.cx / EMU), round(a.ext.cy / EMU)
            else:
                continue
            data = img._data()
            fmt = img.format or ("png" if data[:4] == b"\x89PNG" else "jpeg")
            src = f"data:image/{fmt};base64,{base64.b64encode(data).decode()}"
            size = f"width:{w}px;height:{h}px;" if w and h else ""
            out.append(f'<img src="{src}" style="position:absolute;left:{left}px;top:{top}px;{size}z-index:5" />')
        except Exception:
            continue   # 单张图片失败不影响整体渲染
    return "".join(out)


def _sheet_html(ws, only_rows: list[int] | None = None,
                pad_before: int | None = None, pad_count: int = 0, pad_h: float | None = 16.0) -> str:
    """单个工作表 → <table>（used range 内；合并单元格转 colspan/rowspan）。
    only_rows 给定行号（1 起）时只渲染这些行（分页切片用）：跨块合并按渲染范围截断。
    pad_before/pad_count/pad_h：在该行前插入 pad_count 个带网格边框的补空行（填满格子高度用）。"""
    max_c = ws.max_column or 1
    all_rows = sorted(only_rows) if only_rows else list(range(1, (ws.max_row or 1) + 1))
    rendered = set(all_rows)

    spans = {}    # (r, c) → (rowspan, colspan)（仅左上角）
    covered = set()
    for mr in ws.merged_cells.ranges:
        if only_rows is None:
            rs, cs = mr.max_row - mr.min_row + 1, mr.max_col - mr.min_col + 1
            span_rows = list(range(mr.min_row, mr.max_row + 1))
        else:
            if mr.min_row not in rendered:
                continue   # 锚点行不在本页 → 该合并块整体不渲染（内容随锚点所在页显示）
            # 截断：合并范围内被渲染的行数（输出行相邻，跨切片的空洞在渲染序列里不存在）
            span_rows = [r for r in all_rows if mr.min_row <= r <= mr.max_row]
            rs, cs = len(span_rows), mr.max_col - mr.min_col + 1
        spans[(mr.min_row, mr.min_col)] = (rs, cs)
        # 覆盖单元格 = 同行横向 + 被 rowspan 覆盖的输出行（截断后）
        for c in range(mr.min_col + 1, mr.max_col + 1):
            covered.add((mr.min_row, c))
        for r in span_rows[1:rs]:
            for c in range(mr.min_col, mr.max_col + 1):
                covered.add((r, c))

    cols = []
    total_px = 0
    for c in range(1, max_c + 1):
        dim = ws.column_dimensions.get(get_column_letter(c))
        px = round(dim.width * 7 + 5) if dim and dim.width else 64   # 字符宽 ≈ 7px+5
        cols.append(f'<col style="width:{px}px">')
        total_px += px

    # 补空行：沿用明细末行的边框样式
    filler_html = ""
    if pad_count and pad_before:
        src = pad_before - 1
        tds = []
        for c in range(1, max_c + 1):
            if (src, c) in covered:
                continue
            cell = ws.cell(row=src, column=c)
            b = cell.border
            css = []
            for side_name, css_prop in (("left", "border-left"), ("right", "border-right"),
                                        ("top", "border-top"), ("bottom", "border-bottom")):
                side = getattr(b, side_name, None)
                if side and side.style:
                    line = _BORDER_CSS.get(side.style, "1px solid")
                    color = _rgb(side.color) or "#000000"
                    css.append(f"{css_prop}:{line} {color}")
            attrs = f' style="{";".join(css)}"' if css else ""
            tds.append(f"<td{attrs}>&nbsp;</td>")
        tr_style = f' style="height:{pad_h}pt"' if pad_h else ""
        filler_html = f"<tr{tr_style}>" + "".join(tds) + "</tr>"

    rows_html = []
    for r in all_rows:
        if pad_before and r == pad_before and filler_html:
            rows_html.append(filler_html * pad_count)
        dim = ws.row_dimensions.get(r)
        tr_style = f' style="height:{dim.height}pt"' if dim and dim.height else ""
        tds = []
        for c in range(1, max_c + 1):
            if (r, c) in covered:
                continue
            cell = ws.cell(row=r, column=c)
            span = spans.get((r, c))
            attrs = ""
            if span:
                rs, cs = span
                if rs > 1:
                    attrs += f' rowspan="{rs}"'
                if cs > 1:
                    attrs += f' colspan="{cs}"'
            style = _cell_css(cell)
            raw = cell.value
            value = "&nbsp;" if raw in (None, "") else _html.escape(str(raw))
            # 竖排文字（Excel text_rotation=255，如右侧联注）：绝对定位脱离布局流，
            # 否则窄列里的长竖排文本按横排换行计算会把整行撑高变形。
            # position:relative 必须和样式合并进同一个 style 属性——拆成两个 style 属性时
            # 浏览器只认第一个，relative 丢失后 .vt 会相对页面容器定位，竖排文字跑版
            if value != "&nbsp;" and cell.alignment and cell.alignment.text_rotation == 255:
                style = f"position:relative;{style}" if style else "position:relative"
                value = f'<span class="vt"><span>{_vt_stack(str(raw))}</span></span>'
            if style:
                attrs += f' style="{style}"'
            tds.append(f"<td{attrs}>{value}</td>")
        rows_html.append(f"<tr{tr_style}>{''.join(tds)}</tr>")

    table = (f'<table style="border-collapse:collapse;table-layout:fixed;width:{total_px}px">'
             + "".join(cols) + "".join(rows_html) + "</table>")

    # 图片（logo 等）：按锚点绝对定位，盖在表格上方
    def _row_y(row0: int) -> float:
        # 锚点行（0 起）之前的渲染行高合计（px）
        y = 0.0
        for r in all_rows:
            if r > row0:
                break
            y += _tr_px(ws, r)
        if pad_before and pad_count and row0 + 1 >= pad_before:
            y += pad_count * ((pad_h or 16.0) * 96 / 72)
        return y

    images = _images_html(ws, _row_y)
    if images:
        return f'<div style="position:relative">{table}{images}</div>'
    return table


def _tr_px(ws, r: int) -> float:
    """渲染行高（px）：显式行高 pt→px，否则 21px（约 11pt 文字行）。"""
    dim = ws.row_dimensions.get(r)
    return float(dim.height) * 96 / 72 if dim and dim.height else 21.0

def _row_pt(ws, r: int) -> float:
    """行高估算（pt）：显式行高优先，否则按 16pt（约 21px）。"""
    dim = ws.row_dimensions.get(r)
    return float(dim.height) if dim and dim.height else 16.0


def _pad_plan(ws, region, slot_pt: float):
    """计算补空行方案：让单据内容恰好填满格子高度（2等分/3等分）。
    返回 (在该行之前插入, 行数, 行高pt)；内容已超出格子时不补。返回 (None, 0, 0) 表示不补。"""
    max_r = ws.max_row or 1
    detail_h = sum(_row_pt(ws, r) for r in range(region[0], region[1]))
    head_h = sum(_row_pt(ws, r) for r in range(1, region[0]))
    tail_h = sum(_row_pt(ws, r) for r in range(region[1], max_r + 1))
    content_h = head_h + detail_h + tail_h
    if content_h >= slot_pt:
        return (None, 0, 0)
    dim = ws.row_dimensions.get(region[1] - 1)
    filler_h = float(dim.height) if dim and dim.height else None   # 无显式行高 → 补空行也不设高（自动同内容行）
    count = min(int((slot_pt - content_h) / (filler_h or 16.0)), 60)
    if count <= 0:
        return (None, 0, 0)
    return (region[1], count, filler_h)   # 插入位置：明细区末尾（合计/签名之前）


# 纸张模式（参照一彩打印软件）：A4 = 每条记录一张整页；
# 2等分/3等分 = 单据明细区末尾补空行，使每份单据恰好占满 1/2、1/3 页高
_PAGE_CSS = {
    "a4": ".page{min-height:269mm;page-break-after:always}.page:last-child{page-break-after:auto}",
    "half": (".page{min-height:133mm;page-break-after:auto}"
             ".page:nth-child(2n){page-break-after:always}"
             ".page:nth-child(odd){border-bottom:1px dashed #aaa}"),
    "third": (".page{min-height:89mm;page-break-after:auto}"
              ".page:nth-child(3n){page-break-after:always}"
              ".page:not(:nth-child(3n)){border-bottom:1px dashed #aaa}"),
}

# A4 打印区高约 277mm ≈ 785pt；半页/三分之一页目标高度（pt）
_SLOT_PT = {"half": 392, "third": 261}


def render_fill_html(tpl_file: Path, mt: MetaTable, fields: list[MetaField],
                     records: list[dict], title: str = "打印", paper: str = "a4") -> str:
    """模板 xlsx + 记录列表 → 完整 HTML 打印页。

    a4：每条记录一整页。half/third：每条记录明细区末尾补空行，
    使单据恰好占满 1/2、1/3 页高（一张 A4 排 2/3 份，虚线为裁切参考）。
    """
    from .print_excel import fill_workbook_meta

    wb, regions = fill_workbook_meta(tpl_file, mt, fields, records)
    slot_pt = _SLOT_PT.get(paper)
    pages = []
    for ws, region in zip(wb.worksheets, regions):
        pad_before, pad_count, pad_h = (None, 0, 0)
        if slot_pt and region:
            pad_before, pad_count, pad_h = _pad_plan(ws, region, slot_pt)
        table = _sheet_html(ws, pad_before=pad_before, pad_count=pad_count, pad_h=pad_h)
        pages.append(f'<div class="page"><div class="doc">{table}</div></div>')
    page_css = _PAGE_CSS.get(paper, _PAGE_CSS["a4"])
    pages_html = "".join(pages)
    return f"""<!DOCTYPE html>
<html lang="zh"><head><meta charset="utf-8"><title>{_html.escape(title)}</title>
<style>
  body {{ margin: 0; background: #e8e8e8; font-family: 'Microsoft YaHei', SimSun, sans-serif; }}
  .page {{ background: #fff; width: 190mm; min-height: 100mm; margin: 12px auto; padding: 6mm 4mm;
           box-shadow: 0 1px 4px rgba(0,0,0,.25); box-sizing: border-box; }}
  .doc table {{ margin: 0 auto; }}
  .doc td {{ overflow: hidden; padding: 1px 3px; }}
  .doc td .vt {{ position: absolute; inset: 1px; overflow: hidden; text-align: center;
                 display: flex; align-items: center; justify-content: center; }}
  {page_css}
  @media print {{
    body {{ background: #fff; }}
    .page {{ width: auto; min-height: 0; margin: 0; padding: 0; box-shadow: none; }}
  }}
  @page {{ size: A4; margin: 10mm; }}
</style></head>
<body>{pages_html}</body></html>"""
