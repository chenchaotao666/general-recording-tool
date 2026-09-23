"""记录/模板导出为 xlsx（AI 助手生成 Excel 用；后续记录页导出可复用）。"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


def _save(wb: Workbook) -> BytesIO:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_blank_xlsx(label: str, fields: list[dict], sample_rows: list[list]) -> BytesIO:
    """空白模板：表头（加粗）+ 可选示例行。fields: [{field_name, label, ...}]"""
    wb = Workbook()
    ws = wb.active
    ws.title = (label or "Sheet1")[:31]
    head = Font(bold=True)
    for ci, f in enumerate(fields, start=1):
        ws.cell(row=1, column=ci, value=f.get("label") or f.get("field_name")).font = head
    for ri, row in enumerate(sample_rows or [], start=2):
        for ci, v in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=v)
    return _save(wb)


def export_records_xlsx(title: str, columns: list[dict], rows: list[dict]) -> BytesIO:
    """记录清单导出：表头行 + 数据行。columns: [{prop, label}]"""
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "记录")[:31]
    head = Font(bold=True)
    for ci, c in enumerate(columns, start=1):
        ws.cell(row=1, column=ci, value=c["label"]).font = head
    for ri, r in enumerate(rows, start=2):
        for ci, c in enumerate(columns, start=1):
            ws.cell(row=ri, column=ci, value=r.get(c["prop"]))
    return _save(wb)
