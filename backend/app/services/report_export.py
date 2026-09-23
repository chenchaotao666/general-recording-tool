"""报表导出：Excel（openpyxl，含原生图表）和自包含 HTML（echarts CDN 渲染 + 数据表兜底）。"""
import html as html_mod
import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

DEFAULT_ECHARTS_CDN = "https://cdn.jsdelivr.net/npm/echarts@5.5.0/dist/echarts.min.js"


def _esc(v) -> str:
    return html_mod.escape("" if v is None else str(v))


# ---------- Excel ----------

def export_xlsx(result: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "报表"
    title_font = Font(bold=True, size=14)
    head_font = Font(bold=True)
    row = 1

    ws.cell(row=row, column=1, value=result["name"]).font = title_font
    row += 1
    ws.cell(row=row, column=1, value=f"{result['range']['label']} · 生成于 {result['generated_at']}")
    row += 2

    for b in result["blocks"]:
        t = b["type"]
        if t == "stat":
            ws.cell(row=row, column=1, value=b["title"]).font = head_font
            ws.cell(row=row, column=2, value=b["value"])
            row += 1
        elif t == "chart":
            row += 1
            ws.cell(row=row, column=1, value=b["title"]).font = head_font
            row += 1
            data_start = row
            series = b.get("series") or [{"name": "值", "values": b.get("values") or []}]
            ws.cell(row=row, column=1, value="分组").font = head_font
            for si, s in enumerate(series):
                ws.cell(row=row, column=2 + si, value=s["name"]).font = head_font
            row += 1
            for li, label in enumerate(b["labels"]):
                ws.cell(row=row, column=1, value=label)
                for si, s in enumerate(series):
                    ws.cell(row=row, column=2 + si, value=s["values"][li])
                row += 1
            data_end = row - 1
            if data_end >= data_start:
                cls = {"bar": BarChart, "line": LineChart, "pie": PieChart, "area": AreaChart}[b["chart_type"]]
                chart = cls()
                chart.title = b["title"]
                if b.get("stack") and len(series) > 1 and b["chart_type"] in ("bar", "line", "area"):
                    chart.grouping = "stacked"
                data = Reference(ws, min_col=2, min_row=data_start,
                                 max_col=1 + len(series), max_row=data_end)
                cats = Reference(ws, min_col=1, min_row=data_start + 1, max_row=data_end)
                chart.add_data(data, titles_from_data=True)
                chart.set_categories(cats)
                chart.width, chart.height = 16, 9
                ws.add_chart(chart, f"{get_column_letter(3 + len(series))}{data_start}")
        elif t == "pivot":
            row += 1
            ws.cell(row=row, column=1, value=b["title"]).font = head_font
            row += 1
            totals = b.get("totals")
            ws.cell(row=row, column=1, value="行＼列").font = head_font
            for ci, cl in enumerate(b["col_labels"]):
                ws.cell(row=row, column=2 + ci, value=cl).font = head_font
            if totals:
                ws.cell(row=row, column=2 + len(b["col_labels"]), value="合计").font = head_font
            row += 1
            for ri, rl in enumerate(b["row_labels"]):
                ws.cell(row=row, column=1, value=rl)
                for ci, v in enumerate(b["cells"][ri]):
                    ws.cell(row=row, column=2 + ci, value=v)
                if totals:
                    ws.cell(row=row, column=2 + len(b["col_labels"]), value=b["row_totals"][ri]).font = head_font
                row += 1
            if totals:
                ws.cell(row=row, column=1, value="合计").font = head_font
                for ci, v in enumerate(b["col_totals"]):
                    ws.cell(row=row, column=2 + ci, value=v).font = head_font
                ws.cell(row=row, column=2 + len(b["col_labels"]), value=b["grand_total"]).font = head_font
                row += 1
        elif t == "table":
            row += 1
            ws.cell(row=row, column=1, value=f"{b['title']}（共 {b['total']} 条" + ("，仅导出前 %d 条" % len(b["rows"]) if b["truncated"] else "）")).font = head_font
            row += 1
            for ci, c in enumerate(b["columns"], start=1):
                ws.cell(row=row, column=ci, value=c["label"]).font = head_font
            row += 1
            for r in b["rows"]:
                for ci, c in enumerate(b["columns"], start=1):
                    ws.cell(row=row, column=ci, value=r.get(c["prop"]))
                row += 1
        elif t == "text":
            row += 1
            ws.cell(row=row, column=1, value=b["content"])
            row += 1
        row += 1

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 16
    for ci in range(3, 10):
        ws.column_dimensions[get_column_letter(ci)].width = 14

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------- HTML ----------

_HTML_SHELL = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>{title}</title>
<script src="{cdn}"></script>
<style>
body {{ font-family: 'Helvetica Neue', 'PingFang SC', 'Microsoft YaHei', Arial, sans-serif;
       max-width: 960px; margin: 0 auto; padding: 24px 16px; color: #303133; background: #f5f7fa; }}
.meta {{ color: #909399; font-size: 13px; margin: 4px 0 20px; }}
.stats {{ display: flex; flex-wrap: wrap; gap: 12px; margin-bottom: 8px; }}
.stat {{ background: #fff; border-radius: 8px; padding: 16px 24px; min-width: 140px;
         box-shadow: 0 1px 3px rgba(0,0,0,.06); }}
.stat .t {{ font-size: 13px; color: #909399; }}
.stat .v {{ font-size: 28px; font-weight: 600; margin-top: 4px; }}
.block {{ background: #fff; border-radius: 8px; padding: 16px 20px; margin: 16px 0;
          box-shadow: 0 1px 3px rgba(0,0,0,.06); }}
.block h3 {{ margin: 0 0 12px; font-size: 15px; }}
.chart {{ width: 100%; height: 320px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
th, td {{ border: 1px solid #e4e7ed; padding: 6px 10px; text-align: left; }}
th {{ background: #f5f7fa; }}
.note {{ color: #909399; font-size: 12px; }}
.text-block {{ color: #606266; line-height: 1.7; white-space: pre-wrap; }}
</style>
</head>
<body>
<h2>{title}</h2>
<div class="meta">{range_label} · 生成于 {generated_at}</div>
{body}
<script>
const REPORT = {data_json};
if (window.echarts) {{
  REPORT.blocks.filter(b => b.type === 'chart').forEach(b => {{
    const el = document.getElementById('chart-' + b.id);
    if (!el) return;
    const ch = echarts.init(el);
    let option;
    if (b.chart_type === 'pie') {{
      option = {{ tooltip: {{ trigger: 'item' }}, legend: {{ bottom: 0 }},
        series: [{{ type: 'pie', radius: ['35%', '65%'],
          data: b.labels.map((l, i) => ({{ name: l, value: b.values[i] }})) }}] }};
    }} else {{
      const seriesList = (b.series && b.series.length) ? b.series : [{{ name: '值', values: b.values }}];
      const stack = b.stack && seriesList.length > 1 ? 'total' : undefined;
      option = {{ tooltip: {{ trigger: 'axis' }}, grid: {{ left: 48, right: 24, top: 24, bottom: seriesList.length > 1 ? 56 : 48 }},
        legend: seriesList.length > 1 ? {{ bottom: 0 }} : undefined,
        xAxis: {{ type: 'category', data: b.labels }},
        yAxis: {{ type: 'value' }},
        series: seriesList.map(s => ({{
          name: s.name, type: b.chart_type === 'area' ? 'line' : b.chart_type, data: s.values,
          smooth: true, barMaxWidth: 40,
          ...(stack ? {{ stack }} : {{}}),
          ...(b.chart_type === 'area' ? {{ areaStyle: {{}} }} : {{}}),
        }})) }};
    }}
    ch.setOption(option);
    window.addEventListener('resize', () => ch.resize());
  }});
}}
</script>
</body>
</html>
"""


def export_html(result: dict, echarts_cdn: str | None = None) -> str:
    """单文件 HTML：echarts CDN 渲染图表；每个图表下附数据表兜底（CDN 不可达时数字仍可见）。"""
    parts = []
    stats = [b for b in result["blocks"] if b["type"] == "stat"]
    if stats:
        cards = "".join(
            f'<div class="stat"><div class="t">{_esc(b["title"])}</div><div class="v">{_esc(b["value"])}</div></div>'
            for b in stats
        )
        parts.append(f'<div class="stats">{cards}</div>')

    for b in result["blocks"]:
        t = b["type"]
        if t == "stat":
            continue
        if t == "chart":
            series = b.get("series") or [{"name": "值", "values": b.get("values") or []}]
            if len(series) > 1:
                head = "<tr><th>分组</th>" + "".join(f"<th>{_esc(s['name'])}</th>" for s in series) + "</tr>"
                rows = "".join(
                    f"<tr><td>{_esc(l)}</td>" + "".join(f"<td>{_esc(s['values'][i])}</td>" for s in series) + "</tr>"
                    for i, l in enumerate(b["labels"])
                )
            else:
                head = "<tr><th>分组</th><th>值</th></tr>"
                rows = "".join(
                    f"<tr><td>{_esc(l)}</td><td>{_esc(v)}</td></tr>" for l, v in zip(b["labels"], series[0]["values"])
                )
            parts.append(
                f'<div class="block"><h3>{_esc(b["title"])}</h3>'
                f'<div class="chart" id="chart-{_esc(b["id"])}"></div>'
                f'<details><summary class="note">数据明细</summary>'
                f"<table>{head}{rows}</table></details></div>"
            )
        elif t == "pivot":
            totals = b.get("totals")
            head = "<th>行＼列</th>" + "".join(f"<th>{_esc(c)}</th>" for c in b["col_labels"])
            if totals:
                head += "<th>合计</th>"
            rows = ""
            for ri, rl in enumerate(b["row_labels"]):
                cells = "".join(f"<td>{_esc(v)}</td>" for v in b["cells"][ri])
                if totals:
                    cells += f"<td><b>{_esc(b['row_totals'][ri])}</b></td>"
                rows += f"<tr><td>{_esc(rl)}</td>{cells}</tr>"
            if totals:
                cells = "".join(f"<td><b>{_esc(v)}</b></td>" for v in b["col_totals"])
                cells += f"<td><b>{_esc(b['grand_total'])}</b></td>"
                rows += f'<tr><td><b>合计</b></td>{cells}</tr>'
            parts.append(
                f'<div class="block"><h3>{_esc(b["title"])}</h3>'
                f"<table><tr>{head}</tr>{rows}</table></div>"
            )
        elif t == "table":
            head = "".join(f"<th>{_esc(c['label'])}</th>" for c in b["columns"])
            rows = "".join(
                "<tr>" + "".join(f"<td>{_esc(r.get(c['prop']))}</td>" for c in b["columns"]) + "</tr>"
                for r in b["rows"]
            )
            note = f'<div class="note">共 {b["total"]} 条，仅显示前 {len(b["rows"])} 条</div>' if b["truncated"] else ""
            parts.append(
                f'<div class="block"><h3>{_esc(b["title"])}</h3>'
                f"<table><tr>{head}</tr>{rows}</table>{note}</div>"
            )
        elif t == "text":
            parts.append(f'<div class="block"><h3>{_esc(b["title"])}</h3><div class="text-block">{_esc(b["content"])}</div></div>')

    return _HTML_SHELL.format(
        title=_esc(result["name"]),
        range_label=_esc(result["range"]["label"]),
        generated_at=_esc(result["generated_at"]),
        cdn=_esc(echarts_cdn or DEFAULT_ECHARTS_CDN),
        body="".join(parts),
        data_json=json.dumps(result, ensure_ascii=False).replace("</", "<\\/"),
    )
