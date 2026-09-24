"""冒烟测试：不走 LLM，验证 上传解析 -> 建表导入 -> 动态CRUD -> 删表 全链路。"""
import io
import os
import sys

# 用独立的测试数据库，避免污染 grt.db
os.environ["GRT_DATABASE_URL"] = "sqlite:///./smoke_test.db"

import openpyxl
from fastapi.testclient import TestClient

from app.main import app

cm = TestClient(app)
client = cm.__enter__()  # 进入上下文以触发 lifespan（建元数据表 + 种子管理员）

# 0. 登录：未带 token 一律 401；错误密码 401；默认管理员 admin/admin123 可登录
r = client.get("/api/tables")
assert r.status_code == 401, r.text
r = client.post("/api/auth/login", json={"username": "admin", "password": "wrong"})
assert r.status_code == 401, r.text
r = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
assert r.status_code == 200, r.text
token = r.json()["token"]
assert r.json()["user"]["role"] == "admin"
assert {"share", "create_physical_table", "max_tables"} <= set(r.json()["user"]["perms"])  # admin 拥有全部权限
client.headers.update({"Authorization": f"Bearer {token}"})
print("登录成功，token 已注入请求头")

# 注册：新用户可注册并直接登录；重名/弱密码被拒绝
r = client.post("/api/auth/register", json={"username": "zhangsan", "password": "secret123"})
assert r.status_code == 200 and r.json()["user"]["role"] == "user", r.text
assert "share" not in r.json()["user"]["perms"]  # 普通用户默认无分享权限
r = client.post("/api/auth/register", json={"username": "zhangsan", "password": "secret123"})
assert r.status_code == 400, r.text
r = client.post("/api/auth/register", json={"username": "lisi", "password": "123"})
assert r.status_code == 400, r.text
r = client.post("/api/auth/login", json={"username": "zhangsan", "password": "secret123"})
assert r.status_code == 200, r.text
print("注册/重名校验/弱密码校验通过")

# 修改密码：原密码错误 400，弱密码 400，改后旧密码失效、新密码可登录（以 zhangsan 身份操作）
r = client.post("/api/auth/login", json={"username": "zhangsan", "password": "secret123"})
zs_headers = {"Authorization": f"Bearer {r.json()['token']}"}
saved_headers = dict(client.headers)
client.headers = zs_headers
r = client.put("/api/auth/password", json={"old_password": "wrong", "new_password": "newpass123"})
assert r.status_code == 400, r.text
r = client.put("/api/auth/password", json={"old_password": "secret123", "new_password": "123"})
assert r.status_code == 400, r.text
r = client.put("/api/auth/password", json={"old_password": "secret123", "new_password": "newpass123"})
assert r.status_code == 200, r.text
client.headers = saved_headers
r = client.post("/api/auth/login", json={"username": "zhangsan", "password": "secret123"})
assert r.status_code == 401, r.text
r = client.post("/api/auth/login", json={"username": "zhangsan", "password": "newpass123"})
assert r.status_code == 200, r.text
print("修改密码通过")

# Excel 序列日期：单元格未设日期格式时读出为数字（45200 → 2023-10-01）
from app.services.typemap import coerce_value

ok, val, err = coerce_value(45200, "date")
assert ok and str(val) == "2023-10-01", (val, err)
ok, val, err = coerce_value("45200", "date")
assert ok and str(val) == "2023-10-01", (val, err)
ok, val, err = coerce_value(45200.5, "datetime")
assert ok and val.hour == 12, (val, err)
ok, _, _ = coerce_value("2026", "date")   # 年份不当作序列日期
assert not ok
print("Excel 序列日期转换通过")

# 1. 构造测试 Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "客户"
ws.append(["客户姓名", "联系电话", "下次跟进日期", "意向金额", "已成交"])
ws.append(["张三", "13800001234", "2026-10-01", 15000.5, "是"])
ws.append(["李四", "13900005678", "2026/10/05", 8000, "否"])
ws.append(["王五", "13700009012", "不是日期", None, "是"])  # 这一行日期转换失败
buf = io.BytesIO()
wb.save(buf)
buf.seek(0)

# 2. 上传解析
r = client.post("/api/excel/upload", files={"file": ("客户跟进.xlsx", buf.getvalue())})
assert r.status_code == 200, r.text
up = r.json()
print("upload:", up["file_name"], "sheets:", [(s["name"], s["header_row"], s["total_rows"]) for s in up["sheets"]])
assert up["sheets"][0]["headers"] == ["客户姓名", "联系电话", "下次跟进日期", "意向金额", "已成交"]

# 3. 确认结构建表 + 导入（模拟用户确认后的字段定义）
fields = [
    {"source_header": "客户姓名", "field_name": "customer_name", "label": "客户姓名", "data_type": "varchar", "length": 64, "nullable": False, "widget": "input"},
    {"source_header": "联系电话", "field_name": "phone", "label": "联系电话", "data_type": "varchar", "length": 32, "widget": "input"},
    {"source_header": "下次跟进日期", "field_name": "next_follow_date", "label": "下次跟进日期", "data_type": "date", "widget": "date-picker"},
    {"source_header": "意向金额", "field_name": "amount", "label": "意向金额", "data_type": "decimal", "widget": "number"},
    {"source_header": "已成交", "field_name": "is_deal", "label": "已成交", "data_type": "bool", "widget": "switch"},
]
r = client.post("/api/tables", json={
    "label": "客户跟进记录",
    "fields": fields,
    "source": {"file_id": up["file_id"], "sheet_name": "客户", "header_row": 1},
})
assert r.status_code == 200, r.text
created = r.json()
tid = created["table"]["id"]
report = created["import_report"]
print("建表:", created["table"]["name"], "导入:", report["success"], "成功 /", report["failed"], "失败")
assert report["success"] == 2 and report["failed"] == 1, report
assert "无法转换为 date" in report["failures"][0]["error"]

# 4. 动态 CRUD
r = client.get(f"/api/dyn/{tid}/records")
assert r.status_code == 200 and r.json()["total"] == 2, r.text
first = r.json()["items"][0]
print("列表首条:", first)
assert first["next_follow_date"] in ("2026-10-01", "2026-10-05")

# 筛选：contains
r = client.get(f"/api/dyn/{tid}/records", params={"filters": '[{"field":"customer_name","op":"contains","value":"张"}]'})
assert r.json()["total"] == 1, r.text

# 筛选：bool eq + 日期范围（李四 is_deal=false 不命中，只有张三命中）
r = client.get(f"/api/dyn/{tid}/records", params={"filters": '[{"field":"is_deal","op":"eq","value":true},{"field":"next_follow_date","op":"gte","value":"2026-10-01"}]'})
assert r.json()["total"] == 1, r.text

# 新增
r = client.post(f"/api/dyn/{tid}/records", json={"customer_name": "赵六", "next_follow_date": "2026-11-01", "amount": 3000, "is_deal": False})
assert r.status_code == 200, r.text
rid = r.json()["id"]

# 必填校验
r = client.post(f"/api/dyn/{tid}/records", json={"phone": "123"})
assert r.status_code == 422, r.text

# 编辑 + 删除
r = client.put(f"/api/dyn/{tid}/records/{rid}", json={"customer_name": "赵六改"})
assert r.status_code == 200 and r.json()["customer_name"] == "赵六改", r.text
r = client.delete(f"/api/dyn/{tid}/records/{rid}")
assert r.status_code == 200, r.text

# 5. 图片识别（mock LLM provider，验证压缩/校验/留痕链路）
from PIL import Image

import app.services.llm.gateway as gw


class FakeProvider:
    def complete(self, prompt, system=None):
        return "{}"

    def recognize(self, images, prompt, system=None):
        assert len(images) == 1 and images[0][:2] == b"\xff\xd8"  # 已压缩为 JPEG
        return '{"fields": {"customer_name": "图片客户", "amount": 6666, "next_follow_date": "2026-12-01", "unknown_field": "x"}, "conflicts": [], "notes": "识别完成"}'


gw.get_default_provider = lambda db: FakeProvider()

img = Image.new("RGB", (300, 120), "white")
img_buf = io.BytesIO()
img.save(img_buf, "PNG")

r = client.post(
    "/api/vision/recognize",
    data={"table_id": str(tid), "current": '{"customer_name": "旧名字"}'},
    files={"images": ("test.png", img_buf.getvalue(), "image/png")},
)
assert r.status_code == 200, r.text
vr = r.json()
print("识别结果:", vr["fields"], "冲突:", vr["conflicts"])
assert vr["fields"]["customer_name"] == "图片客户"
assert vr["fields"]["amount"] == 6666.0
assert vr["fields"]["next_follow_date"] == "2026-12-01"
assert "unknown_field" not in vr["fields"]  # 字段清单外的字段被丢弃
assert any(c["field_name"] == "customer_name" for c in vr["conflicts"])  # 与当前值冲突被检出

# 采纳留痕
r = client.put(f"/api/vision/logs/{vr['log_id']}/adopt", json={"adopted": ["amount", "next_follow_date"]})
assert r.status_code == 200, r.text

# 6. 任务引擎：结构化条件 + 站内通知
r = client.post("/api/tasks", json={
    "name": "10月前需要跟进的客户",
    "table_id": tid,
    "enabled": True,
    "condition_mode": "structured",
    "condition": {"logic": "AND", "rules": [
        {"field": "next_follow_date", "op": "lte", "value": "2026-10-31"},
        {"field": "is_deal", "op": "eq", "value": False},
    ]},
    "schedule": {"type": "interval", "minutes": 60},
    "action": {"type": "notify", "template": "客户【{customer_name}】计划 {next_follow_date} 跟进",
               "recipients": {"type": "fixed", "value": ""}},
    "cooldown_hours": 24,
})
assert r.status_code == 200, r.text
rule = r.json()
print("任务创建:", rule["name"])

# 试运行：李四未成交且跟进日期<=10-31，命中 1 条
r = client.post(f"/api/tasks/{rule['id']}/test")
assert r.status_code == 200, r.text
tr = r.json()
print("试运行: 命中", tr["matched"], "将触发", tr["would_fire"])
assert tr["matched"] == 1 and tr["would_fire"] == 1, tr
assert tr["samples"][0]["customer_name"] == "李四"

# 手动执行：生成站内通知
r = client.post(f"/api/tasks/{rule['id']}/run")
assert r.status_code == 200, r.text
assert r.json()["fired"] == 1 and r.json()["failed"] == 0, r.text
r = client.get("/api/notify/unread_count")
assert r.json()["count"] == 1, r.text
r = client.get("/api/notify")
assert "李四" in r.json()[0]["content"] and "2026-10-05" in r.json()[0]["content"], r.text

# 冷却语义：手动执行跳过冷却（设计如此，便于立即验证）；定时调度路径遵守冷却窗口
r = client.post(f"/api/tasks/{rule['id']}/run")
assert r.json()["matched"] == 1 and r.json()["fired"] == 1, r.text
from app.services.task_engine import execute_rule as _exec_rule
r2 = _exec_rule(rule["id"], trigger="schedule")  # 冷却期内 → 不再触发
assert r2["matched"] == 1 and r2["fired"] == 0, r2

# 时间操作符：within_days（跟进日期在未来 60 天内 → 两条都命中）
r = client.post("/api/tasks", json={
    "name": "历史任务", "table_id": tid,
    "condition_mode": "structured",
    "condition": {"logic": "AND", "rules": [{"field": "next_follow_date", "op": "within_days", "value": 60}]},
    "schedule": {"type": "cron", "expr": "0 9 * * *"},
    "action": {"type": "notify", "template": "x", "recipients": {"type": "fixed", "value": ""}},
})
r2 = client.post(f"/api/tasks/{r.json()['id']}/test")
assert r2.json()["matched"] == 2, r2.text
client.delete(f"/api/tasks/{r.json()['id']}")

# older_than_days：created_at 早于 0 天前（即此刻之前创建）→ 两条都命中
r = client.post("/api/tasks", json={
    "name": "older", "table_id": tid, "condition_mode": "structured",
    "condition": {"logic": "AND", "rules": [{"field": "created_at", "op": "older_than_days", "value": 0}]},
    "schedule": {"type": "interval", "minutes": 60},
    "action": {"type": "notify", "template": "x", "recipients": {"type": "fixed", "value": ""}},
})
r2 = client.post(f"/api/tasks/{r.json()['id']}/test")
assert r2.json()["matched"] == 2, r2.text
client.delete(f"/api/tasks/{r.json()['id']}")

# 7. LLM 判断条件（mock provider 返回 matched_ids）
class JudgeFakeProvider(FakeProvider):
    def complete(self, prompt, system=None):
        return '{"matched_ids": [2], "reason": "test"}'

gw.get_default_provider = lambda db: JudgeFakeProvider()
r = client.post("/api/tasks", json={
    "name": "LLM判断", "table_id": tid,
    "condition_mode": "llm",
    "condition": {"description": "意向金额较高的客户"},
    "schedule": {"type": "interval", "minutes": 30},
    "action": {"type": "notify", "template": "高额客户 {customer_name}", "recipients": {"type": "fixed", "value": ""}},
})
assert r.status_code == 200, r.text
r2 = client.post(f"/api/tasks/{r.json()['id']}/test")
assert r2.json()["matched"] == 1 and r2.json()["samples"][0]["id"] == 2, r2.text
client.delete(f"/api/tasks/{r.json()['id']}")

# 非法 cron 被拒绝
r = client.post("/api/tasks", json={
    "name": "bad", "table_id": tid, "condition_mode": "structured",
    "condition": {"logic": "AND", "rules": [{"field": "customer_name", "op": "contains", "value": "x"}]},
    "schedule": {"type": "cron", "expr": "not-a-cron"},
    "action": {"type": "notify", "template": "x", "recipients": {"type": "fixed", "value": ""}},
})
assert r.status_code == 400, r.text

# 8. 通用设置（SMTP 密码保留逻辑）
r = client.put("/api/settings/general", json={"smtp": {"host": "smtp.example.com", "port": 465, "username": "a@b.com", "password": "secret", "use_ssl": True}})
assert r.status_code == 200, r.text
r = client.get("/api/settings/general")
assert r.json()["smtp"]["host"] == "smtp.example.com" and r.json()["smtp"]["password"] is None and r.json()["smtp"]["has_password"]
r = client.put("/api/settings/general", json={"smtp": {"host": "smtp2.example.com"}})
r = client.get("/api/settings/general")
assert r.json()["smtp"]["host"] == "smtp2.example.com" and r.json()["smtp"]["has_password"]  # 密码保留

# 运行日志（2 次手动 + 1 次定时）
r = client.get(f"/api/tasks/{rule['id']}/runs")
assert r.status_code == 200 and len(r.json()) == 3 and r.json()[0]["trigger"] == "schedule", r.text

# 通知已读
r = client.post("/api/notify/read", json={"all": True})
assert client.get("/api/notify/unread_count").json()["count"] == 0

# 删除任务
r = client.delete(f"/api/tasks/{rule['id']}")
assert r.status_code == 200

# 9. 表列表 + 删除
r = client.get("/api/tables")
assert r.status_code == 200 and r.json()[0]["record_count"] == 2, r.text
r = client.delete(f"/api/tables/{tid}")
assert r.status_code == 200, r.text
r = client.get(f"/api/dyn/{tid}/records")
assert r.status_code == 404

# ========== 混合存储 + 多租户 + 分享权限 ==========

PFIELDS = [
    {"field_name": "customer_name", "label": "客户姓名", "data_type": "varchar", "length": 64, "nullable": False, "widget": "input"},
    {"field_name": "next_follow_date", "label": "下次跟进日期", "data_type": "date", "widget": "date-picker"},
    {"field_name": "amount", "label": "意向金额", "data_type": "decimal", "widget": "number"},
    {"field_name": "is_deal", "label": "已成交", "data_type": "bool", "widget": "switch"},
]
PRECS = [
    {"customer_name": "张三", "next_follow_date": "2026-10-01", "amount": 15000.5, "is_deal": True},
    {"customer_name": "李四", "next_follow_date": "2026-10-05", "amount": 8000, "is_deal": False},
    {"customer_name": "王五", "next_follow_date": "2026-09-15", "amount": 20000.25, "is_deal": True},
    {"customer_name": "赵六", "amount": 500, "is_deal": False},
]


def make_table(label, storage_mode):
    r = client.post("/api/tables", json={"label": label, "storage_mode": storage_mode, "fields": PFIELDS})
    assert r.status_code == 200, r.text
    return r.json()["table"]["id"]


def fill(t):
    for rec in PRECS:
        r = client.post(f"/api/dyn/{t}/records", json=rec)
        assert r.status_code == 200, r.text


# 10. 双模式 parity：同数据集 json vs physical，筛选/排序/报表逐项一致
tj = make_table("parity-json", "json")
tp = make_table("parity-physical", "physical")
fill(tj)
fill(tp)

FILTERS = [
    '[{"field":"customer_name","op":"contains","value":"张"}]',
    '[{"field":"amount","op":"gt","value":9000}]',
    '[{"field":"next_follow_date","op":"gte","value":"2026-10-01"}]',
    '[{"field":"is_deal","op":"eq","value":true}]',
    '[{"field":"next_follow_date","op":"null"}]',
    '[{"field":"customer_name","op":"startswith","value":"李"}]',
    '[{"field":"amount","op":"in","value":"8000,500"}]',
    '[{"field":"next_follow_date","op":"older_than_days","value":3}]',
]
def _strip_ts(items):
    # json/physical 是两次独立插入，created_at/updated_at 必然不同，比较时剔除
    return [{k: v for k, v in r.items() if k not in ("created_at", "updated_at")} for r in items]


for flt in FILTERS:
    a = client.get(f"/api/dyn/{tj}/records", params={"filters": flt}).json()
    b = client.get(f"/api/dyn/{tp}/records", params={"filters": flt}).json()
    assert a["total"] == b["total"] and _strip_ts(a["items"]) == _strip_ts(b["items"]), (flt, a, b)
print("parity 筛选一致:", len(FILTERS), "组")

for params in ({"sort_by": "amount", "sort_order": "asc"}, {"sort_by": "amount", "sort_order": "desc"},
               {"sort_by": "next_follow_date", "sort_order": "desc"}, {}):
    a = client.get(f"/api/dyn/{tj}/records", params=params).json()
    b = client.get(f"/api/dyn/{tp}/records", params=params).json()
    assert _strip_ts(a["items"]) == _strip_ts(b["items"]), params
print("parity 排序一致")

BLOCKS = [
    {"id": "b1", "type": "stat", "title": "总额", "agg": "sum", "field": "amount", "filters": {"logic": "AND", "rules": []}},
    {"id": "b2", "type": "chart", "title": "成交对比", "chart_type": "bar", "group": {"kind": "field", "field": "is_deal"}, "agg": "sum", "field": "amount", "top_n": 8, "filters": {"logic": "AND", "rules": []}},
    {"id": "b3", "type": "chart", "title": "按月", "chart_type": "line", "group": {"kind": "month", "field": "next_follow_date"}, "agg": "count", "filters": {"logic": "AND", "rules": []}},
    {"id": "b4", "type": "table", "title": "明细", "columns": ["customer_name", "amount", "next_follow_date"], "sort_by": "amount", "sort_order": "desc", "limit": 10, "filters": {"logic": "AND", "rules": []}},
]
rep_ids = []
for t in (tj, tp):
    r = client.post("/api/reports", json={
        "name": "parity", "table_id": t, "enabled": False,
        "range": {"mode": "this_month", "date_field": "created_at"},
        "blocks": BLOCKS, "schedule": {}, "push": {},
    })
    assert r.status_code == 200, r.text
    rep_ids.append(r.json()["id"])
ra = client.post(f"/api/reports/{rep_ids[0]}/run").json()
rb = client.post(f"/api/reports/{rep_ids[1]}/run").json()
for ba, bb in zip(ra["blocks"], rb["blocks"]):
    ba.pop("id", None)
    bb.pop("id", None)
    assert ba == bb, (ba, bb)
print("parity 报表四区块一致")
for rid_ in rep_ids:
    client.delete(f"/api/reports/{rid_}")

# 记录导出：整表导出 + 带筛选导出，行数与口径一致
r = client.get(f"/api/dyn/{tj}/export")
assert r.status_code == 200 and r.content[:2] == b"PK", r.status_code
assert "attachment" in r.headers.get("content-disposition", "")
assert openpyxl.load_workbook(io.BytesIO(r.content)).active.max_row == 5  # 表头 + 4 条
r = client.get(f"/api/dyn/{tj}/export", params={"filters": '[{"field":"is_deal","op":"eq","value":true}]'})
assert r.status_code == 200, r.text
assert openpyxl.load_workbook(io.BytesIO(r.content)).active.max_row == 3  # 表头 + 张三/王五
print("记录导出（整表/筛选）通过")

# 导出回归：>200 条不被列表页大小上限（MAX_PAGE_SIZE=200）截断，json/physical 两种引擎都验证
exp_ids = []
for mode in ("json", "physical"):
    et = make_table(f"export-cap-{mode}", mode)
    for i in range(205):
        r = client.post(f"/api/dyn/{et}/records", json={"customer_name": f"客{i:03d}"})
        assert r.status_code == 200, r.text
    # 列表接口仍受 200 页上限保护（FastAPI 参数校验 422；page_size=200 正常返回 200 条）
    assert client.get(f"/api/dyn/{et}/records", params={"page_size": 500}).status_code == 422
    assert len(client.get(f"/api/dyn/{et}/records", params={"page_size": 200}).json()["items"]) == 200
    r = client.get(f"/api/dyn/{et}/export")
    assert r.status_code == 200, r.text
    assert openpyxl.load_workbook(io.BytesIO(r.content)).active.max_row == 206  # 表头 + 205 条
    exp_ids.append(et)
for et in exp_ids:
    client.delete(f"/api/tables/{et}")
print("导出上限（>200 条不截断，双引擎）通过")

# 10.5 报表 P0：新时间口径 / 去重计数 / 占比 / 环比 / 查看端筛选
BLOCKS_P0 = [
    {"id": "b1", "type": "stat", "title": "记录数", "agg": "count", "compare": True,
     "filters": {"logic": "AND", "rules": []}},
    {"id": "b2", "type": "stat", "title": "去重客户", "agg": "count_distinct", "field": "customer_name",
     "filters": {"logic": "AND", "rules": []}},
    {"id": "b3", "type": "stat", "title": "成交占比", "agg": "ratio",
     "filters": {"logic": "AND", "rules": [{"field": "is_deal", "op": "eq", "value": True}]}},
    {"id": "b4", "type": "chart", "title": "按成交去重客户", "chart_type": "bar",
     "group": {"kind": "field", "field": "is_deal"}, "agg": "count_distinct", "field": "customer_name",
     "filters": {"logic": "AND", "rules": []}},
]

# 校验：图表不支持占比、去重计数必须有字段、查看端筛选字段必须是表内字段
for bad_blocks, bad_ff in (
    ([{"id": "x", "type": "chart", "chart_type": "bar", "group": {"kind": "field", "field": "is_deal"}, "agg": "ratio"}], None),
    ([{"id": "x", "type": "stat", "agg": "count_distinct"}], None),
    ([{"id": "x", "type": "stat", "agg": "count"}], ["ghost_field"]),
):
    r = client.post("/api/reports", json={
        "name": "bad", "table_id": tj, "range": {"mode": "today"},
        "blocks": bad_blocks, "filter_fields": bad_ff or [],
    })
    assert r.status_code == 400, r.text

p0_ids = []
for t in (tj, tp):
    r = client.post("/api/reports", json={
        "name": "P0报表", "table_id": t, "enabled": False,
        "range": {"mode": "today", "date_field": "created_at"},
        "blocks": BLOCKS_P0, "filter_fields": ["customer_name", "is_deal"],
        "schedule": {}, "push": {},
    })
    assert r.status_code == 200 and r.json()["filter_fields"] == ["customer_name", "is_deal"], r.text
    p0_ids.append(r.json()["id"])


def _blocks_by_id(res):
    return {b["id"]: b for b in res["blocks"]}


# 两种存储引擎结果一致 + 新聚合语义（tj=json / tp=physical，数据均为今天创建）
runs = [client.post(f"/api/reports/{rid}/run").json() for rid in p0_ids]
for res in runs:
    b = _blocks_by_id(res)
    assert res["range"]["label"].startswith("今天"), res["range"]
    assert b["b1"]["value"] == 4 and b["b1"]["compare"] == {"prev": 0, "delta": 4, "delta_pct": None}, b["b1"]
    assert b["b2"]["value"] == 4      # 4 个不同客户名
    assert b["b3"]["value"] == 50.0   # 4 条中 2 条成交
    assert sorted(b["b4"]["values"]) == [2, 2]
for bid in ("b1", "b2", "b3"):
    assert _blocks_by_id(runs[0])[bid] == _blocks_by_id(runs[1])[bid], bid

# 查看端筛选：只筛已成交 → 记录数 2、占比 100%（占比分母是查看者筛选后的集合）
r = client.post(f"/api/reports/{p0_ids[0]}/run", json={
    "filters": {"logic": "AND", "rules": [{"field": "is_deal", "op": "eq", "value": True}]}})
b = _blocks_by_id(r.json())
assert b["b1"]["value"] == 2 and b["b3"]["value"] == 100.0 and sum(b["b4"]["values"]) == 2, b
# 未声明的字段不允许查看端筛选
r = client.post(f"/api/reports/{p0_ids[0]}/run", json={
    "filters": {"logic": "AND", "rules": [{"field": "amount", "op": "gt", "value": 1}]}})
assert r.status_code == 400, r.text

# 新时间口径
for mode, label_prefix, expect in (
    ("yesterday", "昨天", 0), ("past_7d", "近7天", 4), ("past_30d", "近30天", 4),
    ("this_quarter", "本季度", 4), ("this_year", "今年", 4),
):
    res = client.post(f"/api/reports/{p0_ids[0]}/run", json={"range": {"mode": mode}}).json()
    assert res["range"]["label"].startswith(label_prefix), res["range"]
    assert _blocks_by_id(res)["b1"]["value"] == expect, (mode, res["range"])

# 环比：换业务日期字段 + 自定义区间（10 月 2 条 vs 9 月 1 条），两引擎一致
cmp_range = {"range": {"mode": "custom", "start": "2026-10-01", "end": "2026-10-31", "date_field": "next_follow_date"}}
res_py = client.post(f"/api/reports/{p0_ids[0]}/run", json=cmp_range).json()
res_phy = client.post(f"/api/reports/{p0_ids[1]}/run", json=cmp_range).json()
b1 = _blocks_by_id(res_py)["b1"]
assert b1["value"] == 2 and b1["compare"]["prev"] == 1 and b1["compare"]["delta_pct"] == 100.0, b1
assert b1 == _blocks_by_id(res_phy)["b1"]

for rid_ in p0_ids:
    client.delete(f"/api/reports/{rid_}")
print("报表 P0（口径/去重/占比/环比/查看端筛选）通过")

# 10.6 报表 Webhook 推送：企微/钉钉机器人协议、渠道失败汇总、无渠道报错
from app.services import report_engine as re_mod

# 校验：不支持的类型 / 非 http(s) 地址
for bad_wh in ([{"type": "slack", "url": "https://x"}], [{"type": "wecom", "url": "not-a-url"}]):
    r = client.post("/api/reports", json={
        "name": "bad", "table_id": tj, "range": {"mode": "today"},
        "blocks": [{"id": "x", "type": "stat", "agg": "count"}], "push": {"webhooks": bad_wh},
    })
    assert r.status_code == 400, r.text

r = client.post("/api/reports", json={
    "name": "推送报表", "table_id": tj, "enabled": True,
    "range": {"mode": "today", "date_field": "created_at"}, "blocks": BLOCKS_P0,
    "schedule": {"type": "interval", "minutes": 60},
    "push": {"webhooks": [
        {"type": "wecom", "url": "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=ok1"},
        {"type": "dingtalk", "url": "https://oapi.dingtalk.com/robot/send?access_token=ok2"},
    ]},
})
assert r.status_code == 200, r.text
wh_id = r.json()["id"]

wh_calls = []


class _FakeResp:
    status_code = 200

    def __init__(self, payload):
        self._p = payload
        self.text = str(payload)

    def json(self):
        return self._p


def _fake_post(url, json=None, timeout=None, **kw):
    wh_calls.append((url, json))
    # key=bad 的机器人模拟业务报错；其余成功
    return _FakeResp({"errcode": 93000 if "key=bad" in url else 0, "errmsg": "mock"})


_orig_post = re_mod.httpx.post
re_mod.httpx.post = _fake_post
try:
    r = client.post(f"/api/reports/{wh_id}/test-push")
finally:
    re_mod.httpx.post = _orig_post
assert r.status_code == 200 and r.json()["sent"] == 2, r.text
# 企微/钉钉各发一次 markdown 消息，内容含报表名与统计卡
assert len(wh_calls) == 2
wecom_payload = next(p for u, p in wh_calls if "qyapi" in u)
ding_payload = next(p for u, p in wh_calls if "oapi" in u)
assert wecom_payload["msgtype"] == "markdown" and "推送报表" in wecom_payload["markdown"]["content"]
assert "记录数" in wecom_payload["markdown"]["content"] and "成交占比" in wecom_payload["markdown"]["content"]
assert ding_payload["msgtype"] == "markdown" and ding_payload["markdown"]["title"].startswith("【推送报表】")

# 机器人返回 errcode != 0 → test-push 400，错误落推送日志
r = client.put(f"/api/reports/{wh_id}", json={
    "name": "推送报表", "table_id": tj, "enabled": True,
    "range": {"mode": "today", "date_field": "created_at"}, "blocks": BLOCKS_P0,
    "schedule": {"type": "interval", "minutes": 60},
    "push": {"webhooks": [{"type": "wecom", "url": "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=bad"}]},
})
assert r.status_code == 200, r.text
re_mod.httpx.post = _fake_post
try:
    r = client.post(f"/api/reports/{wh_id}/test-push")
finally:
    re_mod.httpx.post = _orig_post
assert r.status_code == 400 and "93000" in r.json()["detail"], r.text
runs = client.get(f"/api/reports/{wh_id}/runs").json()
assert runs[0]["error"] and "93000" in runs[0]["error"] and runs[0]["sent_count"] == 0, runs[0]

# 邮箱与 webhook 都没配 → 400 提示未配置渠道
r = client.put(f"/api/reports/{wh_id}", json={
    "name": "推送报表", "table_id": tj, "enabled": True,
    "range": {"mode": "today", "date_field": "created_at"}, "blocks": BLOCKS_P0,
    "schedule": {"type": "interval", "minutes": 60}, "push": {},
})
assert r.status_code == 200, r.text
r = client.post(f"/api/reports/{wh_id}/test-push")
assert r.status_code == 400 and "未配置推送渠道" in r.json()["detail"], r.text

client.delete(f"/api/reports/{wh_id}")
print("报表 Webhook 推送（协议/失败汇总/无渠道报错）通过")

# 10.7 报表 P1：多系列图表（多指标 / 二级分组 / 堆叠 / 面积图）
# 校验：饼图不支持多系列、metrics 与 group2 互斥、指标超限、ratio 不可用于图表、二级分组字段约束
for bad in (
    [{"id": "x", "type": "chart", "chart_type": "pie", "group": {"kind": "field", "field": "is_deal"},
      "metrics": [{"agg": "count"}]}],
    [{"id": "x", "type": "chart", "chart_type": "bar", "group": {"kind": "field", "field": "is_deal"},
      "metrics": [{"agg": "count"}], "group2": {"field": "customer_name"}}],
    [{"id": "x", "type": "chart", "chart_type": "bar", "group": {"kind": "field", "field": "is_deal"},
      "metrics": [{"agg": "count"}] * 6}],
    [{"id": "x", "type": "chart", "chart_type": "bar", "group": {"kind": "field", "field": "is_deal"},
      "metrics": [{"agg": "ratio"}]}],
    [{"id": "x", "type": "chart", "chart_type": "bar", "group": {"kind": "field", "field": "is_deal"},
      "group2": {"field": "ghost"}}],
    [{"id": "x", "type": "chart", "chart_type": "bar", "group": {"kind": "field", "field": "is_deal"},
      "group2": {"field": "next_follow_date"}}],
):
    r = client.post("/api/reports", json={"name": "bad", "table_id": tj, "range": {"mode": "today"}, "blocks": bad})
    assert r.status_code == 400, r.text

MS_BLOCKS = [
    {"id": "c1", "type": "chart", "title": "多指标", "chart_type": "bar", "stack": True,
     "group": {"kind": "field", "field": "is_deal"},
     "metrics": [{"agg": "sum", "field": "amount", "title": "总额"},
                 {"agg": "count", "title": "笔数"},
                 {"agg": "count_distinct", "field": "customer_name", "title": "去重客户数"}]},
    {"id": "c2", "type": "chart", "title": "按月×成交", "chart_type": "area",
     "group": {"kind": "month", "field": "next_follow_date"}, "agg": "count",
     "group2": {"field": "is_deal"}},
]
ms_ids = []
for t in (tj, tp):
    r = client.post("/api/reports", json={
        "name": "多系列", "table_id": t, "range": {"mode": "past_30d", "date_field": "created_at"},
        "blocks": MS_BLOCKS,
    })
    assert r.status_code == 200, r.text
    ms_ids.append(r.json()["id"])


def _series_map(b):
    """{label: {系列名: 值}}，行序/系列序不依赖实现，便于两引擎对比"""
    return {str(l): {s["name"]: s["values"][i] for s in b["series"]} for i, l in enumerate(b["labels"])}


runs_ms = [client.post(f"/api/reports/{rid}/run").json() for rid in ms_ids]
for res in runs_ms:
    c1 = _blocks_by_id(res)["c1"]
    assert c1["stack"] is True and len(c1["series"]) == 3
    assert [s["name"] for s in c1["series"]] == ["总额", "笔数", "去重客户数"]
    assert c1["values"] == c1["series"][0]["values"]  # values 兼容字段 = 首系列
    m1 = _series_map(c1)
    # 成交组（张三15000.5+王五20000.25）与未成交组（李四8000+赵六500）
    assert sorted((v["总额"], v["笔数"], v["去重客户数"]) for v in m1.values()) == [(8500, 2, 2), (35000.75, 2, 2)], m1

    c2 = _blocks_by_id(res)["c2"]
    assert c2["chart_type"] == "area" and c2["stack"] is False and len(c2["series"]) == 2
    assert c2["labels"] == ["2026-09", "2026-10", "（空）"], c2["labels"]  # 时间升序，空值最后
    m2 = _series_map(c2)
    deal_name = max(m2["2026-09"], key=lambda n: m2["2026-09"][n])  # 2026-09 只有王五（成交）
    other_name = next(n for n in m2["2026-09"] if n != deal_name)
    assert m2["2026-09"] == {deal_name: 1, other_name: 0}
    assert m2["2026-10"] == {deal_name: 1, other_name: 1}
    assert m2["（空）"] == {deal_name: 0, other_name: 1}

# 两引擎多系列结果一致
for bid in ("c1", "c2"):
    a, b = _blocks_by_id(runs_ms[0])[bid], _blocks_by_id(runs_ms[1])[bid]
    assert a["labels"] == b["labels"] and _series_map(a) == _series_map(b), bid

# 导出兼容多系列：xlsx/html 均正常
r = client.get(f"/api/reports/{ms_ids[0]}/export?format=xlsx")
assert r.status_code == 200, r.text
r = client.get(f"/api/reports/{ms_ids[0]}/export?format=html")
assert r.status_code == 200 and "总额" in r.text and "笔数" in r.text, r.status_code

for rid_ in ms_ids:
    client.delete(f"/api/reports/{rid_}")
print("报表多系列图表（多指标/二级分组/堆叠/面积图）通过")

# 10.8 图表下钻：分组/系列序号取明细记录
DRILL_BLOCKS = MS_BLOCKS + [
    {"id": "t1", "type": "table", "title": "明细", "columns": ["customer_name"], "sort_by": "id", "sort_order": "desc", "limit": 10},
]
drill_ids = []
for t in (tj, tp):
    r = client.post("/api/reports", json={
        "name": "下钻", "table_id": t, "range": {"mode": "past_30d", "date_field": "created_at"},
        "blocks": DRILL_BLOCKS, "filter_fields": ["is_deal", "customer_name"],
    })
    assert r.status_code == 200, r.text
    drill_ids.append(r.json()["id"])

run0 = client.post(f"/api/reports/{drill_ids[0]}/run").json()
c1 = _blocks_by_id(run0)["c1"]
assert c1["group2"] is False and _blocks_by_id(run0)["c2"]["group2"] is True

# c1（按是否成交分组，多指标）：成交组 = 首指标"总额"最大的组（张三+王五）
gi_deal = c1["series"][0]["values"].index(max(c1["series"][0]["values"]))
r = client.post(f"/api/reports/{drill_ids[0]}/drill", json={"block_id": "c1", "group_index": gi_deal})
d = r.json()
assert r.status_code == 200 and d["total"] == 2, r.text
assert {row["customer_name"] for row in d["rows"]} == {"张三", "王五"}
assert any(c["prop"] == "customer_name" for c in d["columns"]) and any(c["prop"] == "created_at" for c in d["columns"])

# 多指标图的系列序号不影响记录集（指标不是筛选维度）
r = client.post(f"/api/reports/{drill_ids[0]}/drill", json={"block_id": "c1", "group_index": gi_deal, "series_index": 1})
assert r.json()["total"] == 2, r.text

# 查看端筛选叠加并收窄：只含"张" → 成交组只剩张三
r = client.post(f"/api/reports/{drill_ids[0]}/drill", json={
    "block_id": "c1", "group_index": 0,
    "filters": {"logic": "AND", "rules": [{"field": "customer_name", "op": "contains", "value": "张"}]}})
assert r.status_code == 200 and r.json()["total"] == 1 and r.json()["rows"][0]["customer_name"] == "张三", r.text
# 未开放字段的查看端筛选被拒
r = client.post(f"/api/reports/{drill_ids[0]}/drill", json={
    "block_id": "c1", "group_index": gi_deal,
    "filters": {"logic": "AND", "rules": [{"field": "amount", "op": "gt", "value": 1}]}})
assert r.status_code == 400, r.text

# c2（按月 × 成交二级分组）：2026-10 桶 + 成交系列 → 张三；"（空）"桶 + 未成交系列 → 赵六。两引擎各自定位序号后结果一致
for rid in drill_ids:
    res = client.post(f"/api/reports/{rid}/run").json()
    c2r = _blocks_by_id(res)["c2"]
    gi_oct = c2r["labels"].index("2026-10")
    i_sep = c2r["labels"].index("2026-09")
    si_deal = next(i for i, s in enumerate(c2r["series"]) if s["values"][i_sep] == 1)
    si_nodeal = 1 - si_deal
    d = client.post(f"/api/reports/{rid}/drill",
                    json={"block_id": "c2", "group_index": gi_oct, "series_index": si_deal}).json()
    assert d["total"] == 1 and d["rows"][0]["customer_name"] == "张三", (rid, d)
    gi_null = c2r["labels"].index("（空）")
    d = client.post(f"/api/reports/{rid}/drill",
                    json={"block_id": "c2", "group_index": gi_null, "series_index": si_nodeal}).json()
    assert d["total"] == 1 and d["rows"][0]["customer_name"] == "赵六", (rid, d)

# 非法序号 / 非图表区块 / 不存在的区块
r = client.post(f"/api/reports/{drill_ids[0]}/drill", json={"block_id": "c1", "group_index": 99})
assert r.status_code == 400, r.text
r = client.post(f"/api/reports/{drill_ids[0]}/drill", json={"block_id": "t1", "group_index": 0})
assert r.status_code == 400, r.text
r = client.post(f"/api/reports/{drill_ids[0]}/drill", json={"block_id": "ghost", "group_index": 0})
assert r.status_code == 400, r.text

for rid_ in drill_ids:
    client.delete(f"/api/reports/{rid_}")
print("报表图表下钻（分组/系列/查看端筛选/非法参数）通过")

# 10.9 透视表区块：行×列交叉聚合 + 行列合计 + 下钻
# 校验：行/列字段不存在、行列同分组、ratio 不可用、top_n 越界、时间型维度需日期字段
for bad in (
    [{"id": "x", "type": "pivot", "row": {"kind": "field", "field": "ghost"},
      "col": {"kind": "field", "field": "is_deal"}, "agg": "count"}],
    [{"id": "x", "type": "pivot", "row": {"kind": "field", "field": "is_deal"},
      "col": {"kind": "field", "field": "is_deal"}, "agg": "count"}],
    [{"id": "x", "type": "pivot", "row": {"kind": "field", "field": "is_deal"},
      "col": {"kind": "month", "field": "next_follow_date"}, "agg": "ratio"}],
    [{"id": "x", "type": "pivot", "row": {"kind": "field", "field": "is_deal"},
      "col": {"kind": "month", "field": "next_follow_date"}, "agg": "sum", "field": "customer_name"}],
    [{"id": "x", "type": "pivot", "row": {"kind": "month", "field": "customer_name"},
      "col": {"kind": "field", "field": "is_deal"}, "agg": "count"}],
    [{"id": "x", "type": "pivot", "row": {"kind": "field", "field": "is_deal"},
      "col": {"kind": "field", "field": "customer_name"}, "agg": "count", "row_top_n": 0}],
    [{"id": "x", "type": "pivot", "row": {"kind": "field", "field": "is_deal"},
      "col": {"kind": "field", "field": "customer_name"}, "agg": "count", "col_top_n": 99}],
):
    r = client.post("/api/reports", json={"name": "bad", "table_id": tj, "range": {"mode": "today"}, "blocks": bad})
    assert r.status_code == 400, r.text

PV_BLOCKS = [
    {"id": "p1", "type": "pivot", "title": "成交×月份 金额",
     "row": {"kind": "field", "field": "is_deal"},
     "col": {"kind": "month", "field": "next_follow_date"},
     "agg": "sum", "field": "amount"},
    {"id": "p2", "type": "pivot", "title": "月份×成交 笔数（无合计）",
     "row": {"kind": "month", "field": "next_follow_date"},
     "col": {"kind": "field", "field": "is_deal"},
     "agg": "count", "totals": False},
    {"id": "p3", "type": "pivot", "title": "成交 top1 + 其他",
     "row": {"kind": "field", "field": "is_deal"}, "row_top_n": 1,
     "col": {"kind": "month", "field": "next_follow_date"},
     "agg": "sum", "field": "amount"},
    {"id": "c1", "type": "chart", "title": "成交对比", "chart_type": "bar",
     "group": {"kind": "field", "field": "is_deal"}, "agg": "count"},
]
pv_ids = []
for t in (tj, tp):
    r = client.post("/api/reports", json={
        "name": "透视表", "table_id": t, "range": {"mode": "today", "date_field": "created_at"},
        "blocks": PV_BLOCKS,
    })
    assert r.status_code == 200, r.text
    pv_ids.append(r.json()["id"])


def _pivot_map(b):
    """{行标签: {列标签: 值, "＃合计": 行合计}, "＃合计": {...}}，行列序不依赖实现，便于两引擎对比"""
    m = {str(rl): {str(cl): v for cl, v in zip(b["col_labels"], b["cells"][i])}
         for i, rl in enumerate(b["row_labels"])}
    if b.get("totals"):
        for i, rl in enumerate(b["row_labels"]):
            m[str(rl)]["＃合计"] = b["row_totals"][i]
        m["＃合计"] = {str(cl): v for cl, v in zip(b["col_labels"], b["col_totals"])}
        m["＃合计"]["＃合计"] = b["grand_total"]
    return m


runs_pv = [client.post(f"/api/reports/{rid}/run").json() for rid in pv_ids]
for res in runs_pv:
    p1 = _blocks_by_id(res)["p1"]
    # 时间型列维度确定性排序：升序 + 空值最后；行维度计数并列按 key 升序（False < True）
    assert p1["col_labels"] == ["2026-09", "2026-10", "（空）"], p1["col_labels"]
    assert p1["row_labels"] == ["False", "True"], p1["row_labels"]
    assert _pivot_map(p1) == {
        "True": {"2026-09": 20000.25, "2026-10": 15000.5, "（空）": 0, "＃合计": 35000.75},
        "False": {"2026-09": 0, "2026-10": 8000, "（空）": 500, "＃合计": 8500},
        "＃合计": {"2026-09": 20000.25, "2026-10": 23000.5, "（空）": 500, "＃合计": 43500.75},
    }, p1

    p2 = _blocks_by_id(res)["p2"]
    assert p2["totals"] is False and "row_totals" not in p2 and "grand_total" not in p2
    assert p2["row_labels"] == ["2026-09", "2026-10", "（空）"], p2["row_labels"]
    assert _pivot_map(p2) == {
        "2026-09": {"True": 1, "False": 0},
        "2026-10": {"True": 1, "False": 1},
        "（空）": {"True": 0, "False": 1},
    }, p2

    # p3：field 型行维度 top_n=1 折叠"其他"（并列次序确定性：False 在前，True 并入"其他"）
    p3 = _blocks_by_id(res)["p3"]
    assert p3["row_labels"] == ["False", "其他"], p3["row_labels"]
    assert p3["row_totals"] == [8500, 35000.75]
    assert p3["col_totals"] == [20000.25, 23000.5, 500] and p3["grand_total"] == 43500.75
    assert p3["cells"][1] == [20000.25, 15000.5, 0], p3  # "其他"行 = True 组的折叠
    for i in range(2):  # count/sum 语义下单元格之和 = 行合计
        assert sum(p3["cells"][i]) == p3["row_totals"][i], p3

# 两引擎透视表结果一致
for bid in ("p1", "p2", "p3"):
    a, b = _blocks_by_id(runs_pv[0])[bid], _blocks_by_id(runs_pv[1])[bid]
    assert a["row_labels"] == b["row_labels"] and a["col_labels"] == b["col_labels"], bid
    assert _pivot_map(a) == _pivot_map(b), bid

# 透视表下钻：数据格 / 行合计 / 列合计 / 总计，两引擎一致
for rid in pv_ids:
    res = client.post(f"/api/reports/{rid}/run").json()
    p1 = _blocks_by_id(res)["p1"]
    ri_deal = p1["row_labels"].index("True")
    ci_oct = p1["col_labels"].index("2026-10")
    ci_sep = p1["col_labels"].index("2026-09")
    # 数据格：成交 × 2026-10 → 张三
    d = client.post(f"/api/reports/{rid}/drill",
                    json={"block_id": "p1", "group_index": ri_deal, "series_index": ci_oct}).json()
    assert d["total"] == 1 and d["rows"][0]["customer_name"] == "张三", (rid, d)
    # 行合计列（series_index=None）：成交行全部 → 张三、王五
    d = client.post(f"/api/reports/{rid}/drill", json={"block_id": "p1", "group_index": ri_deal}).json()
    assert d["total"] == 2 and {r_["customer_name"] for r_ in d["rows"]} == {"张三", "王五"}, (rid, d)
    # 列合计行（group_index=None）：2026-09 列全部 → 王五
    d = client.post(f"/api/reports/{rid}/drill", json={"block_id": "p1", "series_index": ci_sep}).json()
    assert d["total"] == 1 and d["rows"][0]["customer_name"] == "王五", (rid, d)
    # 总计格（两个下标都 None）→ 全部 4 条
    d = client.post(f"/api/reports/{rid}/drill", json={"block_id": "p1"}).json()
    assert d["total"] == 4, (rid, d)
    # 非法行/列序号
    assert client.post(f"/api/reports/{rid}/drill", json={"block_id": "p1", "group_index": 99}).status_code == 400
    assert client.post(f"/api/reports/{rid}/drill", json={"block_id": "p1", "series_index": 99}).status_code == 400

# chart 区块下钻仍要求 group_index（为 None 时 400）
r = client.post(f"/api/reports/{pv_ids[0]}/drill", json={"block_id": "c1"})
assert r.status_code == 400, r.text

# 导出兼容透视表：xlsx/html 均正常且含矩阵与合计
r = client.get(f"/api/reports/{pv_ids[0]}/export?format=xlsx")
assert r.status_code == 200, r.text
r = client.get(f"/api/reports/{pv_ids[0]}/export?format=html")
assert r.status_code == 200 and "成交×月份 金额" in r.text and "合计" in r.text, r.status_code

for rid_ in pv_ids:
    client.delete(f"/api/reports/{rid_}")
print("报表透视表区块（交叉聚合/行列合计/其他桶/下钻/导出）通过")

# 10.10 AI 助手：对话生成动作预览 + 执行落库（mock provider 顺序应答）
class AssistantFakeProvider:
    responses = []

    def complete(self, prompt, system=None):
        assert AssistantFakeProvider.responses, "没有预设的应答"
        return AssistantFakeProvider.responses.pop(0)


gw.get_default_provider = lambda db: AssistantFakeProvider()

# 纯聊天：action 为 null
AssistantFakeProvider.responses = ['{"reply": "你好，有什么可以帮你？", "action": null}']
r = client.post("/api/assistant/chat", json={"message": "你好"})
assert r.status_code == 200 and r.json()["reply"] and r.json()["action_card"] is None, r.text

# 智能填表：未知字段丢弃 + 坏值置空 + 有效值转换
AssistantFakeProvider.responses = [
    '{"reply": "整理出 1 条记录", "action": {"type": "fill_records", "table_id": %d, "records": ['
    '{"customer_name": "钱七", "amount": "3万", "is_deal": true, "ghost_field": "x", "next_follow_date": "2026-11-01"}]}}' % tj
]
r = client.post("/api/assistant/chat", json={"message": "记一笔：钱七，金额3万，已成交", "context": {"table_id": tj}})
card = r.json()["action_card"]
assert r.status_code == 200 and card["type"] == "fill_records" and card["table_id"] == tj, r.text
rec = card["payload"]["records"][0]
assert rec["customer_name"] == "钱七" and rec["is_deal"] is True and rec["next_follow_date"] == "2026-11-01", rec
assert rec["amount"] is None and "ghost_field" not in rec, rec  # "3万"无法转 decimal 置空；未知字段丢弃
assert any("ghost_field" in w for w in card["warnings"]) and any("amount" in w or "意向金额" in w for w in card["warnings"]), card

# 执行填表：落库并可查
r = client.post("/api/assistant/execute", json={
    "type": "fill_records",
    "payload": {"table_id": tj, "records": [{"customer_name": "钱七", "amount": 30000, "is_deal": True}]},
})
res = r.json()
assert r.status_code == 200 and res["ok"] == 1 and not res["fail"], r.text
items = client.get(f"/api/dyn/{tj}/records", params={"filters": '[{"field":"customer_name","op":"eq","value":"钱七"}]'}).json()
assert items["total"] == 1 and items["items"][0]["amount"] == 30000, items
# 执行入口重新鉴权：不存在的表 404；空记录 400；类型非法 400
assert client.post("/api/assistant/execute", json={"type": "fill_records", "payload": {"table_id": 99999, "records": [{"a": 1}]}}).status_code == 404
assert client.post("/api/assistant/execute", json={"type": "fill_records", "payload": {"table_id": tj, "records": []}}).status_code == 400
assert client.post("/api/assistant/execute", json={"type": "hack", "payload": {}}).status_code == 400
client.delete(f"/api/dyn/{tj}/records/{res['record_ids'][0]}")

# AI 建表：非法类型降级 + 执行落库
AssistantFakeProvider.responses = [
    '{"reply": "设计了 2 个字段", "action": {"type": "create_table", "label": "供应商台账", "fields": ['
    '{"field_name": "supplier_name", "label": "供应商", "data_type": "varchar", "nullable": false},'
    '{"field_name": "level", "label": "等级", "data_type": "magic", "widget": "select", "options": {"options": ["A", "B"]}}]}}'
]
r = client.post("/api/assistant/chat", json={"message": "帮我建个供应商台账"})
card = r.json()["action_card"]
assert card["type"] == "create_table" and len(card["payload"]["fields"]) == 2, r.text
assert card["payload"]["fields"][1]["data_type"] == "varchar", card  # 非法类型降级
r = client.post("/api/assistant/execute", json={"type": "create_table", "payload": card["payload"]})
new_tid = r.json()["table_id"]
assert r.status_code == 200 and new_tid, r.text
t_meta = client.get(f"/api/tables/{new_tid}").json()
assert t_meta["label"] == "供应商台账" and len(t_meta["fields"]) == 2, t_meta

# 数据问答：stat 受控聚合（全部时间 count = 4 条 PRECS）
AssistantFakeProvider.responses = [
    '{"reply": "查一下", "action": {"type": "query", "table_id": %d, "spec": {"kind": "stat", "agg": "sum", "field": "amount"}}}' % tj
]
r = client.post("/api/assistant/chat", json={"message": "意向金额一共多少", "context": {"table_id": tj}})
card = r.json()["action_card"]
assert card["type"] == "query_answer" and card["result"]["value"] == 43500.75, r.text
assert card["result"]["range_label"] == "全部时间", card["result"]
# 问答带筛选与时间口径（本月按 created_at：4 条都是今天创建）
AssistantFakeProvider.responses = [
    '{"reply": "查一下", "action": {"type": "query", "table_id": %d, "spec": {"kind": "stat", "agg": "count",'
    ' "filters": {"rules": [{"field": "is_deal", "op": "eq", "value": true}]}, "range": {"mode": "this_month"}}}}' % tj
]
r = client.post("/api/assistant/chat", json={"message": "本月成交几笔", "context": {"table_id": tj}})
assert r.json()["action_card"]["result"]["value"] == 2, r.text
# 问答 spec 非法（sum 文本字段）→ 卡片被丢弃
AssistantFakeProvider.responses = [
    '{"reply": "试试", "action": {"type": "query", "table_id": %d, "spec": {"kind": "stat", "agg": "sum", "field": "customer_name"}}}' % tj
]
r = client.post("/api/assistant/chat", json={"message": "客户名求和", "context": {"table_id": tj}})
assert r.json()["action_card"] is None, r.text

# 生成 Excel：blank 模板 + export 导出，下载均 200
AssistantFakeProvider.responses = [
    '{"reply": "模板好了", "action": {"type": "gen_excel", "mode": "blank", "label": "客户登记模板", "fields": ['
    '{"field_name": "customer_name", "label": "客户姓名", "data_type": "varchar"}], "sample_rows": [["张三"]]}}'
]
r = client.post("/api/assistant/chat", json={"message": "给我一个客户登记 Excel 模板"})
card = r.json()["action_card"]
assert card["type"] == "gen_excel" and card["payload"]["mode"] == "blank", r.text
r = client.post("/api/assistant/execute", json={"type": "gen_excel", "payload": card["payload"]})
fid = r.json()["file_id"]
assert r.status_code == 200 and fid, r.text
r = client.get(f"/api/assistant/download/{fid}")
assert r.status_code == 200 and r.content[:2] == b"PK", r.status_code

AssistantFakeProvider.responses = [
    '{"reply": "导好了", "action": {"type": "gen_excel", "mode": "export", "table_id": %d,'
    ' "filters": {"rules": [{"field": "is_deal", "op": "eq", "value": true}]}}}' % tj
]
r = client.post("/api/assistant/chat", json={"message": "把成交客户导成 Excel", "context": {"table_id": tj}})
card = r.json()["action_card"]
assert card["type"] == "gen_excel" and card["payload"]["mode"] == "export", r.text
r = client.post("/api/assistant/execute", json={"type": "gen_excel", "payload": card["payload"]})
fid = r.json()["file_id"]
r = client.get(f"/api/assistant/download/{fid}")
assert r.status_code == 200 and r.content[:2] == b"PK", r.status_code

# 创建报表：chat 意图 → assist_report 生成配置 → 执行落库可直接运行
AssistantFakeProvider.responses = [
    '{"reply": "周报设计好了", "action": {"type": "create_report", "table_id": %d, "description": "本周客户跟进周报，含新增客户数"}}' % tj,
    '{"name": "客户跟进周报", "range": {"mode": "this_week"}, "blocks": [{"type": "stat", "title": "新增客户数", "agg": "count"}], "notes": ""}',
]
r = client.post("/api/assistant/chat", json={"message": "帮我创建客户跟进表的周报", "context": {"table_id": tj}})
card = r.json()["action_card"]
assert card["type"] == "create_report" and card["payload"]["name"] == "客户跟进周报", r.text
r = client.post("/api/assistant/execute", json={"type": "create_report", "payload": card["payload"]})
rep_id = r.json()["report_id"]
assert r.status_code == 200 and rep_id, r.text
r = client.post(f"/api/reports/{rep_id}/run")
assert r.status_code == 200 and r.json()["blocks"][0]["value"] == 4, r.text  # 本周 4 条 PRECS
client.delete(f"/api/reports/{rep_id}")

# 创建任务规则：chat 意图 → assist_task 生成配置 → 执行落库默认停用
AssistantFakeProvider.responses = [
    '{"reply": "任务设计好了", "action": {"type": "create_task", "table_id": %d, "description": "每天早上提醒已成交客户"}}' % tj,
    '{"name": "成交客户提醒", "condition_mode": "structured",'
    ' "condition": {"logic": "AND", "rules": [{"field": "is_deal", "op": "eq", "value": true}]},'
    ' "schedule": {"type": "cron", "expr": "0 9 * * *"},'
    ' "action": {"type": "notify", "template": "成交客户 {customer_name}", "recipients": {"type": "fixed", "value": ""}},'
    ' "cooldown_hours": 24, "notes": ""}',
]
r = client.post("/api/assistant/chat", json={"message": "帮我创建成交提醒任务", "context": {"table_id": tj}})
card = r.json()["action_card"]
assert card["type"] == "create_task", r.text
r = client.post("/api/assistant/execute", json={"type": "create_task", "payload": card["payload"]})
task_id = r.json()["task_id"]
assert r.status_code == 200 and task_id, r.text
t = next(x for x in client.get("/api/tasks").json() if x["id"] == task_id)
assert t["enabled"] is False and t["condition"]["rules"][0]["field"] == "is_deal", t
client.delete(f"/api/tasks/{task_id}")

# 恢复后续测试使用的 LLM mock（智能判断条件）
gw.get_default_provider = lambda db: JudgeFakeProvider()

client.delete(f"/api/tables/{new_tid}")
print("AI 助手（聊天/填表/建表/问答/生成Excel/权限校验）通过")

# 10.10.5 助手粘贴图片：走视觉模型（recognize），无文字也可发送；非法图片被拒
import base64


class VisionChatFakeProvider(AssistantFakeProvider):
    def recognize(self, images, prompt, system=None):
        assert images and images[0][:2] == b"\xff\xd8", "图片应以 JPEG 字节传入"
        assert "附上了 1 张图片" in prompt, prompt[-200:]
        return AssistantFakeProvider.responses.pop(0)


gw.get_default_provider = lambda db: VisionChatFakeProvider()
AssistantFakeProvider.responses = [
    '{"reply": "从图片中整理出 1 条记录", "action": {"type": "fill_records", "table_id": %d,'
    ' "records": [{"customer_name": "图客户", "amount": 100}]}}' % tj
]
img_data_url = "data:image/jpeg;base64," + base64.b64encode(b"\xff\xd8\xff\xd9").decode()
r = client.post("/api/assistant/chat", json={"message": "", "images": [img_data_url]})
card = r.json()["action_card"]
assert r.status_code == 200 and card["type"] == "fill_records", r.text
assert card["payload"]["records"][0]["customer_name"] == "图客户", card
# 无文字且无图片 → 400；非法图片格式 → 400
assert client.post("/api/assistant/chat", json={"message": ""}).status_code == 400
assert client.post("/api/assistant/chat", json={"message": "x", "images": ["not-a-data-url"]}).status_code == 400
gw.get_default_provider = lambda db: JudgeFakeProvider()
print("AI 助手图片对话（视觉模型/空消息/非法图片）通过")

# 10.11 AI 助手联网搜索：搜索执行 → 二次调用组织回答 + 来源卡片；搜索不可用降级
import app.services.web_search as ws_mod

gw.get_default_provider = lambda db: AssistantFakeProvider()

# 第一次调用输出 web_search 动作；第二次调用是基于搜索结果的纯文本回答
AssistantFakeProvider.responses = [
    '{"reply": "", "action": {"type": "web_search", "queries": ["珠海市人民医院 电话", "珠海市人民医院 骨科主任"]}}',
    "珠海市人民医院总机电话是 0756-2222569（来源：医院官网），骨科主任为蒋煜文。",
]
_orig_search = ws_mod.web_search
ws_mod.web_search = lambda db, q, count=8: [
    {"title": f"结果-{q}", "url": f"https://example.com/{q}", "snippet": "摘要"},
]
try:
    r = client.post("/api/assistant/chat", json={"message": "珠海市人民医院电话和骨科主任"})
finally:
    ws_mod.web_search = _orig_search
out = r.json()
assert r.status_code == 200 and out["reply"].startswith("珠海市人民医院总机电话"), r.text
card = out["action_card"]
assert card["type"] == "search_answer" and len(card["payload"]["results"]) == 2, card
assert card["payload"]["queries"] == ["珠海市人民医院 电话", "珠海市人民医院 骨科主任"], card

# 搜索服务未配置/失败 → 如实降级，不崩溃
ws_mod.web_search = lambda db, q, count=8: (_ for _ in ()).throw(ws_mod.SearchError("尚未配置联网搜索服务"))
AssistantFakeProvider.responses = [
    '{"reply": "", "action": {"type": "web_search", "queries": ["今天天气"]}}',
]
try:
    r = client.post("/api/assistant/chat", json={"message": "今天天气怎么样"})
finally:
    ws_mod.web_search = _orig_search
out = r.json()
assert r.status_code == 200 and "尚未配置联网搜索服务" in out["reply"] and out["action_card"] is None, r.text

gw.get_default_provider = lambda db: JudgeFakeProvider()
print("AI 助手联网搜索（二次调用/来源卡片/失败降级）通过")

# 11. 权限矩阵：未分享 404 / 分享者按开关 / 主人与 admin 全权
admin_headers = dict(client.headers)
r = client.post("/api/auth/register", json={"username": "worker", "password": "secret123"})
worker_headers = {"Authorization": f"Bearer {r.json()['token']}"}
r = client.post("/api/auth/register", json={"username": "outsider", "password": "secret123"})
outsider_headers = {"Authorization": f"Bearer {r.json()['token']}"}


def as_user(h):
    client.headers = h


as_user(worker_headers)
assert client.get(f"/api/tables/{tj}").status_code == 404
assert client.get(f"/api/dyn/{tj}/records").status_code == 404
as_user(admin_headers)
r = client.post(f"/api/tables/{tj}/shares", json={"username": "worker", "can_view": True})
assert r.status_code == 200, r.text
assert r.json()["status"] == "pending"  # 直发分享需对方确认
as_user(worker_headers)
assert client.get(f"/api/tables/{tj}").status_code == 404  # 待确认不可见
assert all(t["id"] != tj for t in client.get("/api/tables").json())
r = client.get("/api/shares/pending")
share_id = next(s["id"] for s in r.json() if s["table_id"] == tj)
assert client.post(f"/api/shares/{share_id}/accept").status_code == 200
r = client.get(f"/api/tables/{tj}")
assert r.status_code == 200 and r.json()["my_perms"]["can_create"] is False, r.text
assert client.get(f"/api/dyn/{tj}/records").json()["total"] == 4
assert client.post(f"/api/dyn/{tj}/records", json={"customer_name": "x"}).status_code == 404  # 无新增权
assert client.put(f"/api/tables/{tj}", json={"label": "x"}).status_code == 404  # 非主人不能改结构
assert client.delete(f"/api/tables/{tj}").status_code == 404  # 非主人不能删表
assert client.post(f"/api/tables/{tj}/shares", json={"username": "outsider"}).status_code == 404  # 非主人不能分享
assert any(t["id"] == tj and t["owner_label"] == "admin" for t in client.get("/api/tables").json())  # 分享表可见
as_user(admin_headers)
r = client.post(f"/api/tables/{tj}/shares", json={"username": "worker", "can_view": True, "can_create": True})
assert r.status_code == 200 and r.json()["status"] == "accepted"  # 已接受的分享仅更新权限，不回到待确认
as_user(worker_headers)
assert client.post(f"/api/dyn/{tj}/records", json={"customer_name": "worker新增", "amount": 1}).status_code == 200
as_user(outsider_headers)
assert client.get(f"/api/tables/{tj}").status_code == 404  # 未分享不可见
assert all(t["table_id"] != tj for t in client.get("/api/tasks").json())
print("权限矩阵通过")

# 12. VIP 门槛与角色管理
as_user(worker_headers)
r = client.post("/api/tables", json={"label": "vip表", "storage_mode": "physical", "fields": PFIELDS})
assert r.status_code == 403  # user 不能建独立表
r = client.post("/api/tables", json={"label": "worker的表", "fields": PFIELDS})
worker_tid = r.json()["table"]["id"]
assert client.post(f"/api/tables/{worker_tid}/shares", json={"username": "outsider"}).status_code == 403  # user 不能分享
as_user(admin_headers)
assert all(t["id"] != worker_tid for t in client.get("/api/tables").json())  # admin 列表也不含未分享给自己的表
r = client.get("/api/users")
worker_id = next(u["id"] for u in r.json() if u["username"] == "worker")
assert client.put(f"/api/users/{worker_id}/role", json={"role": "vip"}).status_code == 200
me_id = next(u["id"] for u in client.get("/api/users").json() if u["role"] == "admin")
assert client.put(f"/api/users/{me_id}/role", json={"role": "user"}).status_code == 400  # 不能改自己
as_user(worker_headers)
# 分享对象约束： outsider 还不是好友/同组 → 403；建立好友关系后可分享
assert client.post(f"/api/tables/{worker_tid}/shares", json={"username": "outsider", "can_view": True}).status_code == 403
r = client.post("/api/friends/request", json={"username": "outsider"})
assert r.status_code == 200 and r.json()["status"] == "pending", r.text
assert client.post("/api/friends/request", json={"username": "outsider"}).status_code == 400  # 重复申请
as_user(outsider_headers)
r = client.get("/api/friends/requests")
fid = next(f["id"] for f in r.json() if f["username"] == "worker")
assert client.post(f"/api/friends/{fid}/accept").status_code == 200
as_user(worker_headers)
assert any(f["username"] == "outsider" for f in client.get("/api/friends").json())
assert client.post(f"/api/tables/{worker_tid}/shares", json={"username": "outsider", "can_view": True}).status_code == 200  # vip 可分享（待确认）
r = client.post("/api/tables", json={"label": "vip物理表", "storage_mode": "physical", "fields": PFIELDS})
assert r.status_code == 200  # vip 可建独立表
vip_tid = r.json()["table"]["id"]
print("VIP 门槛通过")

# 13. 归属隔离：规则/报表/通知按 user_id 隔离
r = client.post("/api/dyn/{}/records".format(worker_tid), json={"customer_name": "worker客户", "amount": 100, "is_deal": False})
assert r.status_code == 200, r.text
r = client.post("/api/tasks", json={
    "name": "worker规则", "table_id": worker_tid, "enabled": False,
    "condition_mode": "structured",
    "condition": {"logic": "AND", "rules": [{"field": "customer_name", "op": "contains", "value": "worker"}]},
    "schedule": {"type": "interval", "minutes": 60},
    "action": {"type": "notify", "template": "客户【{customer_name}】", "recipients": {"type": "fixed", "value": ""}},
})
assert r.status_code == 200, r.text
worker_rule = r.json()["id"]
assert any(t["id"] == worker_rule for t in client.get("/api/tasks").json())
as_user(outsider_headers)
assert all(t["id"] != worker_rule for t in client.get("/api/tasks").json())
assert all(t["id"] != rep_ids[0] for t in client.get("/api/reports").json())
as_user(outsider_headers)
assert client.post(f"/api/tasks/{worker_rule}/run").status_code == 404  # 他人规则不可操作
as_user(outsider_headers)
client.post("/api/notify/read", json={"all": True})  # 清掉好友申请通知，避免干扰下面的隔离断言
as_user(worker_headers)
client.post("/api/notify/read", json={"all": True})
r = client.post(f"/api/tasks/{worker_rule}/run")
assert r.status_code == 200 and r.json()["fired"] == 1, r.text
assert client.get("/api/notify/unread_count").json()["count"] >= 1  # worker 收到自己的通知
as_user(outsider_headers)
assert client.get("/api/notify/unread_count").json()["count"] == 0  # 别人收不到
as_user(admin_headers)
assert client.get("/api/notify/unread_count").status_code == 200  # admin 也只看到自己的通知（不报错）
print("归属隔离通过")

# 14. 迁移幂等：重复执行不报错、列不重复加
from app.database import engine as _engine
from app.services.migrate import run_migrations
run_migrations(_engine)
run_migrations(_engine)
print("迁移幂等通过")

# 15. RBAC：自定义权限/角色/授权，数值型权限（数据表上限）生效
role_codes = {x["code"] for x in client.get("/api/roles").json()}
assert {"admin", "vip", "user"} <= role_codes
perm_ids = {p["code"]: p["id"] for p in client.get("/api/permissions").json()}
assert {"share", "create_physical_table", "max_tables"} <= set(perm_ids)

r = client.post("/api/permissions", json={"code": "custom_perm", "name": "自定义权限"})
assert r.status_code == 200, r.text
custom_pid = r.json()["id"]
r = client.post("/api/roles", json={"code": "limited", "name": "受限角色"})
assert r.status_code == 200, r.text
limited_id = r.json()["id"]
# 授权：share + max_tables=1（不给 create_physical_table）
r = client.put(f"/api/roles/{limited_id}/permissions", json={"grants": [
    {"permission_id": perm_ids["share"]},
    {"permission_id": perm_ids["max_tables"], "value": 1},
]})
assert r.status_code == 200, r.text

r = client.post("/api/auth/register", json={"username": "limited_user", "password": "secret123"})
lu_headers = {"Authorization": f"Bearer {r.json()['token']}"}
lu_id = r.json()["user"]["id"]
assert client.put(f"/api/users/{lu_id}/role", json={"role": "limited"}).status_code == 200
assert client.put(f"/api/users/{lu_id}/role", json={"role": "ghost"}).status_code == 400  # 角色不存在

client.headers = lu_headers
r = client.post("/api/tables", json={"label": "限额表1", "fields": PFIELDS})
assert r.status_code == 200, r.text
lu_tid = r.json()["table"]["id"]
r = client.post("/api/tables", json={"label": "限额表2", "fields": PFIELDS})
assert r.status_code == 403, r.text  # 上限 1 张
# 分享对象约束：先与 outsider 成为好友
r = client.post("/api/friends/request", json={"username": "outsider"})
lu_fid = r.json()["id"]
saved_h = dict(client.headers)
client.headers = outsider_headers
assert client.post(f"/api/friends/{lu_fid}/accept").status_code == 200
client.headers = saved_h
assert client.post(f"/api/tables/{lu_tid}/shares", json={"username": "outsider", "can_view": True}).status_code == 200  # 有 share
r = client.post("/api/tables", json={"label": "限额物理表", "storage_mode": "physical", "fields": PFIELDS})
assert r.status_code == 403, r.text  # 无 create_physical_table
client.headers = admin_headers
assert client.delete(f"/api/roles/{limited_id}").status_code == 400  # 角色使用中
assert client.delete(f"/api/permissions/{perm_ids['share']}").status_code == 400  # 内置权限
client.headers = lu_headers
client.delete(f"/api/tables/{lu_tid}")  # 清掉唯一一张表
client.headers = admin_headers
assert client.delete(f"/api/roles/{limited_id}").status_code == 400  # 用户仍挂着该角色
assert client.put(f"/api/users/{lu_id}/role", json={"role": "user"}).status_code == 200  # 改回默认角色
assert client.delete(f"/api/roles/{limited_id}").status_code == 200
assert client.delete(f"/api/permissions/{custom_pid}").status_code == 200

# 给内置 user 角色授予 share：普通用户重新登录后即可分享（登录信息携带权限列表）
r = client.get("/api/roles")
user_role = next(x for x in r.json() if x["code"] == "user")
grants = [{"permission_id": p["id"], "value": p["value"]} for p in user_role["permissions"]]
grants.append({"permission_id": perm_ids["share"]})
r = client.put(f"/api/roles/{user_role['id']}/permissions", json={"grants": grants})
assert r.status_code == 200, r.text
r = client.post("/api/auth/register", json={"username": "normie", "password": "secret123"})
normie_headers = {"Authorization": f"Bearer {r.json()['token']}"}
assert "share" in r.json()["user"]["perms"]  # 登录信息携带权限列表
client.headers = normie_headers
r = client.post("/api/friends/request", json={"username": "outsider"})
normie_fid = r.json()["id"]
client.headers = outsider_headers
assert client.post(f"/api/friends/{normie_fid}/accept").status_code == 200
client.headers = normie_headers
r = client.post("/api/tables", json={"label": "normie的表", "fields": PFIELDS})
normie_tid = r.json()["table"]["id"]
assert client.post(f"/api/tables/{normie_tid}/shares", json={"username": "outsider", "can_view": True}).status_code == 200  # 普通用户被授予 share 后可分享
client.headers = admin_headers
print("RBAC 通过")

# 15.5 分享确认流与组约束
r = client.post("/api/auth/register", json={"username": "stranger", "password": "secret123"})
stranger_headers = {"Authorization": f"Bearer {r.json()['token']}"}

# 用户搜索 scope：all 能搜到任何人；shareable 只搜得到好友/同组
as_user(stranger_headers)
r = client.get("/api/users/search", params={"q": "worker", "scope": "all"})
assert any(u["username"] == "worker" for u in r.json())
r = client.get("/api/users/search", params={"q": "worker", "scope": "shareable"})
assert r.json() == []  # 非好友非同组搜不到

# 拒绝 → 分享者看到已拒绝 → 重新发起重置 pending → 接受后可见
as_user(admin_headers)
r = client.post(f"/api/tables/{vip_tid}/shares", json={"username": "stranger", "can_view": True})
assert r.status_code == 200 and r.json()["status"] == "pending", r.text
as_user(stranger_headers)
sid = next(s["id"] for s in client.get("/api/shares/pending").json() if s["table_id"] == vip_tid)
assert client.post(f"/api/shares/{sid}/reject").status_code == 200
assert client.get(f"/api/tables/{vip_tid}").status_code == 404
as_user(admin_headers)
shares_ = client.get(f"/api/tables/{vip_tid}/shares").json()
assert next(s for s in shares_ if s["target"] == "stranger")["status"] == "rejected"
r = client.post(f"/api/tables/{vip_tid}/shares", json={"username": "stranger", "can_view": True})
assert r.status_code == 200 and r.json()["status"] == "pending"  # 重新发起
as_user(stranger_headers)
sid = next(s["id"] for s in client.get("/api/shares/pending").json() if s["table_id"] == vip_tid)
assert client.post(f"/api/shares/{sid}/accept").status_code == 200
assert client.get(f"/api/tables/{vip_tid}").status_code == 200

# 组约束：只能分享到自己所在的组（admin 豁免）；组分享即时生效
as_user(admin_headers)
r = client.post("/api/groups", json={"name": "临时组"})
tmp_gid = r.json()["id"]
assert client.post(f"/api/groups/{tmp_gid}/members", json={"username": "stranger"}).status_code == 200
as_user(worker_headers)
assert client.post(f"/api/tables/{worker_tid}/shares", json={"group_id": tmp_gid, "can_view": True}).status_code == 403  # 非本组
as_user(admin_headers)
assert client.post(f"/api/groups/{tmp_gid}/members", json={"username": "worker"}).status_code == 200
as_user(worker_headers)
r = client.post(f"/api/tables/{worker_tid}/shares", json={"group_id": tmp_gid, "can_view": True})
assert r.status_code == 200 and r.json()["status"] == "accepted", r.text
as_user(stranger_headers)
assert client.get(f"/api/tables/{worker_tid}").status_code == 200  # 组内即时可见
r = client.get("/api/users/search", params={"q": "worker", "scope": "shareable"})
assert any(u["username"] == "worker" for u in r.json())  # 同组即可搜到
assert any(g["id"] == tmp_gid for g in client.get("/api/groups/mine").json())
as_user(admin_headers)
# admin 分享时可见所有用户和用户组（约束对 admin 豁免）
r = client.get("/api/users/search", params={"q": "stranger", "scope": "shareable"})
assert any(u["username"] == "stranger" for u in r.json())
assert any(g["id"] == tmp_gid for g in client.get("/api/groups/mine").json())
client.delete(f"/api/groups/{tmp_gid}")  # 删组 → 组分享同步失效
as_user(stranger_headers)
assert client.get(f"/api/tables/{worker_tid}").status_code == 404

# 互申自动接受 + 删除好友
as_user(stranger_headers)
r = client.post("/api/friends/request", json={"username": "outsider"})
assert r.status_code == 200 and r.json()["status"] == "pending"
as_user(outsider_headers)
r = client.post("/api/friends/request", json={"username": "stranger"})
assert r.status_code == 200 and r.json()["status"] == "accepted"  # 互申自动成为好友
fid = next(f["id"] for f in client.get("/api/friends").json() if f["username"] == "stranger")
assert client.delete(f"/api/friends/{fid}").status_code == 200
r = client.post("/api/friends/request", json={"username": "stranger"})  # 删除后可重新申请
assert r.status_code == 200 and r.json()["status"] == "pending"
as_user(admin_headers)
print("分享确认流与组约束通过")

# 16. 链接分享 + 用户组
# 用户组
r = client.post("/api/groups", json={"name": "生产部", "description": "生产组成员"})
assert r.status_code == 200, r.text
gid = r.json()["id"]
assert client.post(f"/api/groups/{gid}/members", json={"username": "worker"}).status_code == 200
assert client.post(f"/api/groups/{gid}/members", json={"username": "worker"}).status_code == 400  # 重复入组

# 分享给组
tj2 = make_table("组分享表", "json")
fill(tj2)
r = client.post(f"/api/tables/{tj2}/shares", json={"group_id": gid, "can_view": True, "can_create": True})
assert r.status_code == 200 and r.json()["target_type"] == "group", r.text
as_user(worker_headers)  # worker 是组成员：即时生效 + 收到告知通知
assert any("生产部" in n["content"] for n in client.get("/api/notify").json())
as_user(admin_headers)
as_user(outsider_headers)
assert client.get(f"/api/tables/{tj2}").status_code == 404  # outsider 不在组
as_user(admin_headers)
assert client.post(f"/api/groups/{gid}/members", json={"username": "outsider"}).status_code == 200
as_user(outsider_headers)
assert client.get(f"/api/tables/{tj2}").status_code == 200  # 入组后可见
r = client.post(f"/api/dyn/{tj2}/records", json={"customer_name": "组新增", "amount": 1})
assert r.status_code == 200, r.text  # 组分享的 can_create 生效
rid2 = r.json()["id"]
assert client.delete(f"/api/dyn/{tj2}/records/{rid2}").status_code == 404  # 组分享未给 delete
as_user(admin_headers)

# 链接分享（表，无密码）
r = client.post(f"/api/tables/{tj2}/share-links", json={})
assert r.status_code == 200, r.text
token = r.json()["token"]
saved_h = dict(client.headers)
client.headers = {}
r = client.get(f"/api/share/{token}")
assert r.status_code == 200 and r.json()["total"] >= 4 and r.json()["fields"], r.text
# 带密码 + 有效期
client.headers = saved_h
r = client.post(f"/api/tables/{tj2}/share-links", json={"password": "abcd", "expires_in_days": 7})
token2, lid2 = r.json()["token"], r.json()["id"]
client.headers = {}
assert client.get(f"/api/share/{token2}").status_code == 401  # 需要密码
assert client.get(f"/api/share/{token2}", params={"password": "wrong"}).status_code == 401
assert client.get(f"/api/share/{token2}", params={"password": "abcd"}).status_code == 200
client.headers = saved_h
client.delete(f"/api/tables/{tj2}/share-links/{lid2}")  # 撤销
client.headers = {}
assert client.get(f"/api/share/{token2}", params={"password": "abcd"}).status_code == 404
client.headers = saved_h

# 报表链接分享
r = client.post("/api/reports", json={
    "name": "分享报表", "table_id": tj2, "enabled": False,
    "range": {"mode": "this_month", "date_field": "created_at"},
    "blocks": [{"id": "b1", "type": "stat", "title": "总额", "agg": "sum", "field": "amount", "filters": {"logic": "AND", "rules": []}}],
    "schedule": {}, "push": {},
})
tpl_id2 = r.json()["id"]
r = client.post(f"/api/reports/{tpl_id2}/share-links", json={})
rtoken = r.json()["token"]
client.headers = {}
r = client.get(f"/api/share/{rtoken}")
assert r.status_code == 200 and r.json()["resource_type"] == "report" and r.json()["report"]["blocks"][0]["value"] > 0, r.text
client.headers = saved_h

# 删除组 → 组分享失效
client.delete(f"/api/groups/{gid}")
as_user(outsider_headers)
assert client.get(f"/api/tables/{tj2}").status_code == 404
as_user(admin_headers)
print("链接分享 + 用户组通过")

# 清理测试表
as_user(admin_headers)
for t_ in (tj, tp, worker_tid, vip_tid, tj2, normie_tid):
    client.delete(f"/api/tables/{t_}")

print("\nALL SMOKE TESTS PASSED")
cm.__exit__(None, None, None)
