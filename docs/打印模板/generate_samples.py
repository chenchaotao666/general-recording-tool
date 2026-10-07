# -*- coding: utf-8 -*-
"""按 docs/工单打印截图 逐张复刻 Excel 打印模板（脱敏版）→ docs/打印模板/<行业>/。

脱敏约定：
  公司名 → XXX公司；地址/电话 → 地址：XXX / 电话：XXX；人名（如 系统管理员）→ 删除；
  示例数据（产品/数量/金额/客户名/单号/日期）→ 全部替换为占位符或留空；
  网站水印（一彩送货单打印软件 www.yc620.com）→ 不保留。

占位符语法（与系统 Excel 模板引擎一致）：
  主表  {customer} {address} {phone} {contact} {order_no} {delivery_date} {order_date}
        {contract_no} {vehicle_no} {pay_method} {ship_method} {handler} {remark} {_today}
  明细  {items._index} {items.order_no} {items.product_no} {items.name} {items.spec} {items.model}
        {items.color} {items.material} {items.unit} {items.qty} {items.price} {items.amount}
        {items.weight} {items.spare} {items.remark} ...
  合计  {items.qty#sum} {items.amount#sum} ...    大写 {_total_cn}
注意：上传到数据表前，把占位符字段名改成表里的实际字段名。
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent / "打印模板"

# analyze_grids.py 的产物：从截图网格线实测的列宽像素比例（无覆盖的模板用启发式）
import json as _json

try:
    _PX_WIDTHS = _json.loads((ROOT / "widths_override.json").read_text(encoding="utf-8"))
except FileNotFoundError:
    _PX_WIDTHS = {}

THIN = Side(style="thin")
GRID = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DASH = Side(style="dashed")
RED = "C00000"
SONG = "宋体"

CN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RT = Alignment(horizontal="right", vertical="center")
VERT = Alignment(horizontal="center", vertical="center", text_rotation=255, wrap_text=True)

TOK = {
    "序号": "{items._index}", "序": "{items._index}",
    "订单号": "{items.order_no}", "订单号码": "{items.order_no}", "合同号码": "{items.order_no}",
    "合同号": "{items.order_no}", "订单编号": "{items.order_no}",
    "产品编号": "{items.product_no}", "货号": "{items.product_no}", "项目编号": "{items.product_no}",
    "编号": "{items.product_no}", "产品编码": "{items.product_no}", "物料编码": "{items.product_no}",
    "产品名称": "{items.name}", "项目名称": "{items.name}", "品名及规格": "{items.name}",
    "产品名称、规格": "{items.name}", "产品名称、型号": "{items.name}", "品名": "{items.name}",
    "货物名称": "{items.name}", "物料名称": "{items.name}", "名称": "{items.name}",
    "规格": "{items.spec}", "规格型号": "{items.spec}", "型号": "{items.model}",
    "颜色": "{items.color}", "材料": "{items.material}", "材质": "{items.material}",
    "单位": "{items.unit}", "数量": "{items.qty}", "单价": "{items.price}", "金额": "{items.amount}",
    "重量": "{items.weight}", "备品": "{items.spare}", "备注": "{items.remark}",
    "图号": "{items.drawing_no}", "加工要求": "{items.process}", "尺寸": "{items.size}",
    "长度": "{items.length}", "厚度": "{items.thickness}", "批号": "{items.batch_no}",
    "订单数量": "{items.order_qty}", "送货数量": "{items.qty}", "实收数量": "{items.recv_qty}",
    "件数": "{items.pcs}", "件数(桶)": "{items.pkg_barrel}", "件数(袋)": "{items.pkg_bag}",
    "交货数量": "{items.qty}", "单价元/cm2": "{items.price_cm2}", "单价\n元/cm2": "{items.price_cm2}",
    "生产批次": "{items.batch_no}", "品牌": "{items.brand}", "封装": "{items.package}",
    "料号": "{items.product_no}", "物料名称": "{items.name}", "物料规格及型号": "{items.spec}",
    "产品名称及型号": "{items.name}", "订单": "{items.order_no}", "订货单号": "{items.order_no}",
    # —— 归一化映射（tok() 会先去掉表头里的空格/换行再查表）——
    "產品名稱": "{items.name}", "货品名称": "{items.name}", "货品名称及规格": "{items.name}",
    "商品名称": "{items.name}", "商品名称及规格": "{items.name}", "品名规格": "{items.name}",
    "品名/规格": "{items.name}", "品名(规格)": "{items.name}", "品名": "{items.name}",
    "名稱及規格": "{items.name}", "名称及规格": "{items.name}", "名称/型号": "{items.name}",
    "产品名称及规格": "{items.name}", "产品名称/规格": "{items.name}", "产品名称DESCRIPTION": "{items.name}",
    "产品编号/名称": "{items.name}", "货名及规格": "{items.name}", "货名DESCRIPTION": "{items.name}",
    "品名规格DESCRIPTION": "{items.name}", "产品型号": "{items.model}",
    "规格SPEC": "{items.spec}", "规格(kg)SPEC": "{items.spec}", "出片规格": "{items.spec}",
    "包装规格": "{items.spec}", "包装规格(KG/桶)": "{items.spec}", "尺码": "{items.size}",
    "尺码Dimensions": "{items.size}", "成品尺寸/套数": "{items.size}", "码数": "{items.size_no}",
    "單位": "{items.unit}", "單位UNIT": "{items.unit}",
    "數量": "{items.qty}", "數量Quantity": "{items.qty}", "数量QUANTITY": "{items.qty}",
    "数量(桶)": "{items.qty_barrel}", "数量(桶)QUANTITY": "{items.qty_barrel}", "数量(PCS)": "{items.qty}",
    "單價": "{items.price}", "單價UnitPrice": "{items.price}", "单价UNITPRICE": "{items.price}",
    "单价UNIT PRICE": "{items.price}", "不含稅單價": "{items.price_net}", "价格(元/KG)": "{items.price_kg}",
    "箱价": "{items.box_price}",
    "金額": "{items.amount}", "金額Amount": "{items.amount}", "總金額Amount": "{items.amount}",
    "金额AMOUNT": "{items.amount}", "金额Amount": "{items.amount}", "金额(元)": "{items.amount}",
    "金额(万千百拾元角分)": "{items.amount}", "金额(十万千百拾元角分)": "{items.amount}",
    "小计": "{items.subtotal}",
    "備注": "{items.remark}", "備註Remark": "{items.remark}", "备注REMARKS": "{items.remark}",
    "备注Remark": "{items.remark}", "备注REMARK": "{items.remark}",
    "訂單編號": "{items.order_no}", "订购单号": "{items.order_no}", "发货单号": "{items.ship_no}",
    "工程单号": "{items.project_no}", "客户PO": "{items.customer_po}", "合同PO": "{items.contract_no}",
    "货号ITEM.": "{items.product_no}", "货号NO.": "{items.product_no}", "編號": "{items.product_no}",
    "编號Number": "{items.product_no}", "编码": "{items.product_no}", "代码": "{items.product_no}",
    "工件编号": "{items.product_no}", "零件编号": "{items.product_no}", "货品编号": "{items.product_no}",
    "条形码": "{items.barcode}", "客户料号": "{items.product_no}", "客户号": "{items.customer_no}",
    "序號NO.": "{items._index}", "P数": "{items.pages}", "套数": "{items.sets}",
    "包数": "{items.packs}", "片数": "{items.pieces}", "箱数": "{items.boxes}",
    "色号COLORNO.": "{items.color_no}", "款式": "{items.style}", "款式Model": "{items.style}",
    "类型": "{items.type}", "适用车型": "{items.vehicle_model}", "工艺": "{items.craft}",
    "内容": "{items.content}", "文件名": "{items.file_name}", "文件名 称": "{items.file_name}",
    "服务项目": "{items.service_item}", "材料费": "{items.material_cost}", "衣片部位": "{items.part}",
    "幅数": "{items.frames}", "平方数": "{items.area}", "面积": "{items.area}",
    "包装": "{items.packing}", "稅金": "{items.tax}", "折扣": "{items.discount}",
    "生产日期": "{items.produce_date}", "生产批号": "{items.batch_no}", "有效期": "{items.expiry}",
    "有效日期": "{items.expiry_date}", "保质期": "{items.shelf_life}", "重量(KG)": "{items.weight}",
    "货号ITEM": "{items.product_no}", "金额(十万千百拾元角分)": "{items.amount}",
    "单价UnitPrice": "{items.price}", "单位UNIT": "{items.unit}", "总金额Amount": "{items.amount}",
    "数量Quantity": "{items.qty}", "编號": "{items.product_no}", "订单号FO:NO": "{items.order_no}",
    "訂單號\nPO:NO": "{items.order_no}", "產品名稱\nDESCRIPTION": "{items.name}",
    "產品料號\nP/N": "{items.product_no}", "型    号": "{items.model}", "型  号": "{items.model}",
    "本页小计": None, "合计": None,
}
NUMERIC = ("数量", "单价", "金额", "重量", "订单数量", "送货数量", "实收数量",
           "件数", "件数(桶)", "件数(袋)", "交货数量")


def F(sz=11, bold=False, color=None, name=SONG):
    return Font(name=name, size=sz, bold=bold, color=color)


import re

_NORM_TOK = {re.sub(r"\s", "", k).lower(): v for k, v in TOK.items()}


def tok(header):
    """表头 → 占位符：先原样查，再去掉空格/换行并忽略大小写归一化查；查不到给中文占位符兜底。"""
    h = header.strip()
    return TOK.get(h) or _NORM_TOK.get(re.sub(r"\s", "", h).lower()) or ("{items." + h + "}")


def build(spec: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "打印模板"
    cols = spec["cols"]
    n = len(cols)
    side = spec.get("side")
    ncols = n + (1 if side else 0)
    last = get_column_letter(ncols)
    grid = spec.get("border", "grid") == "grid"
    row = 1

    # ---- 头部：标题 / 地址电话行 / 右上单据名+NO.+日期 ----
    head = spec.get("head", [])          # [(text, size, bold, red, align)] 脱敏后的头部队行
    title = spec.get("title")
    doc = spec.get("doc")                # 右上 [送 货 单]
    no_date = spec.get("no_date", True)  # 右上 NO./日期两行

    if title:
        ws.merge_cells(f"A{row}:{last}{row}")
        t = spec.get("title_size", 18)
        c = ws.cell(row, 1, title)
        c.font = F(t, True, RED if spec.get("title_red") else None)
        c.alignment = CN
        ws.row_dimensions[row].height = t + 12
        row += 1
    for text, red in head:
        ws.merge_cells(f"A{row}:{last}{row}")
        c = ws.cell(row, 1, text)
        c.font = F(10.5, False, RED if red else None)
        c.alignment = CN if spec.get("head_center", True) else LT
        row += 1

    # 单据名居中（如 送 货 单 独占一行大标题）
    if doc and spec.get("doc_center"):
        ws.merge_cells(f"A{row}:{last}{row}")
        c = ws.cell(row, 1, doc)
        c.font = F(15, True)
        c.alignment = CN
        ws.row_dimensions[row].height = 26
        row += 1

    # 右上：单据名 / NO. / 日期（起点列避开 info 区合并范围）
    info = spec.get("info", [])          # [[(text, span)], ...] 从左排；右侧块单独给
    right = spec.get("right", [])        # 右上若干行文本（放最后两列）
    info_rows = max(len(info), len(right))
    info_span_max = max((sum(s for _, s in rowinfo) for rowinfo in info), default=0)
    right_c1 = max(ncols - 2, info_span_max + 1)
    for i in range(info_rows):
        r = row + i
        ws.row_dimensions[r].height = 20
        if i < len(info):
            c1 = 1
            for text, span in info[i]:
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c1 + span - 1)
                c = ws.cell(r, c1, text)
                c.font = F(11)
                c.alignment = LT
                c1 += span
        if i < len(right) and right_c1 <= ncols:
            if right_c1 < ncols:
                ws.merge_cells(start_row=r, start_column=right_c1, end_row=r, end_column=ncols)
            c = ws.cell(r, right_c1, right[i])
            c.font = F(11, i == 0 and doc is None)
            c.alignment = Alignment(horizontal="left", vertical="center")
    row += info_rows

    # ---- 明细表 ----
    head_row = row
    ws.row_dimensions[head_row].height = 20
    for j, h in enumerate(cols):
        c = ws.cell(row, j + 1, h)
        c.font = F(11, True)
        c.alignment = CN
        if grid:
            c.border = GRID
        else:
            c.border = Border(bottom=THIN)
    if side:
        ws.merge_cells(start_row=row, start_column=ncols, end_row=row + 1 + spec.get("blank", 6), end_column=ncols)
        c = ws.cell(row, ncols, side)
        c.font = F(10)
        c.alignment = VERT
    row += 1
    loop_row = row
    ws.row_dimensions[loop_row].height = 20
    for j, h in enumerate(cols):
        c = ws.cell(row, j + 1, tok(h))
        c.font = F(11)
        c.alignment = CN if h in ("序号", "单位", "数量") else LT
        if grid:
            c.border = GRID
        else:
            c.border = Border(bottom=DASH)
    row += 1
    if spec.get("blank_note"):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        c = ws.cell(row, 2, spec["blank_note"])
        c.font = F(10)
        c.alignment = CN
    for _ in range(spec.get("blank", 6)):
        ws.row_dimensions[row].height = 20
        for j in range(ncols if grid else n):
            c = ws.cell(row, j + 1)
            if grid:
                c.border = GRID
            elif j < n:
                c.border = Border(bottom=DASH)
        row += 1
    if not grid:
        for j in range(n):
            ws.cell(row - 1, j + 1).border = Border(bottom=DASH)

    # ---- 合计行：大写标签 + 各合计写在自身列下；small_total 时金额列带 ￥ ----
    total_cn = spec.get("total_cn", True)
    sums = spec.get("sums", [h for h in cols if h in NUMERIC])
    small = spec.get("small_total", True)
    label = spec.get("total_label", "合计金额(大写)：")
    occ = {}
    for h in sums:
        if h in cols:
            occ[cols.index(h) + 1] = f"{{items.{tok(h)[7:-1]}#sum}}"
    if small and "金额" in cols:
        amt_col = cols.index("金额") + 1
        occ[amt_col] = "￥" + occ[amt_col] if amt_col in occ else "{items.amount#sum}"
    r = row
    ws.row_dimensions[r].height = 22
    if total_cn:
        end = max(min(occ) - 1, 1) if occ else n
        if end > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=end)
        c = ws.cell(r, 1, f"{label}{{_total_cn}}")
        c.font = F(11, True)
        c.alignment = LT
    for cidx, v in sorted(occ.items()):
        ws.cell(r, cidx, v).font = F(11, True)
        ws.cell(r, cidx).alignment = RT
    if grid:
        for j in range(1, ncols + 1):
            ws.cell(r, j).border = GRID
    row += 1

    # ---- 备注/注 ----
    for note in spec.get("notes", []):
        ws.merge_cells(f"A{row}:{last}{row}")
        c = ws.cell(row, 1, note)
        c.font = F(10.5)
        c.alignment = LT
        row += 1
    if spec.get("remark_line"):
        ws.merge_cells(f"A{row}:{last}{row}")
        ws.cell(row, 1, f"备注：{{remark}}").font = F(11)
        row += 1

    # ---- 签名 ----
    sigs = spec.get("sigs", [])
    if sigs:
        row += 1
        ws.row_dimensions[row].height = 24
        per = max(1, ncols // len(sigs))
        for i, s in enumerate(sigs):
            c1 = 1 + i * per
            c2 = ncols if i == len(sigs) - 1 else c1 + per - 1
            if c2 > c1:
                ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
            c = ws.cell(row, c1, s)
            c.font = F(11)
            c.alignment = LT
        row += 1
    if spec.get("page_no"):
        ws.merge_cells(start_row=row, start_column=max(ncols - 2, 1), end_row=row, end_column=ncols)
        ws.cell(row, max(ncols - 2, 1), "第1页，共1页").font = F(10)
        ws.cell(row, max(ncols - 2, 1)).alignment = RT
    for note in spec.get("notes2", []):   # 签名区之后的附加行（如底部公司地址电话，已脱敏）
        ws.merge_cells(f"A{row}:{last}{row}")
        c = ws.cell(row, 1, note)
        c.font = F(10.5)
        c.alignment = LT
        row += 1

    # ---- 列宽：优先用截图网格线实测比例（缩放到与启发式总宽一致） ----
    heuristic = []
    for h in cols:
        if h in ("序号", "单位", "数量", "颜色", "重量", "备品"):
            heuristic.append(7)
        elif h in ("单价", "金额", "订单号", "合同号码", "产品编号", "货号"):
            heuristic.append(11)
        elif h in ("产品名称", "项目名称", "品名及规格", "产品名称、规格", "品名"):
            heuristic.append(22)
        else:
            heuristic.append(12)
    if side:
        heuristic.append(3.5)
    widths = spec.get("widths")
    if not widths:
        px = _PX_WIDTHS.get(spec["f"])
        if px and len(px) == len(heuristic):
            scale = min(sum(heuristic), 88.0) / sum(px)
            widths = [max(3.0, round(p * scale, 1)) for p in px]
        else:
            widths = heuristic
    if sum(widths) > 88.0:   # 超出 A4 可打印宽度则整体等比收缩
        k = 88.0 / sum(widths)
        widths = [max(3.0, round(w * k, 1)) for w in widths]
    for j, w in enumerate(widths[:ncols]):
        ws.column_dimensions[get_column_letter(j + 1)].width = w

    out = OUT / f"{spec['f']}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


SPECS = []


def S(f, cols, **kw):
    SPECS.append(dict(f=f, cols=cols, **kw))


ADDR_TEL = [("地址：XXX", False), ("电话：XXX", False)]
ADDR_TEL_RED = [("地址：XXX", True), ("电话：XXX", True)]

# ============================================================ 五金行业
S("五金行业/五金行业送货单H356", ["订单号", "产品名称", "数量", "单价", "金额", "重量", "备品", "备注"],
  title="XXX公司送货单", head=ADDR_TEL, border="grid", blank=8, side="白：存根\n红/蓝/绿：客户\n黄：回单",
  info=[[("客户名称：{customer}", 5)]],
  right=["NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货物已收妥，数量核对无误。"],
  sigs=["收货人签名：", "送货经手人：", "制单：{handler}"])

S("五金行业/五金行业送货单H362", ["序号", "项目名称", "产品编号", "产品名称", "规格", "单位", "数量", "备品", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送 货 单", doc_center=True, border="grid", blank=8, blank_note="以下空白",
  side="(白)存根联(红)客户联(黄)仓库联(绿)回单",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 3)],
        [("客户地址：{address}", 6)]],
  right=["NO.：{order_no}", "送货日期：{delivery_date}"],
  sums=["数量", "金额"], small_total=True, total_cn=False, total_label="合计：",
  sigs=["收货单位(盖章)：", "制单人：{handler}", "送货单位(盖章)："])

S("五金行业/五金行业送货单H372", ["序号", "品名及规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司送货单", border="grid", blank=8,
  head=[("地址：XXX  电话：XXX", False)],
  info=[[("客户名称：{customer}", 4)], [("客户地址：{address}", 4)]],
  right=["单号：{order_no}", "日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("五金行业/五金行业送货单H500", ["序号", "产品编号", "产品名称、规格", "单位", "数量", "单价", "金额"],
  title="XXX公司销售单", border="grid", blank=10,
  info=[[("客户名称：{customer}", 3), ("联系电话：{phone}", 2)],
        [("客户地址：{address}", 5)]],
  right=["送货单号：{order_no}", "开单日期：{delivery_date}"],
  total_label="金额<大写>：", sums=["金额"], small_total=True,
  notes=["开单人：{handler}    送货人：        收货人签名：        付款方式：{pay_method}",
         "地址：XXX    电话：XXX",
         "注：货物当面点清，如有质量问题请于一个星期内凭单办理。白色(存根)，红色(客户)，黄色(仓库)"],
  page_no=True)

S("五金行业/五金行业送货单H546", ["产品名称", "颜色", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_size=16, head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=8,
  side="①记帐\n②客户\n③回单",
  info=[[("客户名称：{customer}", 4)], [("地址电话：{address}", 4)]],
  right=["[送 货 单]", "NO.：{order_no}", "送货车号：{vehicle_no}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("五金行业/五金行业送货单H577", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, border="dash", blank=8,
  info=[[("客户单位：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("联系电话：{phone}", 2)]],
  right=["NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["制单人员：{handler}", "送货人员：", "收货单位签名(盖章)："],
  notes=["1.存根联(白)  2.客户联(红)"])

S("五金行业/五金行业送货单H591", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="[送货单]", border="grid", blank=10,
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送货单]\n(条码位)", "NO.：{order_no}", "PR单号：", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("五金行业/送货单格式H512", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=10,
  side="①白存根\n②红客户\n③黄回单",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("客户电话：{phone}", 2)]],
  right=["[送 货 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["付款方式：{pay_method}",
         "注：以上货品请核对数量，如有质量问题，请在收货后30天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("五金行业/送货单样式H535", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="送 货 单", border="dash", blank=6,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)], [("客户地址：{address}", 5)]],
  right=["送 货 单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["制单人员：{handler}", "送货人员：", "收货单位签名(盖章)："],
  notes=["1.存根联(白)  2.客户联(红)"])

S("五金行业/送货单样式H223", ["合同号码", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=9,
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送 货 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["发货单位盖章：", "收货单位经手人："], page_no=True)

S("五金行业/送货单样式H268", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=[("地址：XXX  电话：XXX", True)], doc="[送 货 单]",
  border="grid", blank=8,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("备注：{remark}", 2)]],
  right=["[送 货 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("五金行业/送货单样式H271", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=9,
  side="①存根\n②回单\n③客户",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送 货 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["发货单位盖章：", "收货单位经手人："])

# ============================================================ 包装行业
S("包装行业/送货单样式H219", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="收 款 收 据", border="grid", blank=7,
  side="①白存根\n②红客户\n③黄仓库\n④绿回单",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("联系人：{contact}", 2)]],
  right=["收 款 收 据", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("包装行业/送货单样式H283", ["货号\nNO.", "货名\nDESCRIPTION", "数量\nQUANTITY", "单位\nUNIT", "单价\nUNITPRICE", "金额\nAMOUNT", "备注\nREMARKS"],
  title="XXX公司", doc="送 货 单", border="grid", blank=6,
  side="①存根白\n②客户红\n③客户绿\n④回单黄",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)]],
  right=["送 货 单", "NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  notes=["1、如有货物退换，请于三天内通知，逾期恕不受理。",
         "2、记帐户请在发单日期起天内付清款，超期按每日5‰缴滞纳金。",
         "3、在货款未清激或支票未过户前上列货物之拥有权仍属本公司。",
         "4、结帐时需公司派人携带公章的收款收据，送货单(收款凭证)和收款人身份证。",
         "5、货送到请收货员检验清楚，签名并盖公司公章。"],
  sigs=["发货人(签章)：", "收货人(签章)："])

S("包装行业/送货单样式H370", ["订单", "名称", "规格", "单位", "数量", "件数", "单价", "金额"],
  title="XXX公司", head=ADDR_TEL, doc="送 货 单", border="grid", blank=7,
  side="①白存根\n②红客户\n③黄回单",
  info=[[("收货单位：{customer}", 4)]],
  right=["送 货 单", "订单号NO.：{order_no}", "日期：{delivery_date}"],
  total_label="金额合计    大写：", sums=["金额"], small_total=True,
  notes=["1、请按约定准时结清货款。",
         "2、此乃送货凭单，收款时另开正式收款凭单。",
         "3、货品一经验收使用后，本司恕不负责一切损失。"],
  sigs=["验收单位/公司：", "送货单位及负责人/盖章："])

S("包装行业/送货单样本H212", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送 货 单", border="grid", blank=7,
  side="①白存根\n②红客户\n③黄仓库\n④绿回单",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("联系人：{contact}", 2)]],
  right=["送 货 单", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("包装行业/送货单格式H197", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("Dongguan XXX Co., Ltd.", False), ("电话：XXX", False)],
  doc="送 货 单", border="grid", blank=7,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("联系人：{contact}", 2)]],
  right=["送 货 单", "NO.：{order_no}", "合同号码：{contract_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["1.存根(白)  2.收货单位(红)  3.回单(蓝)  4.记账(绿)  5.质检(黄)",
         "注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("包装行业/送货单格式H202", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送 货 单", border="grid", blank=7,
  side="①白存根\n②红客户\n③黄仓库\n④绿回单",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("联系人：{contact}", 2)]],
  right=["送 货 单", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("包装行业/送货单格式H211", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="收 款 收 据", border="grid", blank=7,
  side="①白存根\n②红客户\n③黄仓库\n④绿回单",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("联系人：{contact}", 2)]],
  right=["收 款 收 据", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("包装行业/送货单格式H238", ["订单编号", "产品名称", "规格", "数量", "单价", "金额", "重量(KG)"],
  title="XXX公司", head=[("电话：XXX", False)], doc="發  票\nI N V O I C E", border="grid", blank=8,
  side="①請款白\n②客戶紅\n③對帳黃\n④會計綠\n⑤存根藍",
  info=[[("客户：{customer}", 4)]],
  right=["發  票\nI N V O I C E", "NO.：{order_no}", "Date：{delivery_date}"],
  total_label="※收款時需另發正式收據", sums=["金额"], small_total=True,
  sigs=["簽收蓋章：", "制表：{handler}"])

S("包装行业/送货单格式H258", ["订单编号", "料号", "品名规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=8,
  info=[[("客户名称：{customer}", 4)], [("客户电话：{phone}", 4)], [("客户地址：{address}", 4)]],
  right=["[送 货 单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=["金额"], small_total=False, total_label="合计",
  notes=["以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。",
         "第一联白  存根        第二联红  客户        第三联黄  仓库        第四联绿  会计"],
  sigs=["收货单位签名：", "送货人签名：", "制单人：{handler}"])

S("包装行业/送货单格式H281", ["货号\nNO.", "货名\nDESCRIPTION", "数量\nQUANTITY", "单位\nUNIT", "单价\nUNITPRICE", "金额\nAMOUNT", "备注\nREMARKS"],
  title="XXX公司", doc="送 货 单", border="grid", blank=6,
  side="①存根白\n②客户红\n③客户绿\n④回单黄",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)]],
  right=["送 货 单", "NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  notes=["1、如有货物退换，请于三天内通知，逾期恕不受理。",
         "2、记帐户请在发单日期起天内付清款，超期按每日5‰缴滞纳金。",
         "3、在货款未清激或支票未过户前上列货物之拥有权仍属本公司。",
         "4、结帐时需公司派人携带公章的收款收据，送货单(收款凭证)和收款人身份证。",
         "5、货送到请收货员检验清楚，签名并盖公司公章。"],
  sigs=["发货人(签章)：", "收货人(签章)："])

S("包装行业/销售单样式H387", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=8,
  side="①白存根\n②红客户\n③黑财务\n④绿仓库\n⑤蓝记帐",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送 货 单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["发货单位盖章：", "发货单位经手人：", "收货单位经手人："])

S("包装行业/销售单格式H518", ["序号", "产品名称及型号", "规格", "件数(桶)", "件数(袋)", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=7,
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 3), ("收货人：{contact}", 2)]],
  right=["[送 货 单]", "NO.：{order_no}", "送货日期：{delivery_date}", "订单号码：{contract_no}", "联系电话：{phone}"],
  total_cn=False, sums=["件数(桶)", "数量", "金额"], small_total=False, total_label="合    计",
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

# ============================================================ 电子电器
S("电子电器/电子行业送货单H311", ["訂單號\nPO:NO", "產品名稱\nDESCRIPTION", "產品料號\nP/N", "单位\nUNIT", "数量\nQuantity", "单价\nUnit Price", "總金額\nAmount", "备注\nRemark"],
  title="XXX公司", head=ADDR_TEL, doc="送  貨  單", border="grid", blank=8,
  side="①存根\n②红客户\n③财务\n④回单",
  info=[[("客户名称：{customer}", 4)]],
  right=["送  貨  單", "NO.：{order_no}", "送貨日期：{delivery_date}"],
  total_cn=False, sums=["金额"], small_total=False, total_label="合計人民幣 TOTAL ￥",
  notes=["1.此货物所有权在收货方货款未付清或票据未兑现之前，仍属本公司所有；",
         "2.以上货品请核对数量，如有质量问题，请在收货后2日内通知本公司，逾期恕不负责。"],
  sigs=["经手人：{handler}", "收货人(签章)："])

S("电子电器/送货单样式H313", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", doc="收    据", doc_center=True, border="grid", blank=7,
  info=[[("客户名称：{customer}", 4)]],
  right=["日期：{delivery_date}"],
  total_cn=False, sums=["金额"], small_total=False, total_label="合计",
  notes=["大写(人民币)：{_total_cn}"],
  sigs=["收款人：{handler}", "收款单位盖章："])

S("电子电器/送货单样式H321", ["序号", "产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("XXX Co., Ltd.", False), ("地址：XXX  电话：XXX", False)],
  doc="送 货 单", border="grid", blank=7,
  info=[[("客户编号：{customer_no}", 3), ("客户名称：{customer}", 3)],
        [("客户电话：{phone}", 3), ("客户地址：{address}", 3)]],
  right=["送 货 单", "单号：{order_no}", "送货日期：{delivery_date}", "业务员：{handler}", "制单日期：{_today}", "合同号码：{contract_no}"],
  total_cn=False, sums=["数量"], small_total=False, total_label="合计",
  remark_line=True,
  sigs=["客户签名：", "批准人：", "送货人：", "制单人：{handler}"])

S("电子电器/送货单样式H366", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=9,
  side="①白存根\n②红回单\n③黄客户",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("客户电话：{phone}", 2)]],
  right=["[送 货 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  remark_line=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("电子电器/送货单样式H526", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送 货 单]", border="grid", blank=4,
  side="存根白\n①客户红\n②回单黄",
  info=[[("客户名称：{customer}", 3), ("订货单号：{contract_no}", 2)],
        [("订货日期：{order_date}", 3)]],
  right=["[送 货 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，谢谢合作。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("电子电器/送货单格式H324", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送    货    单", border="grid", blank=7,
  blank_note="(以下空白)", side="白联存根\n红联客户\n黄联财务",
  info=[[("客户名称：{customer}", 3)], [("联系人：{contact}", 3)]],
  right=["送    货    单", "单号：{order_no}", "日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  remark_line=True,
  sigs=["核准：", "财务：", "仓管：", "制单人：{handler}", "客户签收："],
  notes=["注意事项：敬请贵公司收到上列产品后及时清点，如有短缺或误装等问题，请于3天内提出，逾期视同无误。"])

S("电子电器/送货单格式H332", ["产品编号", "产品名称", "规格", "单价", "数量", "单位", "金额", "生产批次"],
  title="XXX公司销售出货单", title_size=16, border="grid", blank=8,
  info=[[("日期：{delivery_date}", 2), ("运输方式：{ship_method}", 2), ("收款期限：", 2)],
        [("送货地址：{address}", 4), ("单号：{order_no}", 2)],
        [("购货单位：{customer}", 4), ("客户订单号：{contract_no}", 2)]],
  right=["第1页，共1页"],
  total_label="合计：", sums=["金额"], small_total=True,
  remark_line=True,
  notes=["销货单位：XXX公司    地址/电话：XXX",
         "声明：1.在贵公司未付清应付我司货款之前，此单货物所有权和处理权仍属我司。",
         "      2.贵司如果不按合同约定时间付款，我司有权停止供货，且贵司因此造成的损失由贵司自行承担，与我司无关。",
         "      3.贵司委托收货人在收取货物时请仔细检查货品包装及数量等，签字后即认同此送货单无误，事后不得再提出异议。"],
  sigs=["客户委托收货人签名：", "制单人：{handler}"])

S("电子电器/送货单格式H389", ["产品编号", "产品名称", "规格", "单价\n元/cm2", "单价", "交货数量", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="【送  货  单】", border="grid", blank=4,
  info=[[("客户名称：{customer}", 4)], [("客户地址：{address}", 4)]],
  right=["【送  货  单】", "NO.：{order_no}", "联系人：{contact}", "送货日期：{delivery_date}"],
  total_label="货款合计：(大写)", sums=["金额"], small_total=True,
  sigs=["发货单位：(签字/盖章)", "送货人签字：", "收货单位经手人(签字/盖章)："])

S("电子电器/送货单格式H501", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=4,
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("电子电器/送货单格式H520", ["序号", "型  号", "品牌", "封装", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("XXX Co., Ltd.", False), ("地址：XXX", False)],
  doc="收    据", border="grid", blank=8,
  info=[[("客户名称：{customer}", 4)], [("客户地址：{address}", 2), ("订单编号：{contract_no}", 2)]],
  right=["收    据", "单号：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=["数量", "金额"], small_total=False, total_label="合计",
  notes=["注：1.需方收到货四周内，若对货的品质有异议或认为有品质问题，则请于收到货起四周内传真详细的测试报告供方确认属实，可退货或换货逾期不予处理。",
         "2.货款未全部兑现前，货物所有权仍归属本公司全权所有。",
         "3.如有疑问，请于收货之日起三日内检查提出，否则视同接受本单所载所有事项。"],
  sigs=["客户签名：", "批准人：", "送货人：", "制单人：{handler}"])

S("电子电器/送货单格式H530", ["№", "物料名称", "物料规格及型号", "单位", "数量", "单价", "金额", "备注"],
  title="销售发货单", title_size=16, border="grid", blank=7,
  info=[[("销售单号：{order_no}", 2), ("销售日期：{delivery_date}", 2)],
        [("客户名称：{customer}", 2), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 4)]],
  right=["XXX公司", "电话：XXX", "地址：XXX"],
  total_label="金额合计：", sums=["金额"], small_total=True,
  notes=["备注：① 尊敬的客户，如您对所送货物数量与品质有任何疑问，请于三日内联系我们",
         "      ② 白联-留存  黄联-仓库  红联-客户"],
  sigs=["制单人：{handler}", "核准：", "客户签收："])

S("电子电器/送货单格式H545", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=9,
  side="①存根\n②客户\n③仓库\n④记账",
  info=[[("客户名称：{customer}", 3)], [("地址电话：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

# ============================================================ 贸易行业
S("贸易行业/送货单样式H262", ["货品编号", "品名规格", "数量(PCS)", "单价", "金额", "订单编号"],
  title="XXX公司", head=ADDR_TEL, doc="产品送货单", border="grid", blank=6,
  info=[[("客户名称：{customer}", 4)]],
  right=["产品送货单", "№ {order_no}", "日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["帐期：                币别：人民币",
         "仓管：        发货人：        收货单位签名及盖章：",
         "备注：如有质量及其它问题请三天内提出，逾期视为接受此单。"],
  page_no=True)

S("贸易行业/送货单样式H309", ["序号", "商品名称", "规格型号", "单位", "数量", "单价", "金额", "颜色", "备注"],
  title="XXX公司", doc="出 库 单", head=ADDR_TEL, border="grid", blank=6,
  side="①存根\n②财务\n③仓库\n④网点",
  info=[[("网店名称：{customer}", 3), ("网店电话：{phone}", 2)]],
  right=["出 库 单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  sigs=["开单人：{handler}", "发货人：", "送货人：", "收货人："])

S("贸易行业/送货单样式H345", ["编码", "产品名称", "规格", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="出  货  单", border="grid", blank=7, blank_note="以下空白",
  info=[[("客户名称：{customer}", 3), ("地址电话：{address}", 3)]],
  right=["出  货  单", "NO.：{order_no}", "客户单号：{contract_no}", "日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["备注：1.按桶____个，退桶____个，欠____个，请贵司在收货时检验好质量及数量后再签名盖章。",
         "      2.贵司所购物料请在保质期内用完，超期发生质量问题供方不负任何责任。",
         "      3.付款方式：{pay_method}",
         "第一联：存根(白)  第二联：客户(红)  第三联：财务(蓝)  第四联：仓库(黄)"],
  sigs=["开单：{handler}", "核对：", "送货：", "客户签名(签章)："])

S("贸易行业/送货单样式H513", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=9,
  side="①白存根\n②红客户\n③青回单\n④蓝记账",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("客户电话：{phone}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "合同号码：{contract_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  remark_line=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("贸易行业/送货单样式H525", ["发货单号", "货物名称", "包装规格", "重量", "数量", "实收数量", "备注"],
  title="XXX公司 送货单", title_size=14, border="grid", blank=7,
  info=[[("发货日期：{delivery_date}", 2), ("约定送达日期：", 2), ("地区：", 1), ("车号：{vehicle_no}", 1)],
        [("客户编号：{customer_no}", 1), ("客户名称：{customer}", 2), ("地址：{address}", 2)],
        [("联系人：{contact}", 1), ("电话：{phone}", 2), ("业务经理：{handler}", 2)]],
  right=["NO.：{order_no}"],
  total_cn=False, sums=["数量"], small_total=False, total_label="合    计",
  notes=["物流须送货到门，不得收取任何费用。",
         "以上内容请客户认真填写，一切以本单签署内容为准！"],
  sigs=["制单人：{handler}", "提货人：", "客户："])

S("贸易行业/送货单格式H066", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="收款专用票据", border="grid", blank=5,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)], [("客户地址：{address}", 5)]],
  right=["收款专用票据", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  remark_line=True,
  sigs=["制单人员：{handler}", "收货单位及经手人(盖章)："],
  notes=["1.存根联(白)  2.客户联(红)"])

S("贸易行业/送货单格式H108", ["訂單編號", "產品名稱", "單位", "數量", "單價", "金額", "備注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="[送  貨  單]", border="grid", blank=9,
  side="①存根白\n②财务兰\n③客户红\n④请款黄",
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 5)]],
  right=["[送  貨  單]", "NO.：{order_no}", "送貨日期：{delivery_date}"],
  total_label="合計金額(大寫)：", sums=["金额"], small_total=True,
  notes=["注：以上貨品請核對數量，如有品質問題，請在收貨後3天內通知本公司，逾期恕不負責"],
  sigs=["發貨單位蓋章：", "發貨單位經手人：{handler}", "收貨單位經手人："])

S("贸易行业/送货单格式H232", ["产品名称", "规格", "数量", "面积", "单位", "单价", "金额", "备注"],
  title="XXX公司", doc="送  货  单", border="dash", blank=2,
  info=[[("地址：XXX", 4)], [("货主名称：{customer}", 4)], [("地址：{address}", 4)]],
  right=["送  货  单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["货主名称：{customer}        联系人：{contact}        联系电话：{phone}",
         "说明：1.公司送货上门时，商家结清货款第二、三联由商家留存。2.如果商家因特殊情况仅收货未付清货款第三联由收货人留存告知负责人尚未付款明细。3.第二联结算联是本公司和商家结算的依据，也是上门收款的凭证，请商家妥善保管，谢谢合作。",
         "本单三联：第一联办事处联  第二联结算联  第三联收货联"],
  sigs=["制单人：{handler}", "收货人签名："])

S("贸易行业/送货单格式H341", ["序號\nNO.", "產品名稱\nDESCRIPTION", "规格\nSPEC", "單位\nUNIT", "數量\nQuantity", "單價\nUnit Price", "總金額\nAmount", "備註\nRemark"],
  title="XXX公司", head=[("XXX Co., Ltd.", False)], doc="送  貨  單", border="grid", blank=8,
  side="①存根白\n②回單紅\n③客户藍",
  info=[[("客户名称：{customer}", 3)], [("客户地址：{address}", 3)]],
  right=["送  貨  單", "NO.：{order_no}", "联系人：{contact}", "客户电话：{phone}", "送货日期：{delivery_date}"],
  total_label="金額合計(大寫)：", sums=["金额"], small_total=True,
  sigs=["送貨單位及經手人(蓋章/簽字)：", "收貨單位及經手人(蓋章/簽字)："])

S("贸易行业/送货单格式H523", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("地址：XXX  电话：XXX", False)], doc="[送  货  单]", border="grid", blank=3,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("备注：{remark}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("贸易行业/送货单格式H541", ["序号", "商品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司（送货单）", title_size=16, border="grid", blank=1,
  side="白存根联\n红客户联",
  info=[[("客户名称：{customer}", 3), ("联系电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("送货日期：{delivery_date}", 2)],
        [("代收帐户：", 3), ("账户姓名：", 2)]],
  right=["单号：{order_no}", "结款："],
  total_label="合计  (大写)：", sums=["数量", "金额"], small_total=True,
  notes=["地址：XXX  电话：XXX",
         "经营范围：XXX",
         "农行账号：XXX    户名：XXX",
         "信合账号：XXX    户名：XXX",
         "工商账号：XXX    户名：XXX"],
  sigs=["制单人：{handler}", "送货人：", "收货人签名："])

S("贸易行业/送货单格式H556", ["序号", "产 品 名 称", "规 格", "颜 色", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司送货单", head=ADDR_TEL, border="grid", blank=8,
  side="①白存根\n②红回单\n③黄客户",
  info=[[("收货单位：{customer}", 3), ("联系人：{contact}", 2)],
        [("地址：{address}", 3), ("电话：{phone}", 2)]],
  right=["单号：{order_no}", "日期：{delivery_date}"],
  total_cn=False, total_label="总计：", sums=["金额"], small_total=True,
  remark_line=True,
  notes=["以上货品由送货人员当场核对数量，离点缺少概不负责！"],
  sigs=["制单：{handler}", "发货人：", "送货人：", "送货人电话：", "送货车号：{vehicle_no}", "收货人："])

# ============================================================ 批发行业
S("批发行业/送货单格式H024", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_size=15, head=[("电话：XXX", False)], doc="[送  货  单]", border="grid", blank=2,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("备注：{remark}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("批发行业/送货单格式H094", ["产品名称/规格", "数量", "单位", "单价", "金额", "备注"],
  title="XXX公司供货凭证", title_red=True, head=ADDR_TEL_RED, border="grid", blank=1,
  info=[[("客户：{customer}", 3), ("联系方式：{phone}", 2)],
        [("送货地址：{address}", 3), ("业务员：{handler}", 2)]],
  right=["NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["配送签字：", "客户收货签字："],
  notes=["⑴结算凭证(白)  ⑵客户联(红)  ⑶存根(黄)"])

S("批发行业/送货单格式H277", ["序号", "产品名称", "规格", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("订货电话：XXX", False)], doc="送  货  单", border="grid", blank=12,
  info=[[("客户：{customer}", 3), ("联系电话：{phone}", 2)]],
  right=["送  货  单", "NO.：{order_no}", "日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["①白单(存查)  ②红单(回单)  ③蓝单(客户)",
         "货物请当面点清数量，如有疑问请于当电致，隔日不作处理",
         "全部1页总金额：￥{items.amount#sum}        共1页，第1页"],
  sigs=["制单：{handler}", "送货：", "欠款及经手人："])

S("批发行业/送货单格式H278", ["序号", "产品名称", "规格", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("订货电话：XXX", False)], border="grid", blank=9,
  info=[[("客户：{customer}", 3), ("联系电话：{phone}", 2)]],
  right=["NO.：{order_no}", "日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["货物请当面点清数量，如有疑问请于当电致，隔日不作处理",
         "全部1页总金额：￥{items.amount#sum}"])

S("批发行业/送货单格式H369", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=13,
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。        白单：留底  红单：财务  黄单：客户"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

# ============================================================ 广告行业
S("广告行业/制版行业送货单H845", ["项目名称", "规格", "材料", "数量", "单位", "单价", "金额", "备注"],
  title="XXX公司 接件及送货单", title_size=16, border="grid", blank=5,
  side="①存根\n②客户\n③回单",
  info=[[("客户名称：{customer}", 4)], [("联系电话：{phone}", 4)]],
  right=["NO.：{order_no}", "接件日期：{order_date}", "交货日期：{delivery_date}"],
  total_label="合计人民币(大写)：", sums=["金额"], small_total=True,
  notes=["后期工艺：复膜□  压痕□  模切□  压纹□  压凸□  烫金□  上UV□  上光油□  打孔□",
         "          装订:锁线胶装□  无线胶装□  骑马订□  胶头□  胶边□  包本□  打号□"],
  sigs=["开单：{handler}", "核价：", "设计：", "送货：", "收款："],
  notes2=["地址：XXX", "电话：XXX"])

S("广告行业/送货单格式H011", ["颜色", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX DESIGN\nXXX广告印务", title_size=13, border="grid", blank=8,
  doc="签 收 单",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("电话：{phone}", 2)],
        [("备注：{remark}", 3), ("结算方式：{pay_method}", 2)]],
  right=["送货单号：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=["数量", "金额"], small_total=False, total_label="合计：",
  notes=["金额合计(大写)：{_total_cn}", "温馨所示：1.以上货品已点清。2.如有问题请在收货24小时内投诉。"],
  sigs=["收货经手人：", "送货经手人："],
  notes2=["地址：XXX  电话：XXX"])

S("广告行业/送货单格式H104", ["序号", "产品名称", "数量", "规格", "重量", "单位", "单价", "金额", "备注"],
  title="XXX广告", title_size=14, doc="送  货  单", border="grid", blank=6,
  side="白联：存根\n红联：客户",
  info=[[("客户名称：{customer}", 3), ("电话：{phone}", 2)]],
  right=["地址：XXX", "单号：{order_no}", "开单日期：{order_date}", "交货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["总金额：{items.amount#sum}元  预付款：____元  未付款：____元  付款方式：{pay_method}  交货方式：{ship_method}",
         "注意：请客户认真校稿。画面色彩、文字、内容以客户确认为准，内容涉及侵权、违法、错误由客户自行承担责任。"],
  sigs=["请核对数量和金额无误  客户签字：", "电话：{phone}", "经办人：{handler}"],
  notes2=["电话：XXX  地址：XXX"])

S("广告行业/送货单格式H199", ["序号", "名  称  及  规  格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("XXX ADVERTISING CO., LTD.", False), ("地址：XXX  电话：XXX", False)],
  doc="[送  货  单]", border="grid", blank=8,
  side="白单：回单\n红单：客户\n黄单：存根",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("电话：{phone}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "第1页，共1页", "合同号码：{contract_no}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  notes=["注：请核对以上清单，如有问题，请在3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(签章)：", "制单人：{handler}", "收货单位及经手人(签章)："])

S("广告行业/送货单格式H292", ["编号", "产品名称", "内容", "规格", "幅数", "数量", "单位", "单价", "金额(元)", "备注"],
  title="XXX公司", head=[("电话：XXX", False), ("地址：XXX", False)],
  doc="(送货单)", border="grid", blank=9,
  info=[[("客户名称：{customer}", 3), ("客户地址：{address}", 3)]],
  right=["(送货单)", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计大写金额：", sums=["金额"], small_total=True,
  notes=["1.以上成品已点清，原件、介质盘等已全部取回，来盘加工交件请各自自行备份，若盘内文件不属损坏或丢失，不予赔偿。",
         "2.请客户收货时当面点清，如有错漏，请当场提出，成品收货后概不负责。",
         "3.若对稿件内容有疑问，请及时提出，过时恕不负责，若我方原因负责重印，否则由客户自行负责，本公司不承担任何责任。",
         "4.宽度和高度的单位为米。"],
  sigs=["制单员：{handler}", "送货员：", "收货单位及经手人："])

S("广告行业/送货单格式H360", ["产品编号", "产品名称", "规格", "单位", "数量", "平方数", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="[合 约 单]", border="grid", blank=1,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("备注：{remark}", 2)]],
  right=["[合 约 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("广告行业/送货单格式H367", ["品  名", "规  格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_size=14, border="grid", blank=8,
  side="①存根白\n②客户红",
  info=[[("单位：{customer}", 3), ("签收人：{contact}", 2)],
        [("日期：{delivery_date}", 3), ("电话：{phone}", 2)]],
  right=["NO.：{order_no}"],
  total_label="合计（大写）", sums=["金额"], small_total=True,
  notes=["结算方式：{pay_method}    预付定金：",
         "银行账号：XXX    户名：XXX",
         "地址：XXX  电话：XXX"],
  sigs=["主管：", "财务：", "经手人：{handler}"])

S("广告行业/送货单格式H383", ["产品编号", "产品名称", "文件名", "规格", "数量", "单位", "面积", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="[送  货  单]", border="grid", blank=2,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("备注：{remark}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_cn=False, sums=["数量", "金额"], small_total=False, total_label="合计",
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("广告行业/送货单格式H622", ["名  称  及  规  格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司送货单", title_size=15, head=ADDR_TEL, border="grid", blank=6,
  info=[[("客户名称：{customer}", 4)]],
  right=["№：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["制单人：{handler}", "收货单位及经手人签章："])

# ============================================================ 模具行业
S("模具行业/送货单样式H373", ["序号", "产品名称", "工件编号", "单位", "规格", "材质", "材料费", "数量", "单价", "金额"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=8,
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 5)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计金额：", sums=["金额"], small_total=False,
  remark_line=True,
  notes=["①白存根  ②红回单  ③蓝客户  ④黄记账",
         "注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("模具行业/送货单格式H253", ["工程单号", "客户料号", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="[送  货  单]", border="grid", blank=10,
  side="①白存根\n②红仓库\n③黄客户",
  info=[[("客户名称：{customer}", 3)], [("客户地址：{address}", 3)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, total_label="合    计", sums=["金额"], small_total=True,
  notes=["声明：刀模上机之前，请仔细核对，本公司对刀模产生的后果概不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

# ============================================================ 服装行业
S("服装行业/服装行业送货单H257", ["序号", "颜色", "衣片部位", "码数", "包数", "片数", "备注"],
  title="XXX公司送货单", head=[("XXX Co., Ltd.", False)], border="grid", blank=11,
  side="①存根联\n②客户联\n③回单联",
  info=[[("客户名称：{customer}", 3), ("地址：XXX", 2)],
        [("联系地址：{address}", 3), ("电话：XXX", 2)],
        [("联系人：{contact}", 1), ("联系电话：{phone}", 1), ("传真：XXX", 1)],
        [("NO.：{order_no}", 3), ("款号：{style_no}", 2)]],
  right=[],
  total_cn=False, sums=["包数"], small_total=False, total_label="合计",
  sigs=["日期：{delivery_date}", "制表人：{handler}", "签收人："])

S("服装行业/服装行业送货单H261", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[提  货  单]", border="grid", blank=10,
  info=[[("客户名称：{customer}", 5)]],
  right=["[提  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["发货单位盖章：", "收货单位经手人："])

S("服装行业/服装行业送货单H276", ["订单编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="dash", blank=9,
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 5)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_cn=False, sums=[], small_total=False,
  notes=["此单据经收货单位经办人签收作结款依据"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("服装行业/服装行业送货单H357", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=1,
  info=[[("客户：{customer}", 3), ("联系人：{contact}", 2)],
        [("地址：{address}", 3), ("结算方式：{pay_method}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "订单号：{contract_no}", "日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：请核对以上货品，货品如有质量问题请在三天内通知我司，我们会以最快速度解决问题，逾期者将视为放弃，责任贵司自负。"],
  sigs=["开单：{handler}", "跟单业务：", "客户签章："])

S("服装行业/送货单格式A009", ["合同\nPO", "品名规格\nDESCRIPTION", "色号\nCOLOR NO.", "单位\nUNIT", "数量\nQUANTITY", "单价\nUNIT PRICE", "金额\nAMOUNT"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="送  货  单", border="grid", blank=7,
  info=[[("客户名称：{customer}", 4)]],
  right=["送  货  单", "NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计：TOTAL：", sums=["金额"], small_total=True,
  notes=["以上所送货物敬请验收，并盖上公章。如质量问题，须在七天内通知本司，逾期恕不负责任，多谢合作。"],
  sigs=["经手人：{handler}", "收货人签名盖章："])

# ============================================================ 化工行业
S("化工行业/送货单样式H532", ["序号", "产品编号/名称", "规格", "单位", "单价", "数量", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=10,
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("联系电话：{phone}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "业务跟单：{handler}", "发货日期：{delivery_date}"],
  total_label="合计 大写：", sums=["数量", "金额"], small_total=True,
  notes=["共三联：白联：存根  红联：回单  黄联：客户"],
  sigs=["收货方(盖章)签字：", "发货方(盖章)签字："])

S("化工行业/送货单样式H589", ["型号", "品名", "单位", "包装", "数量", "件数", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送货单", border="grid", blank=9,
  info=[[("购货单位：{customer}", 3)], [("采购单号：{contract_no}", 3)]],
  right=["送货单", "送货单号：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=[], small_total=False,
  notes=["1.第一联：(白)，第二联客户(红)，第三联回单(黄)。",
         "2.如有质量问题，请购货后三日内提出，逾期视为合格。",
         "3.空桶回收。"],
  sigs=["收货人：(签名盖章)", "发货人：(签名盖章)"])

S("化工行业/送货单样式H606", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", doc="送  货  单", border="grid", blank=3,
  side="①存根\n②客户\n③回单",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("联系人：{contact}", 2)]],
  right=["送  货  单", "NO.：{order_no}", "送货车号：{vehicle_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("化工行业/送货单样本H315", ["序号", "产品名称", "规格", "规格", "单位", "数量", "单价", "金额"],
  title="XXX公司", head=ADDR_TEL, doc="送  货  单", border="grid", blank=8,
  side="①白存根\n②红客户\n③蓝仓库\n④黄回单",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("联系人：{contact}", 2)]],
  right=["送  货  单", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  remark_line=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("化工行业/送货单样本H590", ["型号", "品名", "单位", "包装", "数量", "件数", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送货单", border="grid", blank=9,
  info=[[("购货单位：{customer}", 3)], [("采购单号：{contract_no}", 3)]],
  right=["送货单", "送货单号：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=[], small_total=False,
  notes=["1.第一联：(白)，第二联客户(红)，第三联回单(黄)。",
         "2.如有质量问题，请购货后三日内提出，逾期视为合格。",
         "3.空桶回收。"],
  sigs=["收货人：(签名盖章)", "发货人：(签名盖章)"])

S("化工行业/送货单格式H296", ["货号\nITEM.", "货名\nDESCRIPTION", "规格(kg)\nSPEC", "数量(桶)\nQUANTITY", "单价\nUNITPRICE", "金额\nAMOUNT"],
  title="XXX公司", head=ADDR_TEL, doc="供  货  单", doc_center=True, border="grid", blank=7,
  side="①存根白\n②客户红\n③客户绿\n④回单黄",
  info=[[("客户名称：{customer}", 4)]],
  right=["NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  notes=["合约：1.收货单位核对无误后签收，本供货单一经买方或其代理人签字或盖章生效。",
         "     2.收货单位对产品质量有异议，请在收货后五日内书面提出，否则为符合约定。",
         "     3.收货单位应按期支付货款，逾期则按拖欠货款金额的每日万分之五向供货方支付逾期违约金。",
         "     4.如发生纠纷，应协调解决，协调不成由供货方所在地人民法院解决。",
         "     5.本公司业务员无权收款，所有货款均由公司专职人员凭公司专用收据及送货单原件收取货款，否则视为未付。"],
  sigs=["付款方式：{pay_method}", "收货人签名盖章：", "开单员：{handler}", "送货员："])

S("化工行业/送货单格式H586", ["型号", "品名", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送货单", border="grid", blank=9,
  info=[[("购货单位：{customer}", 3)], [("采购单号：{contract_no}", 3)]],
  right=["送货单", "送货单号：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额大写：", sums=["金额"], small_total=True,
  notes=["1.第一联：(白)，第二联客户(红)，第三联回单(黄)。",
         "2.如有质量问题，请购货后三日内提出，逾期视为合格。",
         "3.空桶回收。"],
  sigs=["收货人：(签名盖章)", "发货人：(签名盖章)"])

S("化工行业/送货单格式H588", ["型号", "品名", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送货单", border="grid", blank=9,
  info=[[("购货单位：{customer}", 3)], [("采购单号：{contract_no}", 3)]],
  right=["送货单", "送货单号：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额大写：", sums=["金额"], small_total=True,
  notes=["1.第一联：(白)，第二联客户(红)，第三联回单(黄)。",
         "2.如有质量问题，请购货后三日内提出，逾期视为合格。",
         "3.空桶回收。"],
  sigs=["收货人：(签名盖章)", "发货人：(签名盖章)"])

S("化工行业/送货单格式H600", ["编号", "品    名", "包装规格\n(KG/桶)", "价    格\n(元/KG)", "数量(桶)", "金额", "备注"],
  title="XXX公司发货单", border="grid", blank=9,
  side="①白存根\n②红客户\n③蓝回单",
  info=[[("客户：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户电话：{phone}", 3), ("单号：{order_no}", 1), ("日期：{delivery_date}", 1)]],
  right=["销售代表：{handler}", "销售电话：XXX"],
  total_cn=False, sums=["金额"], small_total=False,
  notes=["收款情况：        金额合计(小写)：{items.amount#sum}        金额合计(大写)：{_total_cn}",
         "备注：收货方依送货单上载明内容，当日清点并签名确认，逾期恕不受理。"],
  sigs=["制单：{handler}", "核价：", "仓库：", "送货人：", "验货签收人："])

# ============================================================ 药品行业
S("药品行业/送货单样式H529", ["代码", "产品名称", "规格", "数量", "单价", "金额", "生产日期", "有效日期", "备注"],
  title="XXX公司送货单", border="grid", blank=9,
  side="①白台账\n②红拣货\n③绿客户\n④蓝记账\n⑤黄仓库",
  info=[[("客户名称：{customer}", 3), ("电话：{phone}", 2)],
        [("客户地址：{address}", 5)]],
  right=["NO.：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=["数量"], small_total=False, total_label="合计：",
  notes=["合计金额(大写)：{_total_cn}"],
  sigs=["收货人：", "审批人：", "发货人：", "业务：", "制单：{handler}"])

S("药品行业/送货单样本H208", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX兽药（批发.零售）", title_red=True, head=ADDR_TEL_RED, doc="收款专用票据",
  border="dash", blank=9,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("付款方式：{pay_method}", 2)]],
  right=["收款专用票据", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["制单人员：{handler}", "收货单位：", "出货单位签名(盖章)："],
  notes=["1.存根联(白)  2.客户联(红)"])

S("药品行业/送货单样本H247", ["产品名称、规格", "单位", "数量", "单价", "金额(万千百拾元角分)", "备注"],
  title="XXX公司", doc="送 货 单", border="grid", blank=9,
  side="①存根(白)\n②客户(红)\n③结算凭证(蓝)",
  info=[[("客户名称：{customer}", 4)]],
  right=["送 货 单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

# ============================================================ 灯饰行业
S("灯饰行业/送货单样本H302", ["序号", "名称及规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX照明", title_red=True, head=ADDR_TEL_RED, doc="[送  货  单]", border="grid", blank=11,
  side="①白存根\n②红回单\n③黄客户",
  info=[[("客户名称：{customer}", 4)], [("客户地址：{address}", 2), ("客户电话：{phone}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("灯饰行业/送货单格式H001", ["型号", "数量", "单价", "金额", "件数", "备注"],
  title="XXX照明订货单", border="grid", blank=10,
  info=[[("订货单位：{customer}", 3), ("供货单位：XXX公司", 2)],
        [("客户签署：", 3), ("门市电话：XXX", 2)],
        [("联系电话：{phone}", 3), ("厂部地址：XXX", 2)],
        [("订货日期：{order_date}", 3), ("厂部电话：XXX", 2)]],
  right=["NO.：{order_no}"],
  total_cn=False, sums=["数量", "金额"], small_total=False, total_label="合计",
  notes=["付款方式：{pay_method}    定金：____    余额：{items.amount#sum}",
         "送货地点：{address}    电话：{phone}    余额大写：{_total_cn}"])

S("灯饰行业/送货单格式H047", ["产品名称", "规格", "单位", "数量", "单价", "金额"],
  title="XXX路灯", title_size=18, border="grid", blank=8,
  info=[[("购货日期：{delivery_date}", 5)],
        [("购货单位：{customer}", 5)],
        [("地址/电话：{address}", 5)]],
  right=["NO.：{order_no}",
         "购货单位须知：一、收取商品请当面验查商品品质及数量，否则破损短缺概不负责。二、样品灯具提走以后超过七天不能退还。三、单价以当日当天为准。"],
  total_label="合计：", sums=["金额"], small_total=True,
  notes2=["地址：XXX    电话：XXX    传真：XXX"])

S("灯饰行业/送货单格式H089", ["货品名称", "规格", "重量", "数量", "单位", "单价", "金额"],
  title="XXX公司", head=[("电话：XXX", False), ("地址：XXX", False)],
  doc="送  货  单", doc_center=True, border="grid", blank=6,
  info=[[("客户名称：{customer}", 4)]],
  right=["单号：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=[], small_total=False, total_label="合计",
  remark_line=True,
  notes=["说明：1、客户对送货单内容如有疑问，请在三天内传真或来电查询更正，否则，此单作结算凭证。",
         "      2、请在结算期内给付货款，如果拖延则按货款金额每日加收3‰滞纳金。",
         "⑴存根(白) ⑵客户联(红) ⑶请款联(黄) ⑷记账联(绿) ⑸收货联(黄)"],
  sigs=["制单人：{handler}", "收货人：", "收货单位盖章："])

# ============================================================ 办公设备行业
S("办公设备行业/送货单格式H251", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX办公设备销售中心", title_red=True, head=ADDR_TEL_RED, doc="[维 修 单]", border="grid", blank=2,
  side="①存根\n②客户\n③回单",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("联系电话：{phone}", 2)]],
  right=["[维 修 单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("办公设备行业/送货单格式H330", ["货号", "货名及规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX办公设备经营部", doc="送 货 单", border="grid", blank=9,
  side="①存根白\n②客户红\n③回单黄",
  info=[[("客户名称：{customer}", 5)]],
  right=["送 货 单", "NO.：{order_no}", "日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["收货单位经手人：", "送货单位经手人："])

# ============================================================ 制版设计行业
S("制版设计行业/制版行业送货单H578", ["服务项目", "文 件 名 称", "规  格", "数量", "金额"],
  title="XXX公司［出库单］", title_size=15, border="grid", blank=3,
  side="①存根白\n②客户红\n③收据黄",
  info=[[("客户：{customer}", 5), ("日期：{delivery_date}", 2), ("NO.：{order_no}", 2)]],
  right=[],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["备注：{remark}", "制单：{handler}", "收货单位经手人(盖章)："])

S("制版设计行业/制版行业送货单H615", ["产品编号", "品 名(规 格)", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("XXX Co., Ltd.", False), ("电话：XXX", False)],
  doc="送  货  单", doc_center=True, border="grid", blank=7,
  info=[[("收货单位：{customer}", 3), ("采购单号：{contract_no}", 3)]],
  right=["NO：{order_no}", "日期：{delivery_date}"],
  total_label="合计人民币：", sums=["金额"], small_total=True,
  sigs=["收货单位及经手人：", "送货单位及经手人："],
  notes=["注明：第一联 存根(白色)  第二联 顾客(红色)  第三联 财会(黄色)  第四联 仓库(绿色)  第五联 采购(蓝色)"])

S("制版设计行业/制版行业送货单H617", ["序号", "订单编号", "文 件 名", "成品尺寸/套数", "颜色", "P数", "单价", "金额"],
  title="XXX公司", title_size=14, doc="送  货  单", border="grid", blank=5,
  side="①存根白\n②回单红\n③客户黄",
  info=[[("客户：{customer}", 3), ("电话：{phone}", 3)],
        [("地址：{address}", 3), ("日期：{delivery_date}", 3)]],
  right=["送  货  单", "№：{order_no}"],
  total_cn=False, sums=["金额"], small_total=False, total_label="合计",
  notes=["注：客户有责任仔细检查菲林，如有错漏我方有责任修改重出，但概不承担一切因错漏页造成的直接或间接的其他经济损失。多谢合作！",
         "地址：XXX  电话：XXX  结款方式：{pay_method}"],
  sigs=["制表：{handler}", "客户签收：", "货款(    )付"])

S("制版设计行业/制版行业送货单H618", ["客户号", "文件名", "出片规格", "套数", "单价", "金额", "客户号"],
  title="XXX设计", title_size=15, head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=4,
  side="①白存根\n②红客户\n③黄回单",
  info=[[("客户名称：{customer}", 3), ("制网前请认真核对胶片，制作、网后概不负责。", 4)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计", sums=["金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("制版设计行业/制版行业送货单H640", ["产品名称", "规格", "单位", "数量", "单价", "金额"],
  title="XXX图文", title_size=16, doc="[送  货  单]", border="grid", blank=1,
  info=[[("客户名称：{customer}", 3), ("项目名称：{project_name}", 3)],
        [("客户地址：{address}", 3), ("制单人员：{handler}", 3)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["备注：收货时请确认质量、数量等是否符合原定要求，有任何问题请于收货当天提出。",
         "公司地址：XXX    公司电话：XXX"],
  sigs=["送货签名：", "客户签名："])

S("制版设计行业/制版行业送货单格式H655", ["序号", "产品名称", "工艺", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("XXX PRINT CO.,LTD", False), ("地址：XXX  电话：XXX", False)],
  doc="送货单 DELIVERY NOTE", doc_center=True, title_size=14, border="grid", blank=7,
  info=[[("客户名称：{customer}", 3), ("单号：{order_no}", 2)],
        [("地址：{address}", 3), ("日期：{order_date}", 2)],
        [("电话：{phone}", 3), ("送货日期：{delivery_date}", 2)]],
  right=[],
  total_label="大写：", sums=["金额"], small_total=True,
  notes=["设计师：        跟单：        业务：",
         "1.以上所送货品请点物品数量敬请验收，并签字。如质量问题，须在5个工作日内通知本司，逾期恕不负责任。  2.本单带有收据及订单效力。金额单位：元"],
  sigs=["客户签收：", "日期：", "送货人：{handler}"])

S("制版设计行业/印刷行业送货单H065", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX印刷有限公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=11,
  side="①白存根\n②红客户\n③兰财务\n④黄请款",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["发货单位盖章：", "发货单位经手人：{handler}", "收货单位经手人："])

S("制版设计行业/印刷行业送货单H659", ["序号", "规  格", "数量", "金额", "折扣", "小计", "备注"],
  title="XXX快印联盟送货单", title_size=14, head=ADDR_TEL, border="grid", blank=8,
  info=[[("出货：{ship_from}", 3), ("业务：{handler}", 2)],
        [("客户：{customer}", 3), ("订单号：{contract_no}", 2)]],
  right=["№：{order_no}", "(条码位)", "送货日期：{delivery_date}"],
  total_label="合计人民币(大写)：", sums=["金额"], small_total=True,
  sigs=["客户签收：", "核收："],
  notes=["样本、单页特价进行中，名片截稿19:00，本公司不负责因送货延迟等相关引申之赔偿责任。"])

S("制版设计行业/送货单样式H237", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="送  货  单", border="grid", blank=9,
  side="第一联存根\n第二联客户\n第三联回单",
  info=[[("客户名称：{customer}", 3), ("业务员：{handler}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["送  货  单", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["备注：尊敬的客户，为了维护双方利益，请您在使用我司生产的印刷版辊进行印刷前，请认真审核图案及文字等是否与意愿相符，如发现问题，请查双方责任加我司责任我司将免费无偿改正重制，若贵司没有审核清楚进行印刷造成的损失我司不负任何损失赔偿责任。本凭证视同购销合同，客户签字后即告合同成立。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("制版设计行业/送货单样本H620", ["序号", "名称/型号", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", doc="销  售  单\nSALES LIST", head=ADDR_TEL, border="grid", blank=9,
  blank_note="以下空白", side="①白存根\n②红客户\n③黄回单",
  info=[[("提货单位：{customer}", 3), ("联系人：{contact}", 2)],
        [("联系地址：{address}", 3), ("联系电话：{phone}", 2)]],
  right=["销  售  单\nSALES LIST", "销售单号：{order_no}", "开单日期：{order_date}"],
  total_label="合计(大写)：", sums=["金额"], small_total=True,
  notes=["小写：￥{items.amount#sum}    订金：￥____    余额：￥{items.amount#sum}",
         "注明：*本据单为交易时使用，购方代表对以上货物的规格、数量、付清方式确认无误后签字认可，并在货物出库和签收时当面点清。",
         "      *本据单签字后有效，如有质量异议，请在收货后一周内提出，过期概不负责。"],
  sigs=["制单人：{handler}", "提货车号：{vehicle_no}", "客户签字："])

S("制版设计行业/送货单格式H228", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX制版送货单", head=[("电话：XXX", False)], border="grid", blank=9,
  side="①白存根\n②红客户\n③黄结算",
  info=[[("购货单位：{customer}", 5)]],
  right=["NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["金额"], small_total=True,
  notes=["备注：尊敬的客户，为了维护双方利益，请您在使用我厂生产的版辊进行印刷前，请认真审核图案、文字等是否与意愿相符，如发现问题，请查阅双方责，如我厂责任我厂将免费修正并重制；若没有审稿进行印刷我厂将不负任何损失与责任。本凭证视同购销合同，客户签字后即合同关系成立。"],
  sigs=["仓库(制单)：{handler}", "送货人：", "验收人：", "收货人："])

S("制版设计行业/送货单格式H515", ["品    名", "客户PO", "规    格", "单位", "数量", "单价", "金额(十万千百拾元角分)"],
  title="XXX制版厂送货单", head=[("电话：XXX", False)], border="grid", blank=9,
  info=[[("客    户：{customer}", 5)]],
  right=["NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计大写：", sums=["金额"], small_total=True,
  notes=["注：客户收货后请即校稿，如有错漏，请在三天内提出更改补做，过期错漏自负，我方概不承担其它任何责任。",
         "[一]请款(白)  [二]送货签收(红)  [三]自存(黄)  [四]留底(绿)"],
  sigs=["经手人：{handler}", "收货单位盖章："], page_no=True)

# ============================================================ 食品行业
S("食品行业/送货单样式H198", ["产品编号", "产品名称", "规格", "单位", "生产批号", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=9,
  side="①白存根\n②红客户\n③黄回单",
  info=[[("客户名称：{customer}", 3), ("送货日期：{delivery_date}", 3)]],
  right=["[送  货  单]", "单据号：{order_no}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  remark_line=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("食品行业/送货单样式H286", ["产品编号", "品 名 规 格", "单位", "数量", "单价", "金额(万千百拾元角分)", "备注"],
  title="XXX食品厂", head=ADDR_TEL, doc="送  货  单", border="grid", blank=6,
  side="白：存根\n红：回单\n黄：客户\n兰：仓库",
  info=[[("顾客名称：{customer}", 5)]],
  right=["送  货  单", "NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["金额"], small_total=True,
  notes=["结算方式及时间：{pay_method}"],
  sigs=["收货单位(签章)：", "制单：{handler}", "仓库：", "经手人："])

S("食品行业/送货单样式H381", ["产品名称", "规格", "单位", "批号", "有效期", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=6,
  info=[[("客户名称：{customer}", 4)], [("客户地址：{address}", 4)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。",
         "第一联 存根        第二联 结款        第三联 客户"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("食品行业/送货单样式H382", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="[送  货  单]", border="grid", blank=2,
  side="①存根\n②客户\n③回单",
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 5)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("食品行业/送货单样式H503", ["产 品 名 称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX食品厂送货单", border="grid", blank=0,
  info=[[("购买单位：{customer}", 4), ("NO.：{order_no}", 2)],
        [("地址：{address}", 4), ("日期：{delivery_date}", 2)]],
  right=[],
  total_cn=False, total_label="合        计", sums=["数量", "金额"], small_total=True,
  sigs=["开单：{handler}", "送货人：", "收货人："],
  notes=["白联存根  红联回单  蓝联客户  黄联仓库"])

S("食品行业/送货单样式H585", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司送货单", head=ADDR_TEL, border="grid", blank=12,
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 5)]],
  right=["NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("食品行业/送货单样本H092", ["产品编号", "产品名称", "单位", "数量", "单价", "金额", "备注"],
  title="XXX食品配销中心", title_red=True, head=[("订购热线：XXX", True)], doc="[送  货  单]",
  border="grid", blank=11,
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "订单号：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  remark_line=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["收货单位及经手人(签字)："])

S("食品行业/食品行业送货单样本H037", ["条形码", "产品名称及规格", "规格", "数量", "箱价", "单价", "金额", "保质期", "生产日期", "备注"],
  title="XXX食品销售单", border="grid", blank=0,
  info=[[("购货单位：{customer}", 2), ("电话：{phone}", 2), ("NO.：{order_no}", 2)],
        [("联系人：{contact}", 2), ("地址：{address}", 2), ("日期：{delivery_date}", 2)]],
  right=["第1页，共1页"],
  total_cn=False, total_label="合计", sums=["数量"], small_total=False,
  notes=["本单应收：{items.amount#sum}",
         "公司地址：XXX    送货地址：{address}",
         "开单人员：{handler}    收货人签名：        欠款人签名：",
         "送货人电话：        传真：XXX        应收小计：{items.amount#sum}",
         "以上商品均已履行进货检查验收法定程序，索验票证齐全，供货者特此声明。"])

# ============================================================ 塑胶行业
S("塑胶行业/送货单样式H077", ["序号", "订单编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送  货  单", border="grid", blank=5, blank_note="(以下空白)",
  side="①白存根\n②红客户\n③黄回单\n④黄合计",
  info=[[("客户名称：{customer}", 4)]],
  right=["送  货  单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["收货单位及经手人(盖章)：", "送货单位及经手人(盖章)：", "制单人员：{handler}"])

S("塑胶行业/送货单样式H209", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX塑料厂送货单", border="grid", blank=2,
  side="①白存根\n②红凭证\n③黄记账",
  info=[[("收货单位：{customer}", 5)]],
  right=["NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["送货人：{handler}", "收货单位经手人："])

S("塑胶行业/送货单样式H298", ["序号", "商品名称及规格", "单位", "数量", "单价", "金额(十万千百拾元角分)", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送  货  单", border="grid", blank=8,
  side="①存根白\n②客户红\n③请款黄\n④财务蓝",
  info=[[("客    户：{customer}", 5)], [("订单号：{contract_no}", 5)]],
  right=["送  货  单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计大写：", sums=["金额"], small_total=True,
  notes=["注：如有质量问题，十日之内致电告知，逾期恕不负责！"],
  sigs=["客户签收及盖章：", "制单人：{handler}"])

S("塑胶行业/送货单样式H403", ["订单编号", "产品编号", "品名及规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="送  货  单", border="grid", blank=9,
  side="①存根\n②客户\n③回单\n④仓库",
  info=[[("客户名称：{customer}", 5)], [("客户电话：{phone}", 5)]],
  right=["送  货  单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("塑胶行业/送货单样本H267", ["序号", "规格", "产品名称", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=9,
  side="一存根(白)\n二回单(黄)\n三客户(红)\n四对帐(兰)",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("塑胶行业/送货单样本H272", ["订单编号", "产品名称", "产品型号", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送  货  单", border="grid", blank=8,
  side="①白存根\n②红客户\n③蓝回单\n④黄财务",
  info=[[("客户名称：{customer}", 4)]],
  right=["送  货  单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("塑胶行业/送货单样本H542", ["产品名称\nDESCRIPTION", "规格\nSPEC", "单位\nUNIT", "数量\nQuantity", "单价\nUnit Price", "金额\nAmount", "备注\nRemark"],
  title="XXX公司", head=[("XXX Co., Ltd.", False), ("地址：XXX  电话：XXX", False)],
  doc="[送  货  单]", border="grid", blank=8,
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 5)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["发货单位盖章：", "发货单位经手人：{handler}", "收货单位经手人："])

S("塑胶行业/送货单样本H592", ["序号", "品名规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=[("地址：XXX", False), ("Tel：XXX", False)], doc="[送  货  单]",
  border="grid", blank=9,
  side="①白财务\n②红请款\n③蓝存根\n④黄客户",
  info=[[("订单号码：{contract_no}", 3), ("制单人员：{handler}", 2)],
        [("付款方式：{pay_method}", 3), ("客户名称：{customer}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("塑胶行业/送货单格式H210", ["订  单  号\nFO:NO", "产品名称\nDESCRIPTION", "规格\nSPEC", "数量\nQuantity", "单价\nUnit Price", "总金额\nAmount", "备  注\nRemark"],
  title="XXX公司", head=[("XXX Plastic&Industry co.,ltd", False)], doc="送  货  单", doc_center=True,
  border="grid", blank=7,
  side="①存根\n②顾款\n③财务\n④回单",
  info=[[("客户名称：{customer}", 3), ("送货车号：{vehicle_no}", 3)]],
  right=["NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：收货方如发现质量问题请于七天内书面向供方提出，否则供货方一律不承担责任。"],
  sigs=["经手人：{handler}", "收货人(签章)："])

# ============================================================ 机械制造行业
S("机械制造行业/送货单样式H323", ["产品名称及规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", doc="送 货 验 收 单", doc_center=True, border="grid", blank=3,
  info=[[("送货人：{handler}", 2), ("送货日期：{delivery_date}", 2), ("编号：{order_no}", 2)],
        [("收货单位：{customer}", 5)],
        [("收货人：{contact}", 2), ("联系电话：{phone}", 2)],
        [("收货地址：{address}", 5)]],
  right=[],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["验收人签名(盖章)：        日期：{delivery_date}",
         "制单：{handler}  地址：XXX  电话：XXX",
         "送货单说明：白色联：存根联  蓝色联：客户联  红色联：结算联  黄色联：仓库联"])

S("机械制造行业/送货单样式H346", ["货品编号", "品名/规格", "数量", "单价", "金额"],
  title="XXX公司", doc="出 货 单", head=[("TEL：XXX", False)], border="grid", blank=3,
  info=[[("订单编号：{contract_no}", 5)],
        [("客户名称：{customer}", 3), ("地址：{address}", 2)],
        [("电话：{phone}", 3), ("联系人：{contact}", 2)],
        [("币别：人民币", 2), ("第一联：存根(白)  第二联：请款(红)  第三联：厂商(黄)", 3)]],
  right=["出 货 单", "本单号码：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=["金额"], small_total=False,
  notes=["本交易为附条件买卖，卖公司占有使用权，本货品尚未付给或兑现全部货款前，货款本公司仍保有货品之所有权，如有违约，本公司不经催告进行解除合约，取回货品并请求赔偿所有之损失。",
         "销项税额：____    合计金额：{items.amount#sum}"],
  sigs=["客户签收：", "承运人：", "仓库：", "核准：", "复核：", "经办：{handler}"])

S("机械制造行业/送货单样式H527", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=7, blank_note="-以下空白-",
  side="①白存根\n②红回单\n③黄客户",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)],
        [("客户地址：{address}", 3), ("客户电话：{phone}", 2)]],
  right=["[送  货  单]", "№：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  remark_line=True,
  notes=["注：以上货品请及时核对数量，如有质量问题，请在收货后3天内通知本公司，逾期本公司恕不负责。"],
  sigs=["收货单位及经手人(盖章)：", "送货单位及经手人(盖章)："])

S("机械制造行业/送货单样本H378", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=9,
  side="①白存根\n②红客户\n③蓝回单",
  info=[[("客户名称：{customer}", 3), ("联系电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("机械制造行业/送货单样本H386", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=9,
  info=[[("客户名称：{customer}", 3), ("联系电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["发货单位经手人：{handler}", "收货单位经手人："],
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"])

S("机械制造行业/送货单样本H390", ["产品名称", "规格", "单位", "数量", "单价", "金额", "件数", "备注"],
  title="XXX公司", head=[("Tel：XXX", False)], doc="[送  货  单]", border="grid", blank=1,
  side="白存根\n红记账\n蓝回单\n黄客户",
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 5)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：{handler}", "收货单位及经手人(盖章)："])

S("机械制造行业/送货单样本H522", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=1,
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["数量"], small_total=False,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("机械制造行业/送货单格式H259", ["产品名称", "适用车型", "类型", "品牌", "单位", "数量", "单价", "金额"],
  title="XXX汽配", title_red=True, head=ADDR_TEL_RED, doc="[送  货  单]", border="grid", blank=1,
  info=[[("客户名称：{customer}", 5)], [("客户地址：{address}", 3), ("客户电话：{phone}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["货品请核对数量！非人为质量问题，不影响二次销售，可退可换。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("机械制造行业/送货单格式H273", ["产品名称", "规格", "单位", "数量", "单价", "金额", "订购单号", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="送  货  单", border="grid", blank=8,
  side="①白存根\n②红回单\n③黄客户",
  info=[[("客户名称：{customer}", 4)], [("客户地址：{address}", 4)]],
  right=["送  货  单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司。"],
  sigs=["制单人员：{handler}", "送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("机械制造行业/送货单格式H287", ["产品编号", "产品名称", "规格", "单位", "数量", "箱数", "单价", "金额", "备注"],
  title="XXX公司", head=[("XXX Co., Ltd.", False)], doc="[送  货  单]", border="grid", blank=0,
  info=[[("客户名称：{customer}", 5)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["①存根(白)  ②客户(红)  ③回单(黄)"],
  sigs=["送货人：", "制单人：{handler}", "收货单位：", "收货人："])

S("机械制造行业/送货单格式H291", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", doc="收款专用票据", border="dash", blank=6,
  info=[[("客户名称：{customer}", 5)]],
  right=["收款专用票据", "NO.：{order_no}", "日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["制单人员：{handler}", "收货单位签名(盖章)："],
  notes=["1.存根联(白)  2.客户联(红)"])

S("机械制造行业/送货单格式H343", ["订单号", "零件编号", "品名规格", "单位", "数量", "箱数", "备注"],
  title="XXX公司", head=[("Tel：XXX", False)], doc="[送  货  单]", border="grid", blank=10,
  side="①白存根\n②红回单\n③黄客户",
  info=[[("客户名称：{customer}", 3), ("联系人：{contact}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, sums=[], small_total=False,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

# ============================================================ 纸品纸箱行业
S("纸品纸箱行业/送货单样式H240", ["编號\nNumber", "尺码\nDimensions", "款式\nModel", "数量\nQuantity", "單價\nUnit Price", "金額\nAmount"],
  title="XXX纸品有限公司", head=ADDR_TEL, doc="送 貨 單\nDelivery Order", doc_center=True, border="grid", blank=5,
  side="①請款白\n②客户紅\n③對帳黃\n④會計綠\n⑤存根藍",
  info=[[("客户：{customer}", 4)]],
  right=["送 貨 單\nDelivery Order", "NO.：{order_no}", "Date：{delivery_date}"],
  total_cn=False, total_label="收款時需另發正式收據    合計 TOTAL RMB￥", sums=["金额"], small_total=True,
  sigs=["收貨人簽名蓋章：", "制單人：{handler}"])

S("纸品纸箱行业/送货单样式H355", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX纸巾厂", head=ADDR_TEL, doc="送货明细单", border="grid", blank=8,
  side="红色联为付款凭证",
  info=[[("购货单位电话：{phone}", 3)], [("购货单位：{customer}", 3)], [("购货单位地址：{address}", 3)]],
  right=["送货明细单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["以上产品数量及质量验收合格"],
  sigs=["验收单位经办人：", "送货单位及经手人：{handler}"])

S("纸品纸箱行业/送货单样式H558", ["商品名称及规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX纸箱厂", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=10,
  info=[[("客户名称：{customer}", 4)], [("客户地址：{address}", 2), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("纸品纸箱行业/送货单样张H402", ["编號", "尺码", "款式", "数量", "不含稅單價", "金額", "稅金"],
  title="XXX纸品有限公司", head=ADDR_TEL, doc="送 貨 單\nDelivery Order", doc_center=True, border="grid", blank=5,
  side="①請款白\n②客户紅\n③對帳黃\n④會計綠\n⑤存根藍",
  info=[[("客户：{customer}", 4)]],
  right=["送 貨 單\nDelivery Order", "NO.：{order_no}", "Date：{delivery_date}"],
  total_cn=False, total_label="收款時需另發正式收據    合計 TOTAL RMB￥", sums=["金额", "稅金"], small_total=False,
  sigs=["收貨人簽名蓋章：", "制單人：{handler}"])

S("纸品纸箱行业/送货单样本H258", ["订单编号", "料号", "品名规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=8,
  info=[[("客户名称：{customer}", 5)], [("客户电话：{phone}", 3)], [("客户地址：{address}", 3)]],
  right=["[送  货  单]", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_cn=False, total_label="合计", sums=["金额"], small_total=False,
  notes=["以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。",
         "第一联白  存根        第二联红  客户        第三联黄  仓库        第四联绿  会计"],
  sigs=["收货单位签名：", "送货人签名：", "制单人：{handler}"])

S("纸品纸箱行业/送货单格式H233", ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX纸业", title_red=True, doc="送  货  单", border="grid", blank=9,
  side="第一联存根\n第二联客户\n第三联回单",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2)],
        [("客户地址：{address}", 3), ("送货日期：{delivery_date}", 2)]],
  right=["送  货  单", "NO.：{order_no}"],
  total_label="金额合计(大写)：", sums=["金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("纸品纸箱行业/送货单格式H236", ["订单编号", "产品名称", "规格", "数量", "单价", "金额", "重量(KG)"],
  title="XXX包装有限公司", head=[("電話：XXX", False)], doc="送  貨  單\nDELIVERY NOTE", doc_center=True,
  border="grid", blank=8,
  side="①請款白\n②客户紅\n③對帳黃\n④會計綠\n⑤存根藍",
  info=[[("客户：{customer}", 4)]],
  right=["送  貨  單\nDELIVERY NOTE", "NO.：{order_no}", "Date：{delivery_date}"],
  total_cn=False, total_label="※收款時需另發正式收據    總金額HKD", sums=["金额"], small_total=True,
  sigs=["簽收蓋章：", "制表：{handler}"])

# ============================================================ 其它行业格式
S("其它行业格式/送货单格式H598", ["产品名称", "数量", "单位", "单价", "金额", "备注"],
  title="XXX公司", head=[("Tel：XXX", False)], doc="【送货单】", border="dash", blank=3,
  info=[[("客户：{customer}", 3), ("联系方式：{phone}", 2), ("结帐方式：{pay_method}", 1)],
        [("送货地址：{address}", 3), ("送货车号：{vehicle_no}", 2), ("送货日期：{delivery_date}", 1)]],
  right=[],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["发货单位经手人：{handler}        收货单位签名(盖章)：",
         "备注：1.客户签字表示购货双方权利义务已确认。",
         "      2.白联存根；红联收款；蓝联和黄联客户。"])

S("其它行业格式/送货单格式H607", ["订单号码", "品名 规格", "颜 色", "数量", "单价", "金额", "备注"],
  title="XXX公司送货单", head=[("地址：XXX  电话：XXX", False)], border="grid", blank=10,
  info=[[("客户名称：{customer}", 3), ("日期：{delivery_date}", 2), ("单号：{order_no}", 2)]],
  right=[],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["声明：本单所送之货物如有色差、损伤或任何异常，请退回我司更换，若经生产加工后发现异常本公司不负任何责任。",
         "①白联存根  ②红联回执  ③黄联客户  ④绿联仓库"],
  sigs=["制单：{handler}", "送货人：", "收货单位签章："])

S("其它行业格式/送货单格式H678", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="【送  货  单】", border="grid", blank=10,
  info=[[("客户名称：{customer}", 3), ("联系电话：{phone}", 2), ("合同号码：{contract_no}", 1)],
        [("客户地址：{address}", 3), ("发货方式：{ship_method}", 2), ("发货日期：{delivery_date}", 1)]],
  right=["【送  货  单】", "NO.：{order_no}"],
  total_label="金额合计(大写)：", sums=["数量", "金额"], small_total=True,
  sigs=["核准：", "财务：", "发货：", "制单：{handler}"],
  notes=["第一联：存根        第二联：客户        第三联：财务        第四联：仓库"])

S("其它行业格式/送货单格式H679", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=2,
  side="①存根白\n②客户红\n③回单黄",
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}", "制单人员：{handler}", "送货日期：{delivery_date}"],
  total_label="合计金额(大写)：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("其它行业格式/送货单格式H893", ["序号", "产品名称", "规格", "颜色", "数量", "单价", "金额", "备注"],
  title="XXX公司", doc="送  货  单", border="grid", blank=6,
  side="①存根\n②客户\n③回单",
  info=[[("客户：{customer}", 3), ("联系人：{contact}", 2), ("电话：{phone}", 1)],
        [("地址：{address}", 3), ("日期：{delivery_date}", 3)]],
  right=["送  货  单", "№：{order_no}"],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责！",
         "地址：XXX    电话：XXX    传真：XXX"],
  sigs=["制表：{handler}", "客户签收："])

S("其它行业格式/送货单格式T011", ["序号", "货号", "货品名称及规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", doc="送  货  单", border="grid", blank=7,
  side="①白存根\n②红客户\n③黄回单",
  info=[[("收货单位：{customer}", 3), ("合同号：{contract_no}", 2), ("单号：{order_no}", 2)],
        [("收货地址：{address}", 3), ("联系人：{contact}", 2), ("日期：{delivery_date}", 2)]],
  right=[],
  total_label="合计 金额", sums=["金额"], small_total=True,
  notes=["备注：以上货品请核对数量，如有质量问题，请在收货后7天内通知本公司，逾期恕不负责。"],
  sigs=["收货单位及经手人(盖章)：", "送货单位及经手人(盖章)："])

S("其它行业格式/送货单格式T026", ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", title_red=True, head=ADDR_TEL_RED, doc="[送  货  单]", border="grid", blank=6,
  info=[[("客户名称：{customer}", 3), ("订单号码：{contract_no}", 2), ("制单人员：{handler}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2), ("送货日期：{delivery_date}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  notes=["注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。"],
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("其它行业格式/送货单格式T035", ["商品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]\n(条码位)", border="grid", blank=9,
  info=[[("客户名称：{customer}", 3), ("合同号码：{contract_no}", 2), ("制单人员：{handler}", 2)],
        [("客户地址：{address}", 3), ("送货日期：{delivery_date}", 2)]],
  right=["[送  货  单]\n(条码位)", "NO.：{order_no}"],
  total_label="合计：", sums=["数量", "金额"], small_total=True,
  sigs=["送货单位及经手人(盖章)：", "收货单位及经手人(盖章)："])

S("其它行业格式/送货单格式T037", ["订单编号", "产品编号", "名称及规格", "单位", "数量", "单价", "金额"],
  title="XXX公司", head=ADDR_TEL, doc="送  货  单", border="grid", blank=9,
  side="①存根(白)\n②客户(红)\n③仓库(蓝)\n④存根(蓝)",
  info=[[("收货单位：{customer}", 4)]],
  right=["送  货  单", "NO.：{order_no}", "送货日期：{delivery_date}"],
  total_label="会计 人民币", sums=["金额"], small_total=True,
  sigs=["收货单位及经手人(盖章)：", "送货单位及经手人："])

S("其它行业格式/送货单格式T055", ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="[送  货  单]", border="grid", blank=10,
  info=[[("订单号码：{contract_no}", 3), ("制单人员：{handler}", 2)],
        [("客户地址：{address}", 3), ("送货车号：{vehicle_no}", 2), ("送货日期：{delivery_date}", 2)]],
  right=["[送  货  单]", "NO.：{order_no}"],
  total_label="合计金额(大写)：", sums=["金额"], small_total=True,
  sigs=["收货单位及经手人(盖章)：", "送货单位及经手人(盖章)："])

S("其它行业格式/送货单格式T060", ["序号", "产品名称", "规格", "颜色", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="定货合同单", border="grid", blank=6,
  side="白色存根联\n蓝色送货联\n红色顾客联\n黄色财务联",
  info=[[("客户名称：{customer}", 3), ("客户电话：{phone}", 2), ("送货日期：{delivery_date}", 2)]],
  right=["定货合同单", "NO.：{order_no}"],
  total_label="合计：", sums=["金额"], small_total=True,
  notes=["1、购货方验收产品质量，应收我厂验收为准。",
         "2、在运输途中碰坏家具，我厂一概不承担责任。",
         "3、自双方签定合同之日起生效。",
         "4、到取货日期后十天不取货者，订金作废。",
         "其它事项："],
  sigs=["主管：", "购方经手人：", "卖方经手人："])

S("其它行业格式/送货单格式T100", ["序号", "名称/型号", "规格", "单位", "数量", "单价", "金额", "备注"],
  title="XXX公司", head=ADDR_TEL, doc="销  售  单\nSALES LIST", border="grid", blank=8, blank_note="以下空白",
  side="①白存根\n②红客户\n③黄回单",
  info=[[("提货单位：{customer}", 3), ("联系人：{contact}", 2)],
        [("联系地址：{address}", 3), ("联系电话：{phone}", 2)]],
  right=["销  售  单\nSALES LIST", "销售单号：{order_no}", "开单日期：{order_date}"],
  total_label="合计(大写)：", sums=["金额"], small_total=True,
  notes=["注明：*本据单为交易时使用，购方代表对以上货物的规格、数量、付款方式确认无误后签字认可，并在货物出库和签收时当面点清。",
         "      *本据单签字后有效，如有质量异议，请在收货后一周内提出，过期概不负责。"],
  sigs=["制单人：{handler}", "提货车号：{vehicle_no}", "客户签字："])

if __name__ == "__main__":
    import sys

    only = sys.argv[1] if len(sys.argv) > 1 else None
    n = 0
    for spec in SPECS:
        if only and only not in spec["f"]:
            continue
        build(spec)
        n += 1
    print(f"生成 {n} 个模板 → {OUT}")
