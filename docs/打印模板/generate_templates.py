# -*- coding: utf-8 -*-
"""按 docs/工单打印截图 复刻 16 个送货单 Excel 打印模板（本系统 Excel 模板引擎的占位符语法）。

占位符：
  主表字段  {order_no} {customer} {address} {contact} {phone} {contract_no} {delivery_date}
            {order_date} {vehicle_no} {pay_method} {ship_method} {handler} {remark}
  内置      {_id} {_today} {_total_cn}（人民币大写）
  明细循环  {items.product_no} {items.order_no} {items.name} {items.spec} {items.color} {items.material}
            {items.unit} {items.qty} {items.price} {items.amount} {items.remark} {items._index}
  明细合计  {items.qty#sum} {items.amount#sum}

注意：上传到自己的数据表前，把占位符里的字段名改成你表里的实际字段名（中文占位符不会被替换）。
公司名/地址/电话为截图中的静态示例文字，请在 Excel/WPS 里替换成自己的。
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent

THIN = Side(style="thin")
B = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DASH = Side(style="dashed")
B_DASH = Border(bottom=DASH)
SONG = "宋体"

CN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RT = Alignment(horizontal="right", vertical="center")
VERT = Alignment(horizontal="center", vertical="center", text_rotation=255, wrap_text=True)


def F(sz=11, bold=False, color=None):
    return Font(name=SONG, size=sz, bold=bold, color=color)


def W(ws, r, c, text, font=None, align=None, border=None):
    cell = ws.cell(r, c, text)
    cell.font = font or F()
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    return cell


def M(ws, r, c1, c2, text, font=None, align=None, border=None):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = W(ws, r, c1, text, font, align)
    if border:
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = border
    return cell


def widths(ws, wl):
    for i, w in enumerate(wl, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def head_row(ws, r, labels, font=None):
    for i, t in enumerate(labels, 1):
        W(ws, r, i, t, font or F(11, True), CN, B)


def loop_row(ws, r, tokens):
    for i, t in enumerate(tokens, 1):
        W(ws, r, i, t, F(), CN, B)


def blank_rows(ws, r0, n, ncols, border=None, dash=False):
    for r in range(r0, r0 + n):
        for c in range(1, ncols + 1):
            ws.cell(r, c).border = B_DASH if dash else (border or B)


def copies_note(ws, r0, r1, col, text):
    """右缘竖排联注（textRotation=255 直排）。"""
    ws.merge_cells(start_row=r0, start_column=col, end_row=r1, end_column=col)
    W(ws, r0, col, text, F(9), VERT)
    ws.column_dimensions[get_column_letter(col)].width = 3.5


def save(ws, name):
    wb = ws.parent
    wb.save(OUT / name)
    print("生成", name)


def new():
    wb = Workbook()
    return wb.active


# ---------------- H228 厦门太豪电子制版送货单 ----------------
def h228():
    ws = new()
    ws.title = "送货单"
    widths(ws, [11, 16, 10, 7, 7, 10, 10, 12, 3.5])
    M(ws, 1, 1, 8, "厦门太豪电子制版送货单", F(16, True), CN)
    M(ws, 2, 1, 5, "电话：(0)13652500551  传真：0769-83608888", F(10), CN)
    M(ws, 2, 6, 8, "NO.：{order_no}", F(11, True), RT)
    M(ws, 3, 1, 5, "购货单位：{customer}", F(), LT)
    M(ws, 3, 6, 8, "{delivery_date}", F(), RT)
    head_row(ws, 4, ["产品编号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 5, ["{items.product_no}", "{items.name}", "{items.spec}", "{items.unit}",
                     "{items.qty}", "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 6, 7, 8)
    M(ws, 13, 1, 4, "合计(大写)：{_total_cn}", F(), LT, B)
    M(ws, 13, 5, 6, "合计(小写)", F(), CN, B)
    W(ws, 13, 7, "{items.amount#sum}", F(), RT, B)
    W(ws, 13, 8, "", F(), CN, B)
    M(ws, 14, 1, 8, "备注：尊敬的客户，为了维护双方利益，请您在使用我厂生产的版辊进行印刷前，请认真审核图案、文字等是否与原稿相符，如发现问题，请查阅双方责，如我厂责任我厂将免费修正并重制；若没有审稿进行印刷我厂将不负任何损失与责任。本凭证视同购销合同，客户签字后即合同关系成立。", F(9), LT)
    ws.row_dimensions[14].height = 40
    M(ws, 15, 1, 2, "仓库(制单)：{handler}", F(), LT)
    M(ws, 15, 3, 4, "送货人：", F(), LT)
    M(ws, 15, 5, 6, "验收人：", F(), LT)
    M(ws, 15, 7, 8, "收货人：", F(), LT)
    copies_note(ws, 4, 12, 9, "①白存根②红客户③黄结算")
    save(ws, "H228-送货单(厦门太豪).xlsx")


# ---------------- H527 天津贝莱特送货单 ----------------
def h527():
    ws = new()
    widths(ws, [7, 16, 10, 7, 7, 10, 10, 12, 3.5])
    M(ws, 1, 1, 6, "天津贝莱特机械销售有限公司", F(16, True), CN)
    M(ws, 1, 7, 8, "[送  货  单]", F(13, True), CN)
    M(ws, 2, 1, 6, "广东省东莞市电子科技大厦A2015室", F(10), CN)
    M(ws, 2, 7, 8, "№：{order_no}", F(11, True), CN)
    M(ws, 3, 1, 6, "Tel: (0)13652500551  Fax: 0769-83608888", F(10), CN)
    M(ws, 4, 1, 3, "客户名称：{customer}", F(), LT)
    M(ws, 4, 4, 5, "联系人：{contact}", F(), LT)
    M(ws, 4, 6, 8, "制单人员：{handler}", F(), LT)
    M(ws, 5, 1, 3, "客户地址：{address}", F(), LT)
    M(ws, 5, 4, 5, "客户电话：{phone}", F(), LT)
    M(ws, 5, 6, 8, "送货日期：{delivery_date}", F(), LT)
    head_row(ws, 6, ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 7, ["{items._index}", "{items.name}", "{items.spec}", "{items.unit}",
                     "{items.qty}", "{items.price}", "{items.amount}", "{items.remark}"])
    M(ws, 8, 1, 8, "-以下空白-", F(), CN, B)
    blank_rows(ws, 9, 5, 8)
    M(ws, 14, 1, 5, "合计：{_total_cn}", F(), LT, B)
    W(ws, 14, 6, "{items.qty#sum}", F(), CN, B)
    W(ws, 14, 7, "{items.amount#sum}", F(), RT, B)
    W(ws, 14, 8, "", F(), CN, B)
    M(ws, 15, 1, 8, "备注：{remark}", F(), LT, B)
    M(ws, 16, 1, 8, "注：以上货品请及时核对数量，如有质量问题，请在收货后3天内通知本公司，逾期本公司恕不负责。", F(9), LT)
    M(ws, 17, 1, 4, "收货单位及经手人(盖章)：", F(), LT)
    M(ws, 17, 5, 8, "送货单位及经手人(盖章)：", F(), LT)
    copies_note(ws, 6, 14, 9, "①白存根②红回单③黄客户")
    save(ws, "H527-送货单(天津贝莱特).xlsx")


# ---------------- H535 南京夏平建材厂（虚线式） ----------------
def h535():
    ws = new()
    widths(ws, [16, 11, 8, 8, 10, 11, 12])
    M(ws, 1, 1, 5, "南京夏平建材厂", F(18, True), CN)
    M(ws, 1, 6, 7, "送  货  单", F(13, True), RT)
    M(ws, 2, 1, 5, "广东省东莞市电子科技大厦A2015室", F(9), CN)
    M(ws, 2, 6, 7, "NO：{order_no}", F(11, True), RT)
    M(ws, 3, 1, 5, "Tel: (0)13652500551  Fax: 0769-83608888", F(9), CN)
    M(ws, 3, 6, 7, "送货日期：{delivery_date}", F(), RT)
    M(ws, 4, 1, 3, "客户名称：{customer}", F(), LT)
    M(ws, 4, 5, 7, "客户电话：{phone}", F(), LT)
    M(ws, 5, 1, 7, "客户地址：{address}", F(), LT)
    for i, t in enumerate(["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"], 1):
        c = W(ws, 6, i, t, F(11, True), CN)
        c.border = Border(bottom=THIN)
    loop = ["{items.name}", "{items.spec}", "{items.unit}", "{items.qty}",
            "{items.price}", "{items.amount}", "{items.remark}"]
    for i, t in enumerate(loop, 1):
        W(ws, 7, i, t, F(), CN, B_DASH)
    blank_rows(ws, 8, 2, 7, dash=True)
    M(ws, 10, 1, 4, "合计：{_total_cn}", F(), LT, Border(top=DASH, bottom=THIN))
    W(ws, 10, 5, "{items.qty#sum}", F(), CN, Border(top=DASH, bottom=THIN))
    W(ws, 10, 6, "{items.amount#sum}", F(), RT, Border(top=DASH, bottom=THIN))
    W(ws, 10, 7, "", F(), CN, Border(top=DASH, bottom=THIN))
    M(ws, 11, 1, 2, "制单人员：{handler}", F(), LT)
    M(ws, 11, 3, 4, "送货人员：", F(), LT)
    M(ws, 11, 5, 7, "收货单位签名(盖章)：", F(), LT)
    M(ws, 12, 1, 7, "1. 存根联(白)  2. 客户联(红)  3. 回单联(黄)", F(9), LT)
    save(ws, "H535-送货单(南京夏平).xlsx")


# ---------------- H598 佛山澜石沙石场（紧凑式） ----------------
def h598():
    ws = new()
    widths(ws, [18, 8, 8, 10, 11, 12])
    M(ws, 1, 1, 4, "佛山市禅城区澜石明窦防汛沙石场", F(15, True), LT)
    M(ws, 1, 5, 6, "【送货单】", F(13, True), RT)
    M(ws, 2, 1, 3, "Tel: (0)13652500551", F(10), LT)
    M(ws, 2, 4, 6, "NO.：{order_no}", F(11, True), RT)
    M(ws, 3, 1, 2, "客户：(CU0002){customer}", F(), LT)
    M(ws, 3, 3, 4, "联系方式：{phone}", F(), LT)
    M(ws, 3, 5, 6, "结帐方式：{pay_method}", F(), LT)
    M(ws, 4, 1, 2, "送货地址：{address}", F(), LT)
    M(ws, 4, 3, 4, "送货车号：{vehicle_no}", F(), LT)
    M(ws, 4, 5, 6, "送货日期：{delivery_date}", F(), LT)
    for i, t in enumerate(["产品名称", "数量", "单位", "单价", "金额", "备注"], 1):
        c = W(ws, 5, i, t, F(11, True), CN)
        c.border = Border(bottom=THIN)
    for i, t in enumerate(["{items.name}", "{items.qty}", "{items.unit}", "{items.price}",
                           "{items.amount}", "{items.remark}"], 1):
        W(ws, 6, i, t, F(), CN, B_DASH)
    M(ws, 7, 1, 4, "合计：{_total_cn}", F(), LT, Border(top=DASH, bottom=THIN))
    W(ws, 7, 5, "{items.amount#sum}", F(), RT, Border(top=DASH, bottom=THIN))
    W(ws, 7, 6, "", F(), CN, Border(top=DASH, bottom=THIN))
    M(ws, 8, 1, 3, "发货单位经手人：{handler}", F(), LT)
    M(ws, 8, 4, 6, "收货单位签名(盖章)：", F(), LT)
    M(ws, 9, 1, 6, "备注：1. 客户签字表示购货双方权利义务已确认。\n2. 白联存根；红联收款；蓝联和黄联客户。", F(9), LT)
    ws.row_dimensions[9].height = 28
    save(ws, "H598-送货单(佛山澜石).xlsx")


# ---------------- H607 深圳天源线厂送货单 ----------------
def h607():
    ws = new()
    widths(ws, [11, 14, 9, 8, 10, 11, 11])
    M(ws, 1, 1, 7, "深圳天源线厂送货单", F(16, True), CN)
    M(ws, 2, 1, 7, "地址：广东省东莞市电子科技大厦A2015室  电话：(0)13652500551  传真：0769-83608888", F(9), CN)
    M(ws, 3, 1, 3, "客户名称：{customer}", F(11, True), LT)
    M(ws, 3, 4, 5, "日期：{delivery_date}", F(11, True), CN)
    M(ws, 3, 6, 7, "单号：{order_no}", F(11, True), RT)
    head_row(ws, 4, ["订单号码", "品名 规格", "颜 色", "数量", "单价", "金额", "备注"])
    loop_row(ws, 5, ["{items.order_no}", "{items.name}", "{items.color}", "{items.qty}",
                     "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 6, 9, 7)
    M(ws, 15, 1, 4, "合计：{_total_cn}", F(), LT, B)
    M(ws, 15, 5, 6, "", F(), CN, B)
    W(ws, 15, 7, "金额：{items.amount#sum}", F(), RT, B)
    M(ws, 16, 1, 7, "声明：本单所送之货物如有色差、损伤或任何异常，请退回我司更换，若经生产加工后发现异常本公司不负任何责任。", F(9), LT)
    M(ws, 17, 1, 2, "制单：{handler}", F(), LT)
    M(ws, 17, 3, 5, "送货人：", F(), LT)
    M(ws, 17, 6, 7, "收货单位签章：", F(), LT)
    M(ws, 18, 1, 7, "※※※※※※※※ ①白联存根  ②红联回执  ③黄联客户  ④绿联仓库 ※※※※※※※※", F(9), CN)
    save(ws, "H607-送货单(深圳天源).xlsx")


# ---------------- H678 宁波开世送货单 ----------------
def h678():
    ws = new()
    widths(ws, [13, 11, 7, 7, 10, 11, 12])
    M(ws, 1, 1, 5, "宁波开世密封科技有限公司", F(16, True), CN)
    M(ws, 1, 6, 7, "【送  货  单】", F(12, True), RT)
    M(ws, 2, 1, 5, "地址：东莞市樟木头镇樟罗先威大道", F(9), CN)
    M(ws, 2, 6, 7, "NO.：{order_no}", F(11, True), RT)
    M(ws, 3, 1, 5, "TEL: 0769-85969959  FAX: 0769-85932906", F(9), CN)
    M(ws, 4, 1, 3, "客户名称：{customer}", F(), LT)
    M(ws, 4, 4, 5, "联系电话：{phone}", F(), LT)
    M(ws, 4, 6, 7, "合同号码：{contract_no}", F(), LT)
    M(ws, 5, 1, 3, "客户地址：{address}", F(), LT)
    M(ws, 5, 4, 5, "发货方式：{ship_method}", F(), LT)
    M(ws, 5, 6, 7, "发货日期：{delivery_date}", F(), LT)
    head_row(ws, 6, ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 7, ["{items.name}", "{items.spec}", "{items.unit}", "{items.qty}",
                     "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 8, 10, 7)
    M(ws, 18, 1, 3, "金额合计(大写)：{_total_cn}", F(), LT, B)
    W(ws, 18, 4, "{items.qty#sum}", F(), CN, B)
    W(ws, 18, 5, "小写金额", F(), CN, B)
    W(ws, 18, 6, "{items.amount#sum}", F(), RT, B)
    W(ws, 18, 7, "", F(), CN, B)
    for i, t in enumerate(["核准：", "财务：", "发货：", "制单：{handler}"], 0):
        M(ws, 19, 1 + i * 2, 2 + i * 2 if i < 3 else 7, t, F(), LT)
    M(ws, 20, 1, 7, "第一联：存根    第二联：客户    第三联：财务    第四联：仓库", F(9), CN)
    save(ws, "H678-送货单(宁波开世).xlsx")


# ---------------- H679 佛山江云送货单 ----------------
def h679():
    ws = new()
    widths(ws, [13, 10, 7, 7, 10, 11, 12, 3.5])
    M(ws, 1, 1, 6, "佛山市江云汽车零部件有限公司", F(15, True), CN)
    M(ws, 1, 7, 7, "[送 货 单]", F(12, True), CN)
    M(ws, 2, 1, 6, "地址：东莞市樟木头镇樟罗先威大道", F(9), CN)
    M(ws, 2, 7, 7, "NO.：{order_no}", F(11, True), CN)
    M(ws, 3, 1, 6, "Tel: 0769-85969959  Fax: 0769-85932906", F(9), CN)
    M(ws, 4, 1, 3, "客户名称：{customer}", F(), LT)
    M(ws, 4, 4, 5, "合同号码：{contract_no}", F(), LT)
    M(ws, 4, 6, 7, "制单人员：{handler}", F(), LT)
    M(ws, 5, 1, 3, "客户地址：{address}", F(), LT)
    M(ws, 5, 4, 5, "送货车号：{vehicle_no}", F(), LT)
    M(ws, 5, 6, 7, "送货日期：{delivery_date}", F(), LT)
    head_row(ws, 6, ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 7, ["{items.name}", "{items.spec}", "{items.unit}", "{items.qty}",
                     "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 8, 4, 7)
    M(ws, 12, 1, 3, "合计金额(大写)：{_total_cn}", F(), LT, B)
    W(ws, 12, 4, "{items.qty#sum}", F(), CN, B)
    W(ws, 12, 5, "小写金额", F(), CN, B)
    W(ws, 12, 6, "{items.amount#sum}", F(), RT, B)
    W(ws, 12, 7, "", F(), CN, B)
    M(ws, 13, 1, 7, "注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。", F(9), LT)
    M(ws, 14, 1, 3, "送货单位及经手人(盖章)：", F(), LT)
    M(ws, 14, 5, 7, "收货单位及经手人(盖章)：", F(), LT)
    copies_note(ws, 6, 12, 8, "①存根(白)②客户(红)③回单(黄)")
    save(ws, "H679-送货单(佛山江云).xlsx")


# ---------------- H845 昆明浩宇接件及送货单 ----------------
def h845():
    ws = new()
    widths(ws, [12, 9, 11, 7, 6, 9, 10, 10, 3.5])
    M(ws, 1, 1, 7, "昆明浩宇图文设计有限公司", F(16, True), CN)
    M(ws, 1, 8, 8, "NO.：{order_no}", F(11, True), CN)
    M(ws, 2, 1, 7, "接 件 及 送 货 单", F(14, True), CN)
    M(ws, 3, 1, 4, "客户名称：{customer}", F(), LT)
    M(ws, 3, 6, 8, "接件日期：{order_date}", F(), LT)
    M(ws, 4, 1, 4, "联系电话：{phone}", F(), LT)
    M(ws, 4, 6, 8, "交货日期：{delivery_date}", F(), LT)
    head_row(ws, 5, ["项目名称", "规格", "材料", "数量", "单位", "单价", "金额", "备注"])
    loop_row(ws, 6, ["{items.name}", "{items.spec}", "{items.material}", "{items.qty}",
                     "{items.unit}", "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 7, 5, 8)
    M(ws, 12, 1, 6, "合计人民币(大写)：{_total_cn}", F(), LT, B)
    M(ws, 12, 7, 8, "￥：{items.amount#sum}", F(), RT, B)
    W(ws, 13, 1, "后期工艺", F(11, True), CN, B)
    M(ws, 13, 2, 8, "覆膜□  压痕□  模切□  压纹□  压凸□  烫金□  上UV□  上光油□  打孔□\n装订：锁线胶装□  无线胶装□  骑马订□  胶头□  胶边□  包本□  打号□", F(9), LT, B)
    ws.row_dimensions[13].height = 30
    for i, t in enumerate(["开单：{handler}", "核价：", "设计：", "送货：", "收款："], 1):
        W(ws, 14, 1 + (i - 1) * 2 if i < 5 else 8, t, F(), LT)
    M(ws, 15, 1, 8, "地址：深圳市宝安区西乡九围雍啟科技园B、C栋", F(9), LT)
    M(ws, 16, 1, 8, "电话：0755-29553163 29553182    传真：0755-29553097", F(9), LT)
    copies_note(ws, 5, 12, 9, "①存根 ②客户 ③回单")
    save(ws, "H845-接件及送货单(昆明浩宇).xlsx")


# ---------------- H893 上海滔业送货单 ----------------
def h893():
    ws = new()
    widths(ws, [6, 15, 10, 8, 7, 10, 11, 11, 3.5])
    M(ws, 1, 1, 7, "上海滔业化学科技有限公司", F(16, True), CN)
    M(ws, 2, 1, 7, "送  货  单", F(14, True), CN)
    M(ws, 2, 8, 8, "№：{order_no}", F(11, True), CN)
    M(ws, 3, 1, 3, "客户：{customer}", F(), LT)
    M(ws, 3, 4, 5, "联系人：{contact}", F(), LT)
    M(ws, 3, 6, 8, "电话：{phone}", F(), LT)
    M(ws, 4, 1, 3, "地址：{address}", F(), LT)
    M(ws, 4, 6, 8, "日期：{delivery_date}", F(), LT)
    head_row(ws, 5, ["序号", "产品名称", "规格", "颜色", "数量", "单价", "金额", "备注"])
    loop_row(ws, 6, ["{items._index}", "{items.name}", "{items.spec}", "{items.color}",
                     "{items.qty}", "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 7, 5, 8)
    M(ws, 12, 1, 4, "合计：{_total_cn}", F(), LT, B)
    M(ws, 12, 5, 6, "小写", F(), CN, B)
    W(ws, 12, 7, "{items.amount#sum}", F(), RT, B)
    W(ws, 12, 8, "", F(), CN, B)
    M(ws, 13, 1, 8, "注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责！", F(9), LT)
    M(ws, 14, 1, 4, "地址：广东省东莞市电子科技大厦A2015室", F(9), LT)
    M(ws, 14, 5, 8, "网址：http://hi.baidu.com/lhk165", F(9), LT)
    M(ws, 15, 1, 8, "电话：13652500551   传真：0769-83608888", F(9), LT)
    M(ws, 16, 1, 4, "制表：{handler}", F(11, True), LT)
    M(ws, 16, 5, 8, "客户签收：", F(11, True), LT)
    copies_note(ws, 5, 12, 9, "①存根 ②客户 ③回单")
    save(ws, "H893-送货单(上海滔业).xlsx")


# ---------------- T011 东莞一彩送货单 ----------------
def t011():
    ws = new()
    widths(ws, [6, 9, 20, 7, 7, 10, 11, 11, 3.5])
    M(ws, 1, 1, 8, "东莞市一彩科技有限公司", F(18, True), CN)
    M(ws, 2, 1, 8, "送  货  单", F(14, True), CN)
    M(ws, 3, 1, 3, "收货单位：{customer}", F(), LT)
    M(ws, 3, 4, 5, "合同号：{contract_no}", F(), LT)
    M(ws, 3, 6, 8, "单号：{order_no}", F(), LT)
    M(ws, 4, 1, 3, "收货地址：{address}", F(), LT)
    M(ws, 4, 4, 5, "联系人：{contact}", F(), LT)
    M(ws, 4, 6, 8, "日期：{delivery_date}", F(), LT)
    head_row(ws, 5, ["序号", "货号", "货品名称及规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 6, ["{items._index}", "{items.product_no}", "{items.name}", "{items.unit}",
                     "{items.qty}", "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 7, 6, 8)
    M(ws, 13, 1, 2, "合 计 金 额", F(11, True), CN, B)
    M(ws, 13, 3, 6, "{_total_cn}", F(11, True), LT, B)
    M(ws, 13, 7, 8, "{items.amount#sum}", F(11, True), RT, B)
    M(ws, 14, 1, 8, "备注：以上货品请核对数量，如有质量问题，请在收货后7天内通知本公司，逾期恕不负责。", F(9), LT)
    M(ws, 15, 1, 4, "收货单位及经手人(盖章)：", F(), LT)
    M(ws, 15, 5, 8, "送货单位及经手人(盖章)：", F(), LT)
    copies_note(ws, 5, 13, 9, "①白存根②红客户③黄回单")
    save(ws, "T011-送货单(东莞一彩).xlsx")


# ---------------- T026 东莞一彩送货单（红头） ----------------
def t026():
    ws = new()
    widths(ws, [16, 10, 7, 7, 10, 11, 12])
    M(ws, 1, 1, 5, "东莞市一彩科技有限公司", F(16, True, "FF0000"), LT)
    M(ws, 1, 6, 7, "[送  货  单]", F(12, True), RT)
    M(ws, 2, 1, 5, "广东省东莞市电子科技大厦A2015室", F(9), LT)
    M(ws, 2, 6, 7, "NO.：{order_no}", F(11, True), RT)
    M(ws, 3, 1, 5, "Tel: (0)13652500551  Fax: 0769-83608888", F(9), LT)
    M(ws, 4, 1, 3, "客户名称：{customer}", F(), LT)
    M(ws, 4, 4, 5, "订单号码：{contract_no}", F(), LT)
    M(ws, 4, 6, 7, "制单人员：{handler}", F(), LT)
    M(ws, 5, 1, 3, "客户地址：{address}", F(), LT)
    M(ws, 5, 4, 5, "送货车号：{vehicle_no}", F(), LT)
    M(ws, 5, 6, 7, "送货日期：{delivery_date}", F(), LT)
    head_row(ws, 6, ["产品名称", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 7, ["{items.name}", "{items.spec}", "{items.unit}", "{items.qty}",
                     "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 8, 7, 7)
    M(ws, 15, 1, 3, "合计：{_total_cn}", F(), LT, B)
    W(ws, 15, 4, "{items.qty#sum}", F(), CN, B)
    W(ws, 15, 5, "", F(), CN, B)
    W(ws, 15, 6, "{items.amount#sum}", F(), RT, B)
    W(ws, 15, 7, "", F(), CN, B)
    M(ws, 16, 1, 7, "注：以上货品请核对数量，如有质量问题，请在收货后3天内通知本公司，逾期恕不负责。", F(9), LT)
    M(ws, 17, 1, 3, "送货单位及经手人(盖章)：", F(), LT)
    M(ws, 17, 5, 7, "收货单位及经手人(盖章)：", F(), LT)
    save(ws, "T026-送货单(红头).xlsx")


# ---------------- T035 东莞一彩送货单（条码位） ----------------
def t035():
    ws = new()
    widths(ws, [15, 10, 7, 7, 10, 11, 12])
    M(ws, 1, 1, 5, "东莞市一彩科技有限公司", F(16, True), LT)
    M(ws, 1, 6, 7, "[送 货 单]", F(12, True), RT)
    M(ws, 2, 1, 5, "广东省东莞市电子科技大厦A2015室", F(9), LT)
    M(ws, 2, 6, 7, "NO.：{order_no}", F(11, True), RT)
    M(ws, 3, 1, 5, "Tel: (0)13652500551  Fax: 0769-83608888", F(9), LT)
    M(ws, 4, 1, 3, "客户名称：{customer}", F(), LT)
    M(ws, 4, 4, 5, "合同号码：{contract_no}", F(), LT)
    M(ws, 4, 6, 7, "制单人员：{handler}", F(), LT)
    M(ws, 5, 1, 3, "客户地址：{address}", F(), LT)
    M(ws, 5, 6, 7, "送货日期：{delivery_date}", F(), LT)
    head_row(ws, 6, ["商品名称", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 7, ["{items.name}", "{items.spec}", "{items.unit}", "{items.qty}",
                     "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 8, 6, 7)
    M(ws, 14, 1, 3, "合计：{_total_cn}", F(), LT, B)
    W(ws, 14, 4, "{items.qty#sum}", F(), CN, B)
    W(ws, 14, 5, "", F(), CN, B)
    W(ws, 14, 6, "{items.amount#sum}", F(), RT, B)
    W(ws, 14, 7, "", F(), CN, B)
    M(ws, 15, 1, 3, "送货单位及经手人(盖章)：", F(), LT)
    M(ws, 15, 5, 7, "收货单位及经手人(盖章)：", F(), LT)
    save(ws, "T035-送货单(条码位).xlsx")


# ---------------- T037 东莞一彩送货单（订单编号列） ----------------
def t037():
    ws = new()
    widths(ws, [10, 10, 18, 7, 7, 10, 11, 3.5])
    M(ws, 1, 1, 7, "东莞市一彩科技有限公司", F(16, True), CN)
    M(ws, 2, 1, 7, "广东省东莞市电子科技大厦A2015室  Tel: (0)13652500551  Fax: 0769-83608888", F(9), CN)
    M(ws, 3, 1, 7, "送    货    单", F(13, True), CN)
    M(ws, 4, 1, 3, "收货单位：{customer}", F(), LT)
    M(ws, 4, 4, 5, "NO.：{order_no}", F(11, True), LT)
    M(ws, 4, 6, 7, "送货日期：{delivery_date}", F(), LT)
    head_row(ws, 5, ["订单编号", "产品编号", "名称及规格", "单位", "数量", "单价", "金额"])
    loop_row(ws, 6, ["{items.order_no}", "{items.product_no}", "{items.name}", "{items.unit}",
                     "{items.qty}", "{items.price}", "{items.amount}"])
    blank_rows(ws, 7, 7, 7)
    M(ws, 14, 1, 2, "金计\n人民币", F(11, True), CN, B)
    M(ws, 14, 3, 6, "{_total_cn}", F(11, True), CN, B)
    W(ws, 14, 7, "{items.amount#sum}", F(11, True), RT, B)
    M(ws, 15, 1, 4, "收货单位及经手人(盖章)：", F(), LT)
    M(ws, 15, 5, 7, "送货单位及经手人：", F(), LT)
    copies_note(ws, 5, 14, 8, "①存根(白) ②客户(红) ③仓库(蓝) ④存根(蓝)")
    save(ws, "T037-送货单(订单编号).xlsx")


# ---------------- T055 东莞一彩送货单（序号版） ----------------
def t055():
    ws = new()
    widths(ws, [6, 16, 10, 7, 7, 10, 11, 12])
    M(ws, 1, 1, 6, "东莞市一彩科技有限公司", F(16, True), LT)
    M(ws, 1, 7, 8, "[送  货  单]", F(12, True), RT)
    M(ws, 2, 1, 6, "广东省东莞市电子科技大厦A2015室", F(9), LT)
    M(ws, 2, 7, 8, "NO.：{order_no}", F(11, True), RT)
    M(ws, 3, 1, 6, "Tel: (0)13652500551  Fax: 0769-83608888", F(9), LT)
    M(ws, 4, 1, 4, "订单号码：{contract_no}", F(), LT)
    M(ws, 4, 5, 8, "制单人员：{handler}", F(), LT)
    M(ws, 5, 1, 4, "客户地址：{address}", F(), LT)
    M(ws, 5, 5, 6, "送货车号：{vehicle_no}", F(), LT)
    M(ws, 5, 7, 8, "送货日期：{delivery_date}", F(), LT)
    head_row(ws, 6, ["序号", "产品名称", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 7, ["{items._index}", "{items.name}", "{items.spec}", "{items.unit}",
                     "{items.qty}", "{items.price}", "{items.amount}", "{items.remark}"])
    M(ws, 8, 1, 5, "合计金额(大写)：{_total_cn}", F(), LT, B)
    M(ws, 8, 6, 8, "小写合计：{items.amount#sum}", F(), RT, B)
    M(ws, 9, 1, 4, "收货单位及经手人(盖章)：", F(), LT)
    M(ws, 9, 5, 8, "送货单位及经手人(盖章)：", F(), LT)
    save(ws, "T055-送货单(序号版).xlsx")


# ---------------- T060 定货合同单 ----------------
def t060():
    ws = new()
    widths(ws, [6, 14, 9, 8, 7, 10, 11, 11, 3.5])
    M(ws, 1, 1, 5, "东莞市一彩科技有限公司", F(15, True), LT)
    M(ws, 1, 6, 8, "定货合同单", F(16, True), RT)
    M(ws, 2, 1, 5, "地址：广东省东莞市电子科技大厦A2015室", F(9), LT)
    M(ws, 2, 6, 8, "NO.{order_no}", F(11, True), RT)
    M(ws, 3, 1, 5, "电话：(0)13652500551  传真：0769-83608888", F(9), LT)
    M(ws, 3, 6, 8, "送货日期：{delivery_date}", F(), RT)
    M(ws, 4, 1, 5, "客户名称：{customer}", F(), LT)
    M(ws, 4, 6, 8, "客户电话：{phone}", F(), LT)
    head_row(ws, 5, ["序号", "产品名称", "规格", "颜色", "数量", "单价", "金额", "备注"])
    loop_row(ws, 6, ["{items._index}", "{items.name}", "{items.spec}", "{items.color}",
                     "{items.qty}", "{items.price}", "{items.amount}", "{items.remark}"])
    blank_rows(ws, 7, 8, 8)
    M(ws, 15, 1, 4, "合计：{_total_cn}", F(11, True), LT, B)
    M(ws, 15, 5, 6, "", F(), CN, B)
    M(ws, 15, 7, 8, "{items.amount#sum}", F(11, True), RT, B)
    M(ws, 16, 1, 4, "1、购货方验收产品质量，应收货厂验收为准。\n2、在运输途中破坏家具，我厂一概不承担责任。\n3、自双方签定合同之日起生效。\n4、到取货日期后十天不取货者，订金作废。", F(9), LT, B)
    M(ws, 16, 5, 8, "其它事项", F(), Alignment(horizontal="left", vertical="top"), B)
    ws.row_dimensions[16].height = 56
    M(ws, 17, 1, 2, "主管：", F(), LT)
    M(ws, 17, 3, 5, "购方经手人：", F(), LT)
    M(ws, 17, 6, 8, "卖方经手人：", F(), LT)
    copies_note(ws, 5, 15, 9, "白色存根联 蓝色送货联 红色顾客联 黄色财务联")
    save(ws, "T060-定货合同单.xlsx")


# ---------------- T100 销售单 ----------------
def t100():
    ws = new()
    widths(ws, [6, 15, 10, 7, 7, 10, 11, 11, 3.5])
    M(ws, 1, 1, 6, "东莞市一彩科技有限公司", F(16, True), CN)
    M(ws, 1, 7, 8, "销  售  单", F(16, True), CN)
    M(ws, 2, 1, 6, "电话：(0)13652500551  传真：0769-83608888", F(9), LT)
    M(ws, 2, 7, 8, "SALES LIST", F(11, True), CN)
    M(ws, 3, 1, 6, "地址：广东省东莞市电子科技大厦A2015室", F(9), LT)
    M(ws, 3, 7, 8, "销售单号：{order_no}", F(11, True), RT)
    M(ws, 4, 1, 6, "网址：", F(9), LT)
    M(ws, 4, 7, 8, "开单日期：{delivery_date}", F(), RT)
    M(ws, 5, 1, 3, "提货单位：{customer}", F(), LT)
    M(ws, 5, 4, 5, "联系人：{contact}", F(), LT)
    head_row(ws, 6, ["序号", "名称/型号", "规格", "单位", "数量", "单价", "金额", "备注"])
    loop_row(ws, 7, ["{items._index}", "{items.name}", "{items.spec}", "{items.unit}",
                     "{items.qty}", "{items.price}", "{items.amount}", "{items.remark}"])
    M(ws, 8, 1, 8, "-以下空白-", F(), CN, B)
    blank_rows(ws, 9, 6, 8)
    M(ws, 15, 1, 3, "合计(大写)：{_total_cn}", F(11, True), LT, B)
    M(ws, 15, 4, 5, "小写", F(11, True), CN, B)
    W(ws, 15, 6, "￥{items.amount#sum}", F(11, True), RT, B)
    M(ws, 15, 7, 8, "", F(), CN, B)
    M(ws, 16, 1, 8, "注明：＊本税单为交易时使用，购方代表对以上货物的规格、数量、付款方式确认无误后签字认可，并在货物出库和验收时当面点清。\n＊本税单签字后有效，如有质量异议，请在收货后一周内提出，过期恕不负责。", F(9), LT)
    ws.row_dimensions[16].height = 30
    M(ws, 17, 1, 3, "制单人：{handler}", F(), LT)
    M(ws, 17, 4, 5, "提货车号：{vehicle_no}", F(), LT)
    M(ws, 17, 6, 8, "客户签字：", F(), LT)
    copies_note(ws, 6, 15, 9, "①白存根 ②红客户 ③黄回单")
    save(ws, "T100-销售单.xlsx")


if __name__ == "__main__":
    for fn in [h228, h527, h535, h598, h607, h678, h679, h845, h893,
               t011, t026, t035, t037, t055, t060, t100]:
        fn()
    print("全部完成 →", OUT)
